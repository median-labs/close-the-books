"""The founder review workbook: the one artifact a non-accountant touches.

Everything else in this repo is plumbing. This is the file a founder opens on a
Sunday, and the whole design rests on it being obvious. One sheet, one row per
transaction, one column they fill in.

THE CONTRACT WITH THE FOUNDER
    Nothing in this file posts anything. It is a proposal. They fill in
    `founder_decision`, save, and approve in their own terminal, and then they
    import into QuickBooks themselves. Three separate acts, each one theirs.

WHY row_hash EXISTS
    An approval is an approval of specific decisions, not of a filename. So the
    approval records a hash of the cells that carry a decision: date,
    description, amount, proposed action, account and founder_decision. Change
    any of those after approving and the approval no longer covers the file,
    which `approval.check` catches before anything reaches the import folder.

    The hash covers those six columns and nothing else, deliberately. Hashing
    the file bytes would make an approval die every time Excel rewrote the file
    on open, and a gate that cries wolf is a gate people learn to bypass.
    Widths, colours, the Why column and the Question text can all change
    without invalidating a decision, because none of them changes what gets
    posted.

SECURITY
    Descriptions are bank descriptor text, which is attacker chosen. openpyxl
    writes a leading `=` as a live formula, so every string cell goes through
    `util.csv_safe` first. A proposal flagged `needs_human` (the descriptor
    contained text aimed at an automated system) is marked in the sheet where
    the founder cannot miss it, because that row is exactly the one an
    automated decision should not be trusted on.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from decimal import Decimal

from .util import ZERO, csv_safe, money, parse_date, plain, sha256_of

QBO_DATE = "%m/%d/%Y"
SHEET = "Review"
META_SHEET = "About"

COLUMNS = (
    "Date", "Description", "Amount", "Proposed action", "Account", "Class",
    "Why", "Confidence", "Question", "founder_decision",
)

# The cells an approval is an approval OF. Order matters: it is hashed.
HASHED_COLUMNS = ("Date", "Description", "Amount", "Proposed action", "Account", "founder_decision")

# A blank decision is a decision not yet made, which is why it is in the list.
DECISIONS = ("", "approve", "change account", "skip", "answer")

NEEDS_HUMAN_MARKER = "NEEDS HUMAN"

_COLUMN_KEYS = {
    "Date": "date", "Description": "description", "Amount": "amount",
    "Proposed action": "action", "Account": "account", "Class": "klass",
    "Why": "why", "Confidence": "confidence", "Question": "question",
    "founder_decision": "founder_decision",
}


class ReviewFileError(ValueError):
    """A workbook is not shaped like a review file."""


@dataclass
class ReviewFile:
    path: str
    batch_tag: str
    rows: int
    needs_human: int
    row_hash: str
    header_row: int
    sheet: str = SHEET
    columns: tuple = field(default_factory=lambda: COLUMNS)


# --------------------------------------------------------------- writing

def _instructions(batch_tag: str, meta: dict, rows: int, flagged: int) -> list:
    company = str(meta.get("company", "") or "your books").strip()
    period = str(meta.get("period", "") or "").strip()
    title = f"{company}: {rows} transactions to review"
    if period:
        title += f" ({period})"
    lines = [
        title,
        "Fill in the last column, founder_decision, on every row. Pick from the dropdown.",
        "approve = book it the way it says. change account = type the right account in "
        "the Account column first, then pick change account. skip = leave this one alone. "
        "answer = you typed a reply in the Question column.",
        "Nothing here touches QuickBooks. Save this file, approve the batch in your own "
        f"terminal (batch-{batch_tag}), then import the file yourself.",
    ]
    if flagged:
        lines.append(
            f"{flagged} row(s) are marked {NEEDS_HUMAN_MARKER} in the Question column. "
            f"The bank description on those rows contains text aimed at an automated "
            f"system, so read them yourself before deciding."
        )
    return lines


def _cell_values(proposal, batch_tag: str) -> dict:
    line = proposal.line
    date = parse_date(getattr(line, "date", None), field="proposal date")
    amount = money(getattr(line, "amount", ZERO), "proposal amount")
    account = (getattr(proposal, "account_full", "") or getattr(proposal, "account", "") or "").strip()
    rule_id = (getattr(proposal, "rule_id", "") or "").strip()
    source = (getattr(proposal, "source", "") or "").strip()
    why = ": ".join(p for p in (rule_id, source) if p)
    question = (getattr(proposal, "question", "") or "").strip()
    if getattr(proposal, "needs_human", False):
        detail = question or (
            "the bank description contains text aimed at an automated system. "
            "Read it and decide yourself."
        )
        question = f"{NEEDS_HUMAN_MARKER}: {detail}"
    decision = (getattr(proposal, "founder_decision", "") or "").strip()
    confidence = getattr(proposal, "confidence", 0.0) or 0.0
    return {
        "Date": date.strftime(QBO_DATE),
        "Description": (getattr(line, "descriptor", "") or "").strip(),
        "Amount": amount,
        "Proposed action": getattr(proposal, "action", "question"),
        "Account": account,
        "Class": (getattr(proposal, "klass", "") or "").strip(),
        "Why": why,
        "Confidence": round(float(confidence), 4),
        "Question": question,
        "founder_decision": decision,
    }


def write_review(path, proposals, batch_tag: str, meta: dict = None) -> ReviewFile:
    """Write the review workbook. Returns a ReviewFile describing what landed."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    meta = dict(meta or {})
    batch_tag = str(batch_tag or "").strip()
    if not batch_tag:
        raise ReviewFileError("a review workbook must name its batch")

    values = [_cell_values(p, batch_tag) for p in proposals]
    flagged = sum(1 for p in proposals if getattr(p, "needs_human", False))

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET

    intro = _instructions(batch_tag, meta, len(values), flagged)
    for i, text in enumerate(intro, start=1):
        cell = ws.cell(row=i, column=1, value=csv_safe(text))
        cell.alignment = Alignment(wrap_text=False, vertical="center")
        if i == 1:
            cell.font = Font(bold=True, size=14)
        elif i == len(intro) and flagged:
            cell.font = Font(bold=True, color="9C0006")

    header_row = len(intro) + 2          # one blank row between the block and the table
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for col, name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col, value=csv_safe(name))
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="bottom", wrap_text=False)

    flag_fill = PatternFill("solid", fgColor="FCE4E4")
    money_format = '#,##0.00;[Red]-#,##0.00'
    for i, row_values in enumerate(values):
        r = header_row + 1 + i
        needs_human = bool(getattr(proposals[i], "needs_human", False))
        for col, name in enumerate(COLUMNS, start=1):
            v = row_values[name]
            cell = ws.cell(row=r, column=col)
            if isinstance(v, Decimal):
                cell.value = v
                cell.number_format = money_format
            elif isinstance(v, float):
                cell.value = v
                cell.number_format = "0%"
            else:
                cell.value = csv_safe(v)
            if name == "Question":
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            if needs_human:
                cell.fill = flag_fill

    widths = (12, 46, 13, 16, 32, 16, 34, 11, 48, 18)
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=header_row, column=i).column_letter].width = width

    last = header_row + len(values)
    decision_col = ws.cell(row=header_row, column=len(COLUMNS)).column_letter
    choices = ",".join(d for d in DECISIONS if d)
    dv = DataValidation(type="list", formula1=f'"{choices}"', allow_blank=True)
    dv.errorTitle = "Pick one of the four"
    dv.error = f"Type one of: {choices}. Leave it blank if you have not decided yet."
    dv.promptTitle = "Your decision"
    dv.prompt = "approve, change account, skip or answer."
    ws.add_data_validation(dv)
    if values:
        dv.add(f"{decision_col}{header_row + 1}:{decision_col}{last}")

    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{decision_col}{max(last, header_row)}"

    about = wb.create_sheet(META_SHEET)
    about.append([csv_safe("field"), csv_safe("value")])
    for cell in about[1]:
        cell.font = Font(bold=True)
    about.append([csv_safe("batch"), csv_safe(f"batch-{batch_tag}")])
    about.append([csv_safe("rows"), len(values)])
    about.append([csv_safe(f"rows marked {NEEDS_HUMAN_MARKER}"), flagged])
    for key in sorted(meta):
        about.append([csv_safe(key), csv_safe(meta[key])])
    about.column_dimensions["A"].width = 26
    about.column_dimensions["B"].width = 70

    wb.save(path)
    return ReviewFile(
        path=str(path),
        batch_tag=batch_tag,
        rows=len(values),
        needs_human=flagged,
        row_hash=row_hash(path),
        header_row=header_row,
        sheet=SHEET,
        columns=COLUMNS,
    )


# --------------------------------------------------------------- reading

def _open(path):
    from openpyxl import load_workbook
    # data_only=False: QuickBooks-authored sheets store numbers as literal
    # formulas with no cached value, and data_only=True reads those as None.
    # util.money strips the leading "=" for us.
    wb = load_workbook(path, data_only=False)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    return wb, ws


def _find_header(ws) -> int:
    """The row holding the column names. Scans, because the block above it grows."""
    for r in range(1, min(ws.max_row, 60) + 1):
        first = ws.cell(row=r, column=1).value
        if first is None or str(first).strip() != COLUMNS[0]:
            continue
        names = [str(ws.cell(row=r, column=c).value or "").strip() for c in range(1, len(COLUMNS) + 1)]
        if names == list(COLUMNS):
            return r
    raise ReviewFileError(
        f"{getattr(ws, 'title', '?')} has no review header row. Expected a row "
        f"reading: {', '.join(COLUMNS)}"
    )


def _read_rows(path) -> list:
    wb, ws = _open(path)
    header_row = _find_header(ws)
    out = []
    for r in range(header_row + 1, ws.max_row + 1):
        cells = {name: ws.cell(row=r, column=c).value for c, name in enumerate(COLUMNS, start=1)}
        if all(v is None or str(v).strip() == "" for v in cells.values()):
            continue
        row = {"row": r}
        for name in COLUMNS:
            row[_COLUMN_KEYS[name]] = cells[name]
        out.append(row)
    wb.close()
    return out


def read_decisions(path) -> list:
    """Read the founder's decisions back. One dict per row, in sheet order.

    Values are normalised for a caller that has to act on them: amounts are
    Decimals, dates are `datetime.date`, and `decision` is lowercased and
    stripped. `valid` is False when the founder typed something that is not one
    of the four choices, which happens whenever a spreadsheet drops the data
    validation on copy and paste.
    """
    out = []
    for raw in _read_rows(path):
        decision = str(raw.get("founder_decision") or "").strip()
        normalised = decision.lower()
        question = str(raw.get("question") or "").strip()
        try:
            date = parse_date(raw.get("date"), field=f"row {raw['row']} date")
        except Exception:
            date = None
        out.append({
            "row": raw["row"],
            "date": date,
            "description": str(raw.get("description") or "").strip(),
            "amount": money(raw.get("amount"), f"row {raw['row']} amount"),
            "action": str(raw.get("action") or "").strip(),
            "account": str(raw.get("account") or "").strip(),
            "klass": str(raw.get("klass") or "").strip(),
            "why": str(raw.get("why") or "").strip(),
            "confidence": raw.get("confidence"),
            "question": question,
            "founder_decision": decision,
            "decision": normalised,
            "valid": normalised in DECISIONS,
            "needs_human": question.startswith(NEEDS_HUMAN_MARKER),
        })
    return out


def unrecognized(decisions) -> list:
    """The rows whose founder_decision is not one of the four. For the gate."""
    return [d for d in decisions if not d["valid"]]


def row_hash(path) -> str:
    """sha256 over the decision-bearing cells only. See WHY row_hash EXISTS.

    Stable across an open-and-save with no edits, because every value is
    canonicalised: the amount through `plain(money(...))` so 1234.5 and
    1234.50 agree, the date through `parse_date` so a text date and a real one
    agree, and the decision casefolded so Approve and approve agree.
    """
    parts = []
    for raw in _read_rows(path):
        row = []
        for name in HASHED_COLUMNS:
            key = _COLUMN_KEYS[name]
            v = raw.get(key)
            if name == "Amount":
                row.append(plain(money(v, "row amount")))
            elif name == "Date":
                d = None
                try:
                    d = parse_date(v, field="row date")
                except Exception:
                    d = None
                row.append(d.isoformat() if d else str(v or "").strip())
            elif name == "founder_decision":
                row.append(str(v or "").strip().lower())
            else:
                row.append(str(v or "").strip())
        parts.append("|".join(row))
    return sha256_of("\n".join(parts))
