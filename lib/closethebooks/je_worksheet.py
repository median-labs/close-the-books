"""The typing worksheet: adjusting entries laid out to be entered by hand.

WHY THIS MODULE IS THE PRIMARY OUTPUT IN THE US
    QuickBooks Online US cannot import journal entries. Intuit says so in
    plain words: "Although data can be imported into the US version of
    QuickBooks Online, journal entry import is currently not an option."
    https://quickbooks.intuit.com/learn-support/en-us/reports-and-accounting/importing-journal-entries-in-quickbooks-online-us-version/00/1312709
    verified-on: 2026-09-09  VERIFIED: intuit-doc

    So in the US there is no file to upload. A person opens +New > Journal
    Entry and types each entry, or the company buys a third party importer.
    `je_csv` still matters for QuickBooks Online Canada and UK, where the
    import exists, and for those third party apps. This module is what a US
    founder actually works from.

DESIGNED FOR THE TYPIST, NOT FOR THE ACCOUNTANT
    Every choice here answers a failure that happens at the keyboard, not one
    that happens in the ledger:

    Order of the fields. Each block asks for journal date, then journal no.,
    then the line rows of account / debits / credits / description / name,
    in exactly the order the QuickBooks Journal Entry screen asks for them.
    A worksheet that reads in a different order than the screen makes the
    typist hunt for each value, and hunting is where transpositions come from.

    A running total on every line, and a bold MUST EQUAL line under the block.
    QuickBooks shows its own running total as you type. Two numbers side by
    side catch a transposed 249.00 / 294.00 before the entry is saved, which
    is the difference between fixing a typo and hunting an out-of-balance
    trial balance a week later.

    A tick column. With forty entries the real failure is not arithmetic, it
    is losing your place: a half entered batch, no record of which half, and
    the only way to find out is to read the register. Column A is a `[ ]` per
    entry. Tick it when the entry is saved, not when you start it.

    Repeats grouped together. Twelve monthly amortisations are one entry typed
    once and repeated eleven times by QuickBooks' own "Make recurring" button,
    which most owners have never noticed. Entries that hit the same accounts on
    the same sides are placed next to each other and the group carries a note
    saying to do exactly that. It is the single biggest time saver available to
    someone facing this worksheet.

    A family is grouped on its ACCOUNTS, not on its amounts. A twelve month
    amortisation whose final month carries the rounding stub (11 x 208.33 and
    one 208.37) is one recurring template with one amount to change, not an
    eleven and a one that loses the template for the whole family. The block
    prints every month's amount and marks the members that differ from the
    first, so the person editing the template knows exactly where to change a
    number.

    A count and an honest time estimate at the top, so the person can decide
    whether they have time to start, rather than discovering the answer forty
    minutes in. Next to it, the plain fact that a third party importer app will
    do the same job, so the trade is theirs to make rather than ours to hide.

    Resumability. `read_ticks` reads the tick column back and reports where the
    person stopped: how many are done, which entry is next, and how long the
    rest will take. The realistic state of a forty entry batch is partly done,
    and a second session that cannot see the first one starts from zero or,
    worse, posts something twice.

WHAT IT REFUSES
    The same things `je_csv` refuses, for the same reasons: an entry that
    fails `check()` (unbalanced, no lines, no recorded basis), a line with
    both a debit and a credit or neither, and an account key that does not
    resolve to a real QuickBooks account name. A worksheet is not a lower bar
    than an import file. Somebody is going to type whatever it says into a
    real company's books, and a wrong account name typed by hand creates the
    same new account the import wizard would have offered to create.

SECURITY
    Descriptions carry bank descriptor text, which is attacker chosen. Every
    string cell goes through `util.csv_safe`, because openpyxl writes a
    leading `=` as a live formula. Amounts are written as Decimals into
    number-formatted cells and are never attacker text.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from decimal import Decimal

# `_line_parts` is private to je_csv, and shared on purpose: the worksheet and
# the import CSV must agree about what a line is, down to the error messages.
from .je_csv import UnresolvedAccount, _line_parts, resolve_full_name
from .util import ZERO, csv_safe, money, parse_date, plain

QBO_DATE = "%m/%d/%Y"
SHEET = "Type these in"

# Column A is the tick column. The rest of the block mirrors the QuickBooks
# Journal Entry screen, plus two running-total columns the screen also shows.
COLUMNS = (
    "Done", "Account", "Debits", "Credits", "Description", "Name", "Class",
    "Running debits", "Running credits",
)
TICK_COLUMN = 1
TICK = "[ ]"

MUST_EQUAL = "MUST EQUAL"
RECURRING_NOTE_PREFIX = "Make recurring:"
CLICK_PATH = "+New > Journal Entry"

# Written next to any member of a recurring family whose amount is not the
# template's. Loud, because the failure it prevents is posting the template's
# amount twelve times and quietly overstating one month.
DIFFERS_MARK = "AMOUNT DIFFERS"

# The header line that prices the job, and the alternative to doing it by hand.
# No app is named and no price is quoted, because we have verified neither, and
# a number we cannot stand behind is worse than no number. The decision is the
# owner's; our job is to make sure they know there is one.
COST_PREFIX = "What this costs you:"
IMPORTER_NOTE = (
    "There is another way. Third party importer apps in the QuickBooks App Store post a "
    "journal entry file into QuickBooks Online US for you, which is the same batch without "
    "the typing. We are not naming one and we are not quoting a price we have not checked. "
    "Most are priced per month, and for a one-off catch-up like this a single month is "
    "usually enough. Weigh a month of an app against the hours above and pick whichever "
    "you would rather spend."
)

MONEY_FORMAT = '#,##0.00;[Red]-#,##0.00'

# How long this actually takes, measured as a model rather than a wish.
# An entry costs the fixed price of opening the screen, setting the date and
# the journal number, saving, and waiting for the save; plus a per-line price
# of finding the account in the dropdown, typing the amount and the
# description, and tabbing on. A repeat done through a recurring template
# costs only opening the template, changing the date and saving.
# These are estimates, not measurements. Nobody has timed a real founder.
# verified-on: 2026-09-09  VERIFIED: no
SECONDS_PER_ENTRY = 90
SECONDS_PER_LINE = 45
SECONDS_PER_RECURRING_COPY = 40
# A repeat whose amount is not the template's costs the copy plus finding and
# retyping the amounts. Same standing as the others: an estimate, not a stopwatch.
SECONDS_PER_RECURRING_AMOUNT_EDIT = 20
SECONDS_TO_MAKE_TEMPLATE = 60
BREAK_EVERY = 10
SECONDS_PER_BREAK = 120


class WorksheetError(ValueError):
    """An entry cannot be laid out as something a person can type."""


@dataclass
class WorksheetFile:
    """What landed on disk, and what it will cost the person who opens it."""
    path: str
    entries: int
    lines: int
    groups: int                 # how many kind blocks
    repeating_groups: int       # how many of those are "Make recurring" families
    repeating_entries: int      # entries inside those families
    seconds: int
    estimate: str
    # Of the entries inside recurring families, how many the template covers
    # exactly and how many need an amount changed as well as a date.
    repeating_exact: int = 0
    repeating_needing_edit: int = 0
    batch_tag: str = ""
    sheet: str = SHEET
    columns: tuple = field(default_factory=lambda: COLUMNS)


# ------------------------------------------------------------------ views

def _view(entry, resolve, batch_tag, unresolved) -> dict:
    """One validated, resolved, render-ready entry. Raises unless postable."""
    entry.check()
    number = str(getattr(entry, "number", "") or "").strip()
    if not number:
        raise WorksheetError("an entry has no number; the typist needs one to write down")
    date = parse_date(getattr(entry, "date", None), field=f"entry {number} date")
    tag = (str(getattr(entry, "batch_tag", "") or "").strip()
           or str(batch_tag or "").strip())

    rows, run_d, run_c = [], ZERO, ZERO
    for i, line in enumerate(entry.lines, start=1):
        where = f"entry {number} line {i}"
        acct, debit, credit, memo, name, klass = _line_parts(line, where)
        d, c = money(debit, f"{where} debit"), money(credit, f"{where} credit")
        if d < ZERO or c < ZERO:
            raise WorksheetError(f"{where}: debits and credits are never negative; flip the side")
        if d != ZERO and c != ZERO:
            raise WorksheetError(f"{where}: has both a debit ({d}) and a credit ({c})")
        if d == ZERO and c == ZERO:
            raise WorksheetError(f"{where}: has neither a debit nor a credit")
        try:
            full = resolve_full_name(acct, resolve, where)
        except UnresolvedAccount as exc:
            unresolved.append(str(exc))
            full = ""
        run_d += d
        run_c += c
        text = str(memo or getattr(entry, "memo", "") or "").strip()
        description = f"{text} [batch-{tag}]".strip() if tag else text
        rows.append({
            "account": full,
            "debit": d,
            "credit": c,
            "description": description,
            "name": str(name or "").strip(),
            "class": str(klass or "").strip(),
            "running_debits": run_d,
            "running_credits": run_c,
        })
    return {
        "number": number,
        "date": date,
        "kind": (str(getattr(entry, "kind", "") or "").strip() or "other"),
        "basis": str(getattr(entry, "basis", "") or "").strip(),
        "lines": rows,
        "debits": run_d,
        "credits": run_c,
        "tag": tag,
    }


def _signature(view) -> tuple:
    """What makes two entries "the same entry in a different month".

    The accounts, the side each one is on, and the kind. NOT the amounts, and
    not the date or the memo.

    Amounts used to be part of this, and that was wrong in the most common
    real case there is. A twelve month amortisation of 2,500.00 is eleven
    months of 208.33 and a final 208.37, because the schedule has to add back
    to the whole. Keying on amounts split that into a family of eleven plus a
    stranger, which cost the family its recurring template, cost the person
    the biggest time saver on the worksheet, and did it for four cents.

    What a recurring template actually needs is the shape: the same accounts
    on the same sides. Changing an amount inside a saved template is one field
    and a save. So the family is the shape, and the members whose amounts are
    not the template's are marked instead of exiled.
    """
    return (view["kind"], tuple(
        (l["account"], "D" if l["debit"] != ZERO else "C") for l in view["lines"]
    ))


def _amounts(view) -> tuple:
    """The amounts of one entry, in line order. What the typist must change."""
    return tuple((str(l["debit"]), str(l["credit"])) for l in view["lines"])


# ---------------------------------------------------------------- grouping

def plan(entries, resolve=None, batch_tag="") -> list:
    """Group and order the entries the way they should be typed.

    Returns a list of kind blocks: {"kind", "families"}, where a family is
    {"signature", "views", "repeating"}. Repeating families come first inside
    a kind, largest first, because that is where the recurring template pays.
    """
    unresolved = []
    views = [_view(e, resolve, batch_tag, unresolved) for e in entries]
    if unresolved:
        raise UnresolvedAccount("; ".join(unresolved))

    kinds, order = {}, []
    for v in views:
        kind = v["kind"]
        if kind not in kinds:
            kinds[kind] = []
            order.append(kind)
        kinds[kind].append(v)

    blocks = []
    for kind in order:
        families, seen = {}, []
        for v in kinds[kind]:
            sig = _signature(v)
            if sig not in families:
                families[sig] = []
                seen.append(sig)
            families[sig].append(v)
        packed = []
        for sig in seen:
            members = sorted(families[sig], key=lambda v: (v["date"], v["number"]))
            repeating = len(members) > 1
            template = _amounts(members[0]) if members else ()
            for v in members:
                # A singleton is never "different from the template": there is
                # no template. Marking one would only add noise to a block the
                # person has to read line by line.
                v["differs"] = bool(repeating and _amounts(v) != template)
            needing_edit = sum(1 for v in members if v["differs"])
            packed.append({
                "signature": sig,
                "views": members,
                "repeating": repeating,
                "template_amounts": template,
                "needing_edit": needing_edit,
                "exact": len(members) - needing_edit if repeating else 0,
            })
        # Repeats first, biggest first; singletons keep the order they arrived
        # in, so a re-run of the engine produces the same worksheet.
        packed.sort(key=lambda f: (0 if f["repeating"] else 1, -len(f["views"])))
        blocks.append({"kind": kind, "families": packed})
    return blocks


# ------------------------------------------------------------------ timing

def estimate_seconds(blocks) -> int:
    """A time estimate with a stated basis, not a guess dressed as a number."""
    total, entries = 0, 0
    for block in blocks:
        for family in block["families"]:
            views = family["views"]
            first = views[0]
            total += SECONDS_PER_ENTRY + SECONDS_PER_LINE * len(first["lines"])
            entries += 1
            if family["repeating"]:
                total += SECONDS_TO_MAKE_TEMPLATE
                for v in views[1:]:
                    total += SECONDS_PER_RECURRING_COPY
                    if v["differs"]:
                        total += SECONDS_PER_RECURRING_AMOUNT_EDIT
                entries += len(views) - 1
            else:
                for v in views[1:]:
                    total += SECONDS_PER_ENTRY + SECONDS_PER_LINE * len(v["lines"])
                    entries += 1
    total += SECONDS_PER_BREAK * (max(entries, 1) - 1) // BREAK_EVERY
    return int(total)


def human_duration(seconds: int) -> str:
    minutes = max(1, round(seconds / 60))
    if minutes < 60:
        return f"about {minutes} minutes"
    hours, rest = divmod(minutes, 60)
    if rest == 0:
        return f"about {hours} hour{'s' if hours > 1 else ''}"
    return f"about {hours} hour{'s' if hours > 1 else ''} {rest} minutes"


def _counts(blocks):
    entries = sum(len(f["views"]) for b in blocks for f in b["families"])
    lines = sum(len(v["lines"]) for b in blocks for f in b["families"] for v in f["views"])
    rep_families = [f for b in blocks for f in b["families"] if f["repeating"]]
    rep_entries = sum(len(f["views"]) for f in rep_families)
    needing_edit = sum(f["needing_edit"] for f in rep_families)
    return entries, lines, len(rep_families), rep_entries, needing_edit


def _amount_roll(family) -> str:
    """Every member's date and total, so the amounts are visible before typing."""
    parts = []
    for v in family["views"]:
        stamp = f"{v['date'].strftime(QBO_DATE)} {plain(v['debits'])}"
        parts.append(f"{stamp} (differs)" if v["differs"] else stamp)
    return ", ".join(parts)


def _recurring_note(family, first_index: int, total: int) -> str:
    n = len(family["views"])
    last_index = first_index + n - 1
    edits = family["needing_edit"]
    exact = family["exact"]
    out = (
        f"{RECURRING_NOTE_PREFIX} entries {first_index} to {last_index} of {total} hit the same "
        f"accounts on the same sides, so they are one entry typed {n} times. Type entry "
        f"{first_index}, then click Make recurring at the bottom of the Journal Entry screen and "
        f"save it as a template. Use the template for the other {n - 1}: change the date, save, "
        f"repeat. Most people never find that button, and it is the biggest time saver on this "
        f"worksheet."
    )
    if edits:
        out += (
            f" The template's amounts are right for {exact} of the {n}. The other {edits} "
            f"{'needs' if edits == 1 else 'need'} the amount changed as well as the date, and "
            f"each is marked {DIFFERS_MARK} on its own block; change the amount in the template "
            f"before you save that one."
        )
    else:
        out += f" The template's amounts are right for all {n}: only the date changes."
    out += f" Amounts by date: {_amount_roll(family)}."
    return out


def _intro(blocks, meta, batch_tag, entries, lines, rep_families, rep_entries, estimate,
           needing_edit=0) -> list:
    company = str((meta or {}).get("company", "") or "your books").strip()
    period = str((meta or {}).get("period", "") or "").strip()
    title = f"{company}: {entries} adjusting entries to type into QuickBooks"
    if period:
        title += f" ({period})"
    out = [
        title,
        f"{entries} entries, {lines} lines, {estimate}. Nothing here is posted. You are typing "
        f"these in yourself, one at a time, at {CLICK_PATH}.",
        "QuickBooks Online US has no journal entry import. Intuit: \"journal entry import is "
        "currently not an option.\" That is why this is a worksheet and not a file to upload. "
        "The Canada and UK versions do import, and so do third party importer apps.",
        f"{COST_PREFIX} {entries} entries, {lines} lines, {estimate} of typing, at a keyboard, "
        f"by you. That is the real price of the US having no import.",
        IMPORTER_NOTE,
        "Tick column A as you save each entry, not as you start it. If you stop halfway, the "
        "ticks are the only record of where you stopped, and stopping halfway is the normal "
        "outcome, not the failure case. Save the file with the ticks in it; read_ticks will "
        "tell you where you are and what is left.",
        "Each block gives the journal date first, then the journal no., then the lines, in the "
        "same order the QuickBooks screen asks for them. The running totals should match the "
        "totals QuickBooks shows you as you type. If the MUST EQUAL line and your screen "
        "disagree, a number is wrong; fix it before you save.",
    ]
    if rep_families:
        text = (
            f"{rep_entries} of these entries repeat across {rep_families} group(s), grouped on "
            f"the accounts they hit rather than on their amounts. Type the first of each group, "
            f"then use Make recurring and reuse the template. See the note above each group."
        )
        if needing_edit:
            text += (
                f" {needing_edit} of them need an amount changed as well as a date, and each one "
                f"says so on its own block."
            )
        out.append(text)
    if batch_tag:
        out.append(
            f"Every description ends [batch-{batch_tag}]. Search that in QuickBooks later to see "
            f"exactly the entries this worksheet produced, and void them if you need the batch gone."
        )
    else:
        out.append(
            "No batch tag was set, so these entries will not be findable as one batch later. "
            "Pass batch_tag= if you want that."
        )
    return out


# ----------------------------------------------------------------- writing

def write_worksheet(path, entries, resolve=None, batch_tag="", meta=None) -> WorksheetFile:
    """Write the typing worksheet. Returns a WorksheetFile describing it.

    Nothing is written unless every entry passes, so a refusal never leaves a
    half worksheet behind: a partial worksheet is worse than none, because the
    person typing has no way to know it is partial.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    blocks = plan(entries, resolve=resolve, batch_tag=batch_tag)
    n_entries, n_lines, rep_families, rep_entries, needing_edit = _counts(blocks)
    seconds = estimate_seconds(blocks)
    estimate = human_duration(seconds)

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET

    intro = _intro(blocks, meta, str(batch_tag or "").strip(), n_entries, n_lines,
                   rep_families, rep_entries, estimate, needing_edit)
    r = 1
    for i, text in enumerate(intro, start=1):
        cell = ws.cell(row=r, column=2, value=csv_safe(text))
        cell.alignment = Alignment(wrap_text=False, vertical="center")
        if i == 1:
            cell.font = Font(bold=True, size=14)
        elif "currently not an option" in text or text.startswith(COST_PREFIX):
            cell.font = Font(bold=True, color="9C0006")
        r += 1
    ws.cell(row=r, column=TICK_COLUMN, value=csv_safe(COLUMNS[0])).font = Font(bold=True)
    ws.cell(row=r, column=2, value=csv_safe(
        "Tick each entry as you save it.")).font = Font(bold=True)
    r += 2

    kind_fill = PatternFill("solid", fgColor="DDEBF7")
    head_fill = PatternFill("solid", fgColor="F2F2F2")
    note_font = Font(bold=True, color="1F5C2E")

    index = 0
    for block in blocks:
        n_in_kind = sum(len(f["views"]) for f in block["families"])
        cell = ws.cell(row=r, column=2, value=csv_safe(
            f"{block['kind'].upper()}: {n_in_kind} entr{'y' if n_in_kind == 1 else 'ies'}"))
        cell.font = Font(bold=True, size=12)
        cell.fill = kind_fill
        r += 1

        for family in block["families"]:
            if family["repeating"]:
                note = ws.cell(row=r, column=2,
                               value=csv_safe(_recurring_note(family, index + 1, n_entries)))
                note.font = note_font
                note.alignment = Alignment(wrap_text=True, vertical="top")
                ws.row_dimensions[r].height = 46
                r += 1

            for view in family["views"]:
                index += 1
                tick = ws.cell(row=r, column=TICK_COLUMN, value=csv_safe(TICK))
                tick.font = Font(bold=True)
                tick.alignment = Alignment(horizontal="center")
                label = f"Entry {index} of {n_entries}"
                if view["differs"]:
                    # The one thing a recurring template gets wrong. Said on the
                    # block itself, not only in the note above the family, because
                    # the person typing is looking at the block.
                    label += (f", {DIFFERS_MARK}: this one's amounts are not the "
                              f"template's, change them as well as the date")
                head = ws.cell(row=r, column=2, value=csv_safe(label))
                head.font = Font(bold=True)
                head.fill = head_fill
                if view["differs"]:
                    mark = ws.cell(row=r, column=5, value=csv_safe(DIFFERS_MARK))
                    mark.font = Font(bold=True, color="9C0006")
                r += 1

                # The QuickBooks screen asks for these two first, in this order.
                ws.cell(row=r, column=2, value=csv_safe("Journal date"))
                ws.cell(row=r, column=3, value=csv_safe(view["date"].strftime(QBO_DATE)))
                r += 1
                ws.cell(row=r, column=2, value=csv_safe("Journal no."))
                ws.cell(row=r, column=3, value=csv_safe(view["number"]))
                r += 1

                for col, name in enumerate(COLUMNS[1:], start=2):
                    c = ws.cell(row=r, column=col, value=csv_safe(name))
                    c.font = Font(bold=True)
                    c.fill = head_fill
                r += 1

                for line in view["lines"]:
                    ws.cell(row=r, column=2, value=csv_safe(line["account"]))
                    for col, key in ((3, "debit"), (4, "credit"),
                                     (8, "running_debits"), (9, "running_credits")):
                        amount = line[key]
                        if col in (3, 4) and amount == ZERO:
                            continue
                        c = ws.cell(row=r, column=col, value=amount)
                        c.number_format = MONEY_FORMAT
                    d = ws.cell(row=r, column=5, value=csv_safe(line["description"]))
                    d.alignment = Alignment(wrap_text=True, vertical="top")
                    ws.cell(row=r, column=6, value=csv_safe(line["name"]))
                    ws.cell(row=r, column=7, value=csv_safe(line["class"]))
                    r += 1

                label = ws.cell(row=r, column=2, value=csv_safe(MUST_EQUAL))
                label.font = Font(bold=True)
                for col, amount in ((3, view["debits"]), (4, view["credits"])):
                    c = ws.cell(row=r, column=col, value=amount)
                    c.number_format = MONEY_FORMAT
                    c.font = Font(bold=True)
                note = ws.cell(row=r, column=5, value=csv_safe(
                    f"{plain(view['debits'])} debits = {plain(view['credits'])} credits. "
                    f"If your screen says anything else, a number is wrong. Fix it before saving."))
                note.font = Font(bold=True)
                r += 1

                if view["basis"]:
                    why = ws.cell(row=r, column=2, value=csv_safe(f"Why: {view['basis']}"))
                    why.alignment = Alignment(wrap_text=True, vertical="top")
                    r += 1
                r += 1     # one blank row between entries

    widths = (7, 42, 13, 13, 46, 18, 16, 15, 15)
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width

    wb.save(path)
    return WorksheetFile(
        path=str(path),
        entries=n_entries,
        lines=n_lines,
        groups=len(blocks),
        repeating_groups=rep_families,
        repeating_entries=rep_entries,
        seconds=seconds,
        estimate=estimate,
        repeating_exact=rep_entries - needing_edit,
        repeating_needing_edit=needing_edit,
        batch_tag=str(batch_tag or "").strip(),
    )


# -------------------------------------------------------------- resuming

# What the reader has to find again in a saved worksheet. These shapes are
# written by `write_worksheet` above; if one changes, change it in both places
# or a resumed session silently reports zero done.
ENTRY_HEAD_RE = re.compile(r"^Entry (\d+) of (\d+)\b")
KIND_HEAD_RE = re.compile(r"^([A-Z][A-Z0-9 _/&+.-]*): \d+ entr(?:y|ies)$")
FAMILY_RANGE_RE = re.compile(r"entries (\d+) to (\d+) of \d+")

# A tick cell holding any of these is untyped. Anything else a person has put
# there ("x", "[x]", "done", a date) means they saved that entry.
UNTICKED = {"", TICK, "[]", "[ ]", "()", "( )", "-"}


@dataclass
class Progress:
    """Where a part-typed worksheet actually stands.

    The realistic state of a forty entry batch is neither nothing nor
    everything. Without this, a second session either starts from the top,
    which double posts, or asks the person to remember, which they cannot.
    """
    path: str
    total: int
    done: int
    remaining: int
    next_index: int = 0            # 1-based position in the worksheet, 0 when finished
    next_number: str = ""          # the journal no. of the next entry
    next_label: str = ""           # a sentence naming it: kind, date, first account
    next_date: str = ""
    seconds_remaining: int = 0
    estimate: str = ""             # human duration of what is left
    unmarked: int = 0              # ticks somebody blanked rather than marked
    entries: list = field(default_factory=list)
    sheet: str = SHEET

    @property
    def finished(self) -> bool:
        return self.total > 0 and self.remaining == 0

    def line(self) -> str:
        """One sentence, for the top of a resumed session."""
        if self.total == 0:
            return f"{self.path} has no entry blocks in it."
        if self.finished:
            return (f"All {self.total} entries in {self.path} are ticked. Nothing left to "
                    f"type. Ticks are the person's own record; they are not proof the "
                    f"entries are in QuickBooks, so check the register before you rely on it.")
        head = (f"You are {self.done} of {self.total} in, {self.remaining} left, "
                f"{self.estimate} remaining. Next is {self.next_label}.")
        if self.unmarked:
            head += (f" {self.unmarked} tick box(es) are blank rather than ticked; those are "
                     f"counted as not done, because typing one twice is cheaper to find than "
                     f"never typing it at all.")
        return head


def _is_ticked(value) -> bool:
    """Blank counts as NOT done, deliberately.

    Both errors cost something. Calling a typed entry untyped risks a double
    post, which shows up in the register and in the trial balance. Calling an
    untyped entry typed leaves an adjusting entry that never gets made, in
    books that now look finished. The second error is the one this whole
    module exists to prevent, so ambiguity resolves towards not done.
    """
    s = "" if value is None else str(value).strip()
    if s.startswith("'"):                      # the csv_safe escape, if present
        s = s[1:].strip()
    return s.lower() not in {u.lower() for u in UNTICKED}


def _is_number(value) -> bool:
    return isinstance(value, (int, float, Decimal)) and not isinstance(value, bool)


def read_ticks(path) -> Progress:
    """Read a saved worksheet's tick column back and say where the person is.

    Reads only what `write_worksheet` wrote: the tick column, the entry heads,
    the journal date and no., the account lines, and the recurring family
    notes. It re-prices what is LEFT with the same model that priced the whole
    job, so "45 minutes remaining" comes from the same arithmetic as the
    "2 hours" the person read before they started.
    """
    from openpyxl import load_workbook

    wb = load_workbook(path)
    try:
        if SHEET not in wb.sheetnames:
            raise WorksheetError(
                f"{path} has no sheet called {SHEET!r}, so it is not one of these "
                f"worksheets. Sheets found: {', '.join(wb.sheetnames)}."
            )
        ws = wb[SHEET]
        rows = [[c.value for c in row] for row in ws.iter_rows()]
    finally:
        wb.close()

    items, kind, family = [], "", None
    current = None
    for row in rows:
        row = list(row) + [None] * (9 - len(row))
        a = row[0]
        b = "" if row[1] is None else str(row[1]).strip()
        if b.startswith("'"):
            b = b[1:]

        m = ENTRY_HEAD_RE.match(b)
        if m:
            index = int(m.group(1))
            current = {
                "index": index,
                "done": _is_ticked(a),
                "blank": (a is None or str(a).strip() == ""),
                "kind": kind,
                "number": "",
                "date": "",
                "account": "",
                "lines": 0,
                "differs": DIFFERS_MARK in b,
                "family": None,
                "is_template": False,
            }
            if family and family[0] <= index <= family[1]:
                current["family"] = family
                current["is_template"] = (index == family[0])
            items.append(current)
            continue

        if KIND_HEAD_RE.match(b):
            kind = KIND_HEAD_RE.match(b).group(1).strip().lower()
            family = None                      # a family never spans two kinds
            continue

        if b.startswith(RECURRING_NOTE_PREFIX):
            r = FAMILY_RANGE_RE.search(b)
            family = (int(r.group(1)), int(r.group(2))) if r else None
            continue

        if current is None:
            continue
        if b == "Journal date":
            current["date"] = "" if row[2] is None else str(row[2]).strip()
        elif b == "Journal no.":
            current["number"] = "" if row[2] is None else str(row[2]).strip()
        elif b == MUST_EQUAL:
            current = None                     # the block is over
        elif b and _is_number(row[7]):
            # Only line rows carry a running total, which is what tells a line
            # apart from a label without depending on where it sits.
            current["lines"] += 1
            if not current["account"]:
                current["account"] = b

    total = len(items)
    done = sum(1 for i in items if i["done"])
    left = [i for i in items if not i["done"]]
    seconds = _remaining_seconds(left)
    prog = Progress(
        path=str(path), total=total, done=done, remaining=len(left),
        seconds_remaining=seconds,
        estimate=human_duration(seconds) if left else "nothing",
        unmarked=sum(1 for i in items if i["blank"] and not i["done"]),
        entries=items,
    )
    if left:
        nxt = left[0]
        prog.next_index = nxt["index"]
        prog.next_number = nxt["number"]
        prog.next_date = nxt["date"]
        bits = [f"entry {nxt['index']} of {total}"]
        if nxt["number"]:
            bits.append(nxt["number"])
        if nxt["kind"]:
            bits.append(nxt["kind"])
        if nxt["date"]:
            bits.append(nxt["date"])
        if nxt["account"]:
            bits.append(nxt["account"])
        label = ", ".join(bits)
        if nxt["differs"]:
            label += f" ({DIFFERS_MARK}: change the amount, not only the date)"
        prog.next_label = label
    return prog


def _remaining_seconds(left) -> int:
    """Price what is left with the model that priced the whole job.

    A repeat still costs a repeat: a member of a recurring family that is not
    the family's first entry is a template copy whether or not the first one
    has been typed yet, because by the time the person reaches it, it has.
    """
    total = 0
    for item in left:
        if item["family"] and not item["is_template"]:
            total += SECONDS_PER_RECURRING_COPY
            if item["differs"]:
                total += SECONDS_PER_RECURRING_AMOUNT_EDIT
            continue
        total += SECONDS_PER_ENTRY + SECONDS_PER_LINE * max(item["lines"], 1)
        if item["family"] and item["is_template"]:
            total += SECONDS_TO_MAKE_TEMPLATE
    total += SECONDS_PER_BREAK * (max(len(left), 1) - 1) // BREAK_EVERY
    return int(total)


# --------------------------------------------------------------- markdown

def render_markdown(entries, resolve=None, batch_tag="", meta=None) -> str:
    """The same worksheet as text, for a terminal, a PR or a printed page.

    Checkboxes rather than a tick column, because that is what a tick column
    is in markdown, and because a founder who works in a text editor should
    not have to open a spreadsheet to type forty entries.
    """
    blocks = plan(entries, resolve=resolve, batch_tag=batch_tag)
    n_entries, n_lines, rep_families, rep_entries, needing_edit = _counts(blocks)
    estimate = human_duration(estimate_seconds(blocks))
    intro = _intro(blocks, meta, str(batch_tag or "").strip(), n_entries, n_lines,
                   rep_families, rep_entries, estimate, needing_edit)

    out = [f"# {intro[0]}", ""]
    for text in intro[1:]:
        out.append(text)
        out.append("")

    index = 0
    for block in blocks:
        n_in_kind = sum(len(f["views"]) for f in block["families"])
        out.append(f"## {block['kind'].upper()}: {n_in_kind} "
                   f"entr{'y' if n_in_kind == 1 else 'ies'}")
        out.append("")
        for family in block["families"]:
            if family["repeating"]:
                out.append(f"> {_recurring_note(family, index + 1, n_entries)}")
                out.append("")
            for view in family["views"]:
                index += 1
                head = f"- [ ] **Entry {index} of {n_entries}**"
                if view["differs"]:
                    head += (f" {DIFFERS_MARK}: this one's amounts are not the template's, "
                             f"change them as well as the date.")
                out.append(head)
                out.append("")
                out.append(f"  Journal date: {view['date'].strftime(QBO_DATE)}  ")
                out.append(f"  Journal no.: {view['number']}")
                out.append("")
                out.append("  | Account | Debits | Credits | Description | Name | Class "
                           "| Running debits | Running credits |")
                out.append("  |---|---|---|---|---|---|---|---|")
                for line in view["lines"]:
                    out.append(
                        "  | {account} | {debit} | {credit} | {description} | {name} | {klass} "
                        "| {rd} | {rc} |".format(
                            account=line["account"],
                            debit=plain(line["debit"]) if line["debit"] != ZERO else "",
                            credit=plain(line["credit"]) if line["credit"] != ZERO else "",
                            description=line["description"].replace("|", "\\|"),
                            name=line["name"], klass=line["class"],
                            rd=plain(line["running_debits"]), rc=plain(line["running_credits"]),
                        ))
                out.append("")
                out.append(f"  **{MUST_EQUAL}: {plain(view['debits'])} debits = "
                           f"{plain(view['credits'])} credits.** If your screen says anything "
                           f"else, a number is wrong. Fix it before saving.")
                if view["basis"]:
                    out.append("")
                    out.append(f"  Why: {view['basis']}")
                out.append("")
    return "\n".join(out).rstrip() + "\n"
