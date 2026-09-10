"""Tests for qbo_exports, against synthetic workbooks that plant every trap.

The fixtures are generated here with openpyxl rather than checked in, so the
shape of a QuickBooks export is stated in code where it can be read and argued
with. Every company, account and vendor name is invented.

Runs two ways, with no test dependency required:
    python3 -m pytest tests/ -q
    python3 tests/test_qbo_exports.py
"""

from __future__ import annotations

import datetime
import os
import sys
import tempfile
from decimal import Decimal

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "lib"))

import openpyxl  # noqa: E402

from closethebooks.model import Account  # noqa: E402
from closethebooks.qbo_exports import (  # noqa: E402
    ExportError,
    discover,
    load_account_list,
    load_all,
    load_general_ledger,
    load_journal,
    load_profit_and_loss_by_month,
    load_trial_balance,
    report_face,
)

D = Decimal

# Every test takes one argument, `tmp`, a directory to write fixtures into.
# Under pytest that is this fixture; under the plain runner at the bottom of
# the file it is a TemporaryDirectory. Neither runner is required to be
# installed for the other to work.
try:
    import pytest

    @pytest.fixture(name="tmp")
    def _tmp_dir(tmp_path):
        return str(tmp_path)

except ImportError:                     # pragma: no cover - pytest is optional
    pytest = None


# --------------------------------------------------------------- fixtures

def write_sheet(path, rows, sheet_title="Sheet1"):
    """Write rows to an xlsx exactly as given.

    A string beginning "=" is stored by openpyxl as a formula with no cached
    value, which is precisely what QuickBooks produces and precisely why
    data_only=True reads these files as empty.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    for row in rows:
        ws.append(list(row))
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    return path


FOOTER_TB = "Monday, Mar 02, 2026 12:01:05 PM GMT-8 - Accrual Basis"
FOOTER_STD = "Accrual Basis Monday, March 02, 2026 08:01 PM GMTZ"
FOOTER_NO_BASIS = " Monday, March 02, 2026 08:02 PM GMTZ"


def trial_balance_rows(expense="=4500.00", total_debit=None, total_credit=None):
    """A trial balance whose every amount is a literal formula string.

    Face layout is the trial balance's own: company first, title second. That
    is the opposite of every other QuickBooks report and the parser has to cope
    with both.
    """
    total_debit = total_debit or "=B6+B7+B8+B9+B10+B11"
    total_credit = total_credit or "=C6+C7+C8+C9+C10+C11"
    return [
        ["Acme Robotics Inc.", "", ""],
        ["Trial Balance", "", ""],
        ["As of December 31, 2025", "", ""],
        ["", "", ""],
        ["", "Debit", "Credit"],
        ["1000 Assets:Cash", "=15000.00", ""],
        ["1200 Assets:Accounts Receivable", "=2500.00", ""],
        ["2000 Liabilities:Accounts Payable", "", "=4000.00"],
        ["3000 Equity:Common Stock", "", "=10000.00"],
        ["4000 Income:Product Revenue", "", "=8000.00"],
        ["6000 Expenses:Contractor Expense", expense, ""],
        ["TOTAL", total_debit, total_credit],
        ["", "", ""],
        [FOOTER_TB, "", ""],
    ]


# A general ledger in the single signed `Amount` layout. Amounts are in each
# account's own natural direction: positive is a debit on Cash and a CREDIT on
# Product Revenue. Only the chart can tell them apart.
GL_AMOUNT_ROWS = [
    ["General Ledger", "", "", "", "", "", "", "", ""],
    ["Acme Robotics Inc.", "", "", "", "", "", "", "", ""],
    ["January-December, 2025", "", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", "", ""],
    ["", "Distribution account", "Transaction date", "Transaction type", "Num",
     "Name", "Memo/Description", "Amount", "Balance"],
    ["1000 Cash", "", "", "", "", "", "", "", ""],
    ["", "Beginning Balance", "", "", "", "", "", "", 1000.00],
    ["", "1000 Cash", "03/04/2025", "Deposit", "", "Northwind Supply", "invoice 4471", 5000.00, 6000.00],
    ["", "1000 Cash", "11/12/2025", "Expense", "", "Northwind Supply", "workshop supplies", -1200.00, 4800.00],
    ["Total for 1000 Cash", "", "", "", "", "", "", "=H8+H9", ""],
    ["4000 Product Revenue", "", "", "", "", "", "", "", ""],
    ["", "Beginning Balance", "", "", "", "", "", "", 0.00],
    ["", "4000 Product Revenue", "03/04/2025", "Deposit", "", "Northwind Supply", "invoice 4471", 5000.00, 5000.00],
    ["Total for 4000 Product Revenue", "", "", "", "", "", "", "=H13", ""],
    ["6000 Contractor Expense", "", "", "", "", "", "", "", ""],
    ["", "Beginning Balance", "", "", "", "", "", "", 0.00],
    ["", "6000 Contractor Expense", "11/12/2025", "Expense", "", "Northwind Supply", "workshop supplies", 1200.00, 1200.00],
    ["Total for 6000 Contractor Expense", "", "", "", "", "", "", "=H17", ""],
    ["", "", "", "", "", "", "", "", ""],
    [FOOTER_STD, "", "", "", "", "", "", "", ""],
]

# The same ledger in the other layout: explicit Debit and Credit columns, which
# are already debit-positive and must NOT be flipped.
GL_DEBIT_CREDIT_ROWS = [
    ["General Ledger", "", "", "", "", "", "", "", ""],
    ["Acme Robotics Inc.", "", "", "", "", "", "", "", ""],
    ["January-December, 2025", "", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", "", ""],
    ["", "Distribution account", "Transaction date", "Transaction type", "Num",
     "Name", "Memo/Description", "Debit", "Credit"],
    ["1000 Cash", "", "", "", "", "", "", "", ""],
    ["", "Beginning Balance", "", "", "", "", "", "", ""],
    ["", "1000 Cash", "03/04/2025", "Deposit", "", "Northwind Supply", "invoice 4471", 5000.00, ""],
    ["", "1000 Cash", "11/12/2025", "Expense", "", "Northwind Supply", "workshop supplies", "", 1200.00],
    ["Total for 1000 Cash", "", "", "", "", "", "", "=H8+H9", "=I8+I9"],
    ["4000 Product Revenue", "", "", "", "", "", "", "", ""],
    ["", "4000 Product Revenue", "03/04/2025", "Deposit", "", "Northwind Supply", "invoice 4471", "", 5000.00],
    ["Total for 4000 Product Revenue", "", "", "", "", "", "", "=H12", "=I12"],
    ["6000 Contractor Expense", "", "", "", "", "", "", "", ""],
    ["", "6000 Contractor Expense", "11/12/2025", "Expense", "", "Northwind Supply", "workshop supplies", 1200.00, ""],
    ["Total for 6000 Contractor Expense", "", "", "", "", "", "", "=H15", "=I15"],
    ["", "", "", "", "", "", "", "", ""],
    [FOOTER_STD, "", "", "", "", "", "", "", ""],
]

JOURNAL_ROWS = [
    ["Journal", "", "", "", "", "", "", "", "", ""],
    ["Acme Robotics Inc.", "", "", "", "", "", "", "", "", ""],
    ["January-December, 2025", "", "", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", "", "", ""],
    ["", "Transaction date", "Transaction type", "Num", "Name", "Memo/Description",
     "Distribution account number", "Account full name", "Debit", "Credit"],
    ["1001", "", "", "", "", "", "", "", "", ""],
    ["", "03/04/2025", "Deposit", "", "Northwind Supply", "invoice 4471", "1000", "1000 Assets:Cash", 5000.00, ""],
    ["", "03/04/2025", "Deposit", "", "Northwind Supply", "invoice 4471", "4000", "4000 Income:Product Revenue", "", 5000.00],
    ["Total for 1001", "", "", "", "", "", "", "", "=I7+I8", "=J7+J8"],
    ["1002", "", "", "", "", "", "", "", "", ""],
    ["", "11/12/2025", "Expense", "", "Northwind Supply", "workshop supplies", "6000", "6000 Expenses:Contractor Expense", 1200.00, ""],
    ["", "11/12/2025", "Expense", "", "Northwind Supply", "workshop supplies", "1000", "1000 Assets:Cash", "", 1200.00],
    ["Total for 1002", "", "", "", "", "", "", "", "=I11+I12", "=J11+J12"],
    ["", "", "", "", "", "", "", "", "", ""],
    [FOOTER_NO_BASIS, "", "", "", "", "", "", "", "", ""],
]

ACCOUNT_LIST_ROWS = [
    ["Account List", "", "", "", "", ""],
    ["Acme Robotics Inc.", "", "", "", "", ""],
    ["", "", "", "", "", ""],
    ["Account #", "Full name", "Type", "Detail type", "Description", "Total balance"],
    ["1000", "Assets:Cash", "Bank", "Checking", "", 15000.00],
    ["2000", "Liabilities:Accounts Payable", "Accounts Payable (A/P)", "", "", -4000.00],
    ["4000", "Income:Product Revenue", "Income", "Sales of Product Income", "", 0.00],
    ["6000", "Expenses:Contractor Expense", "Expenses", "Legal & Professional Fees", "", 4500.00],
    ["", "Petty Cash", "Bank", "Cash on hand", "", ""],
    ["3300", "Assets:Prepaid Insurance", "Other Current Assets", "Prepaid Expenses", "", 0.00],
    ["TOTAL", "", "", "", "", 15500.00],
    ["", "", "", "", "", ""],
    [FOOTER_NO_BASIS, "", "", "", "", ""],
]


def pl_by_month_rows(revenue_total="=B7+C7+D7"):
    return [
        ["Profit and Loss by Month", "", "", "", ""],
        ["Acme Robotics Inc.", "", "", "", ""],
        ["January-March, 2025", "", "", "", ""],
        ["", "", "", "", ""],
        ["Distribution account", "January 2025", "February 2025", "March 2025", "Total"],
        ["Income", "", "", "", ""],
        ["4000 Product Revenue", 1000.00, 2000.00, 3000.00, revenue_total],
        ["Total for Income", "=B7", "=C7", "=D7", "=B8+C8+D8"],
        ["Expenses", "", "", "", ""],
        ["6000 Contractor Expense", 100.00, 200.00, 300.00, "=B10+C10+D10"],
        ["Total for Expenses", "=B10", "=C10", "=D10", "=B11+C11+D11"],
        ["Net Income", "=B8-B11", "=C8-C11", "=D8-D11", "=B12+C12+D12"],
        ["", "", "", "", ""],
        ["Accrual Basis Monday, March 02, 2026 08:04 PM GMTZ", "", "", "", ""],
    ]


CHART = [
    Account(name="Cash", number="1000", full_name="Assets:Cash", type="Bank"),
    Account(name="Product Revenue", number="4000", full_name="Income:Product Revenue", type="Income"),
    Account(name="Contractor Expense", number="6000",
            full_name="Expenses:Contractor Expense", type="Expenses"),
]


# ------------------------------------------------------------------- tests

def test_formula_strings_are_read_and_data_only_true_would_zero_them(tmp):
    """TRAP 1. The reason data_only must stay False.

    Asserts both halves: that the loader reads the real figures, and that the
    naive reading of the very same file returns nothing at all. The second half
    is what stops a future maintainer "simplifying" the flag back.
    """
    path = write_sheet(os.path.join(tmp, "tb.xlsx"), trial_balance_rows(), "Trial Balance")

    tb = load_trial_balance(path)
    assert tb.total_debits == D("22000.00"), tb.total_debits
    assert tb.total_credits == D("22000.00"), tb.total_credits

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    cached = [ws.cell(row=r, column=2).value for r in range(6, 12)]
    wb.close()
    assert all(v is None for v in cached), (
        f"fixture no longer reproduces the trap: data_only=True returned {cached}"
    )

    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb[wb.sheetnames[0]]
    assert ws.cell(row=6, column=2).value == "=15000.00"
    wb.close()


def test_trial_balance_face_is_captured_verbatim(tmp):
    """TRAP 7. period_text is the evidence the date filter applied."""
    path = write_sheet(os.path.join(tmp, "tb.xlsx"), trial_balance_rows(), "Trial Balance")
    tb = load_trial_balance(path)
    assert tb.face.company == "Acme Robotics Inc."
    assert tb.face.title == "Trial Balance"
    assert tb.face.period_text == "As of December 31, 2025"   # verbatim
    assert tb.as_of == datetime.date(2025, 12, 31)
    assert tb.basis == "accrual"
    assert tb.face.generated_at == datetime.date(2026, 3, 2)


def test_face_handles_both_row_orders(tmp):
    """The trial balance puts the company first; every other report puts the
    title first. Identifying the title by vocabulary is what makes both work."""
    tb = report_face(write_sheet(os.path.join(tmp, "tb.xlsx"), trial_balance_rows(), "Trial Balance"))
    gl = report_face(write_sheet(os.path.join(tmp, "gl.xlsx"), GL_AMOUNT_ROWS))
    assert (tb.title, tb.company) == ("Trial Balance", "Acme Robotics Inc.")
    assert (gl.title, gl.company) == ("General Ledger", "Acme Robotics Inc.")
    assert gl.period_text == "January-December, 2025"
    assert gl.period_start == datetime.date(2025, 1, 1)
    assert gl.period_end == datetime.date(2025, 12, 31)


def test_out_of_balance_trial_balance_raises_naming_both_totals(tmp):
    """TRAP 4. A hard gate, not a warning."""
    rows = trial_balance_rows(expense="=4600.00")
    path = write_sheet(os.path.join(tmp, "bad.xlsx"), rows, "Trial Balance")
    try:
        load_trial_balance(path)
    except ExportError as exc:
        msg = str(exc)
        assert "22100.00" in msg and "22000.00" in msg, msg
        assert "100.00" in msg, msg
    else:
        raise AssertionError("an out-of-balance trial balance must raise ExportError")


def test_trial_balance_must_match_its_own_printed_total(tmp):
    """TRAP 8. Equal-but-wrong is real: dropping a section loses a debit and
    its credit together and the remainder still balances."""
    rows = trial_balance_rows(total_debit="=99999.00", total_credit="=99999.00")
    path = write_sheet(os.path.join(tmp, "mismatch.xlsx"), rows, "Trial Balance")
    try:
        load_trial_balance(path)
    except ExportError as exc:
        assert "printed total" in str(exc), str(exc)
    else:
        raise AssertionError("a parse disagreeing with the printed total must raise")


def test_trial_balance_columns_are_found_by_header_not_index(tmp):
    """TRAP 3. Right-aligned columns: same report, shifted two to the right."""
    rows = [
        ["Acme Robotics Inc.", "", "", "", ""],
        ["Trial Balance", "", "", "", ""],
        ["As of December 31, 2025", "", "", "", ""],
        ["", "", "", "", ""],
        ["", "", "", "Debit", "Credit"],
        ["1000 Assets:Cash", "", "", "=15000.00", ""],
        ["2000 Liabilities:Accounts Payable", "", "", "", "=15000.00"],
        ["TOTAL", "", "", "=D6+D7", "=E6+E7"],
        ["", "", "", "", ""],
        [FOOTER_TB, "", "", "", ""],
    ]
    tb = load_trial_balance(write_sheet(os.path.join(tmp, "shifted.xlsx"), rows, "Trial Balance"))
    assert len(tb.rows) == 2
    assert tb.total_debits == D("15000.00") == tb.total_credits
    assert tb.foots


def test_general_ledger_amount_layout_skips_subtotals_and_reads_dates(tmp):
    """TRAPS 2, 5 and 6 in one file."""
    path = write_sheet(os.path.join(tmp, "gl.xlsx"), GL_AMOUNT_ROWS)
    gl = load_general_ledger(path, accounts=CHART)
    assert len(gl) == 4, [l.account for l in gl]        # 3 sections, 3 beginnings, 3 totals all skipped
    assert not any("total" in l.account.lower() for l in gl)
    assert not any("beginning" in (l.memo or "").lower() for l in gl)
    # MM/DD/YYYY, never sliced and never read day-first
    assert gl[0].date == datetime.date(2025, 3, 4)
    assert gl[1].date == datetime.date(2025, 11, 12)
    assert gl.beginning_balances["1000"] == D("1000.00")


def test_general_ledger_amount_is_natural_sign_not_debit_positive(tmp):
    """The costliest trap, and the one the brief did not name.

    A QuickBooks General Ledger signs its Amount in each account's own
    direction: on Product Revenue a positive Amount is a credit. Without the
    chart the set cannot balance and says so; with it, it balances exactly.
    """
    path = write_sheet(os.path.join(tmp, "gl.xlsx"), GL_AMOUNT_ROWS)

    raw = load_general_ledger(path)
    assert raw.sign_basis == "report"
    assert not raw.balanced, "the unflipped ledger is expected not to balance"
    assert raw.notes and "no chart" in raw.notes[0]

    fixed = load_general_ledger(path, accounts=CHART)
    assert fixed.sign_basis == "debit_positive"
    assert fixed.balanced, f"{fixed.total_debits} vs {fixed.total_credits}"
    assert fixed.total_debits == D("6200.00") == fixed.total_credits
    assert fixed.foots
    revenue = [l for l in fixed if l.account == "4000"][0]
    assert revenue.credit == D("5000.00") and revenue.debit == D("0.00")


def test_general_ledger_debit_credit_layout_is_not_flipped(tmp):
    """TRAP 5, the other layout. Explicit sides are already debit-positive."""
    path = write_sheet(os.path.join(tmp, "gl2.xlsx"), GL_DEBIT_CREDIT_ROWS)
    gl = load_general_ledger(path, accounts=CHART)
    assert len(gl) == 4
    assert gl.total_debits == D("6200.00") == gl.total_credits
    assert gl.foots
    revenue = [l for l in gl if l.account == "4000"][0]
    assert revenue.credit == D("5000.00")
    # and it reads the same without a chart, because nothing needs flipping
    assert load_general_ledger(path).balanced


def test_general_ledger_reports_a_section_that_does_not_tie(tmp):
    """TRAP 8, recorded rather than raised, so every bad section shows at once."""
    rows = [list(r) for r in GL_AMOUNT_ROWS]
    rows[9][7] = "=H8"          # the Cash subtotal now omits a row
    path = write_sheet(os.path.join(tmp, "gl_bad.xlsx"), rows)
    gl = load_general_ledger(path, accounts=CHART)
    assert len(gl) == 4, "a bad subtotal must not change what is parsed"
    assert gl.discrepancies, "a section that does not tie must be recorded"
    assert not gl.foots
    assert "1200.00" in gl.discrepancies[0], gl.discrepancies


def test_journal_groups_balance_and_reference_formulas_resolve(tmp):
    """The Journal's subtotals are "=I7+I8" cell references, not numbers.

    util.money cannot read one, so the loader resolves them against the sheet
    it already holds. That is the only way to check a group actually tied.
    """
    path = write_sheet(os.path.join(tmp, "journal.xlsx"), JOURNAL_ROWS)
    j = load_journal(path)
    assert len(j) == 4
    assert sorted({l.txn_id for l in j}) == ["1001", "1002"]
    assert j.total_debits == D("6200.00") == j.total_credits
    assert j.foots and not j.discrepancies
    assert sorted({l.account for l in j}) == ["1000", "4000", "6000"]
    assert j[0].name == "Northwind Supply"
    assert j.face.basis == "", "this report carries no basis and must not invent one"


def test_journal_group_that_does_not_tie_is_recorded(tmp):
    rows = [list(r) for r in JOURNAL_ROWS]
    rows[8][8] = "=99.00"       # group 1001's printed debit no longer matches
    path = write_sheet(os.path.join(tmp, "journal_bad.xlsx"), rows)
    j = load_journal(path)
    assert len(j) == 4
    assert j.discrepancies and "1001" in j.discrepancies[0], j.discrepancies
    assert "99.00" in j.discrepancies[0]
    assert not j.foots


def test_account_list_reads_numbers_types_and_balances(tmp):
    path = write_sheet(os.path.join(tmp, "accounts.xlsx"), ACCOUNT_LIST_ROWS)
    accounts = load_account_list(path)
    assert len(accounts) == 6, [a.full_name for a in accounts]   # the TOTAL row is not an account
    by_key = {a.key: a for a in accounts}
    assert by_key["1000"].balance == D("15000.00")
    assert by_key["2000"].balance == D("-4000.00")
    assert by_key["1000"].parent == "Assets"
    assert by_key["1000"].name == "Cash"
    assert "Petty Cash" in by_key, "an unnumbered account keys off its full name"
    assert accounts.sum_of_balances == accounts.printed_total == D("15500.00")
    assert accounts.foots


def test_account_types_are_normalised_to_the_model_vocabulary(tmp):
    """QuickBooks writes types in the plural; model.py's vocabulary is singular.

    Left raw, Account.debit_normal is False for every expense and every asset
    in the company, which then decides the wrong way to sign a general ledger.
    """
    path = write_sheet(os.path.join(tmp, "accounts.xlsx"), ACCOUNT_LIST_ROWS)
    by_key = {a.key: a for a in load_account_list(path)}
    assert by_key["6000"].type == "expense"          # from "Expenses"
    assert by_key["2000"].type == "accounts payable"  # from "Accounts Payable (A/P)"
    assert by_key["3300"].type == "other current asset"
    assert by_key["6000"].debit_normal is True
    assert by_key["1000"].debit_normal is True
    assert by_key["2000"].debit_normal is False
    assert by_key["3300"].role == "prepaid"
    assert by_key["1000"].role == "bank"


def test_profit_and_loss_by_month_resolves_its_total_column(tmp):
    path = write_sheet(os.path.join(tmp, "plm.xlsx"), pl_by_month_rows())
    pl = load_profit_and_loss_by_month(path)
    assert pl.months == ["2025-01", "2025-02", "2025-03"]
    assert sorted(pl.rows) == ["4000 Product Revenue", "6000 Contractor Expense"]
    assert pl.rows["4000 Product Revenue"]["2025-02"] == D("2000.00")
    # the annual Total column exists only once "=B7+C7+D7" is resolved
    assert pl.row_totals["4000 Product Revenue"] == D("6000.00")
    assert pl.totals["2025-03"] == D("2700.00")      # the report's own Net Income
    assert "Total for Income" in pl.subtotals
    assert pl.foots and not pl.discrepancies


def test_profit_and_loss_by_month_records_a_row_that_does_not_foot(tmp):
    path = write_sheet(os.path.join(tmp, "plm_bad.xlsx"), pl_by_month_rows(revenue_total="=B7+C7"))
    pl = load_profit_and_loss_by_month(path)
    assert not pl.foots
    assert pl.discrepancies and "3000.00" in pl.discrepancies[0], pl.discrepancies
    assert pl.rows["4000 Product Revenue"]["2025-03"] == D("3000.00"), "the parse is unaffected"


def test_discovery_tolerates_plus_signs_spaces_and_duplicate_suffixes(tmp):
    """TRAP 9. Filenames arrive as the browser saved them."""
    folder = os.path.join(tmp, "exports")
    write_sheet(os.path.join(folder, "Acme+Robotics+Inc_Trial+Balance.xlsx"),
                trial_balance_rows(), "Trial Balance")
    write_sheet(os.path.join(folder, "Acme Robotics Inc_General Ledger (1).xlsx"), GL_AMOUNT_ROWS)
    write_sheet(os.path.join(folder, "Acme Robotics Inc_Profit and Loss by Month.xlsx"),
                pl_by_month_rows())
    write_sheet(os.path.join(folder, "Acme Robotics Inc_Profit and Loss Detail.xlsx"),
                pl_by_month_rows())
    found = discover(folder)
    assert list(found["trial_balance"])[0].endswith("Acme+Robotics+Inc_Trial+Balance.xlsx")
    assert len(found["general_ledger"]) == 1
    # "profit and loss" is a prefix of both, so the specific token must win
    assert len(found.get("profit_and_loss_by_month", [])) == 1
    assert len(found.get("profit_and_loss_detail", [])) == 1
    assert "profit_and_loss" not in found


def test_load_all_populates_a_balanced_ledger_without_double_counting(tmp):
    """The Journal and the General Ledger cover the same period. Loading both
    would post every transaction twice; the ledger would still balance."""
    folder = os.path.join(tmp, "all")
    write_sheet(os.path.join(folder, "Acme Robotics Inc_Account List.xlsx"), ACCOUNT_LIST_ROWS)
    write_sheet(os.path.join(folder, "Acme Robotics Inc_General Ledger.xlsx"), GL_AMOUNT_ROWS)
    write_sheet(os.path.join(folder, "Acme Robotics Inc_Journal.xlsx"), JOURNAL_ROWS)
    write_sheet(os.path.join(folder, "Acme+Robotics+Inc_Trial+Balance.xlsx"),
                trial_balance_rows(), "Trial Balance")

    ledger = load_all(folder)
    assert ledger.company == "Acme Robotics Inc."
    assert ledger.basis == "accrual"
    assert len(ledger.accounts) == 6
    assert len(ledger.lines) == 4, "the Journal must not be counted on top of the ledger"
    assert ledger.total_debits() == D("6200.00") == ledger.total_credits()
    assert ledger.foots
    assert ledger.period_start == datetime.date(2025, 1, 1)
    assert ledger.period_end == datetime.date(2025, 12, 31)

    used = {s["file"]: s for s in ledger.sources if s.get("used")}
    assert "Acme Robotics Inc_General Ledger.xlsx" in used
    skipped = [s for s in ledger.sources if s["file"].endswith("Journal.xlsx")]
    assert skipped and "overlaps" in skipped[0]["note"]
    # A BALANCE IS THE OPENING POSITION PLUS THE MOVEMENT, and the two are
    # separately readable. The lines move 1000 by 3,800.00 (5,000 in, 1,200
    # out); the ledger opened at 1,000.00, which the General Ledger states on
    # its own Beginning Balance row and which `load_all` now carries through.
    assert ledger.movement_of("1000") == D("3800.00")
    assert ledger.opening_of("1000") == D("1000.00")
    assert ledger.balance_of("1000") == D("4800.00")
    assert ledger.opening_basis, "the opening basis has to be recorded, not assumed"
    assert "General Ledger" in ledger.opening_basis

    # The trial balance is opened, not glanced at.
    assert len(ledger.trial_balances) == 1
    tb = ledger.trial_balance_as_of(datetime.date(2025, 12, 31))
    assert tb is not None and tb.foots
    tb_source = [s for s in ledger.sources if s["kind"] == "trial_balance"]
    assert tb_source and tb_source[0]["used"] is True, (
        "a trial balance recorded as 'read for its face only' is a trial balance "
        "no test can fail against"
    )


def test_a_missing_or_unreadable_file_raises_export_error(tmp):
    try:
        load_trial_balance(os.path.join(tmp, "nope.xlsx"))
    except ExportError as exc:
        assert "no such file" in str(exc)
    else:
        raise AssertionError("a missing file must raise ExportError")

    junk = os.path.join(tmp, "junk.xlsx")
    with open(junk, "wb") as fh:
        fh.write(b"this is not a workbook")
    try:
        load_trial_balance(junk)
    except ExportError as exc:
        assert "cannot open" in str(exc)
    else:
        raise AssertionError("a non-workbook must raise ExportError")


def test_a_report_with_no_debit_credit_header_raises(tmp):
    rows = [["Acme Robotics Inc.", ""], ["Trial Balance", ""], ["As of December 31, 2025", ""],
            ["", ""], ["Account", "Something else"], ["1000 Cash", "=1.00"]]
    path = write_sheet(os.path.join(tmp, "headerless.xlsx"), rows, "Trial Balance")
    try:
        load_trial_balance(path)
    except ExportError as exc:
        assert "Debit/Credit" in str(exc)
    else:
        raise AssertionError("a trial balance with no Debit/Credit header must raise")


# ------------------------------------------------------------ plain runner

def _tests():
    g = globals()
    return [(n, g[n]) for n in sorted(g) if n.startswith("test_") and callable(g[n])]


def main():
    failures = []
    for name, fn in _tests():
        with tempfile.TemporaryDirectory() as tmp:
            try:
                fn(tmp)
                print(f"ok   {name}")
            except Exception as exc:  # noqa: BLE001 - a test runner reports everything
                failures.append((name, exc))
                print(f"FAIL {name}: {type(exc).__name__}: {exc}")
    total = len(_tests())
    print(f"\n{total - len(failures)} passed, {len(failures)} failed")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
