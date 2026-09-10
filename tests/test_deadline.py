"""The deadline gates: filing-year scope, guided intake, and the three refusals.

Run it with

    python3 tests/test_deadline.py

The fixture reproduces the shape a wind-down file takes when it is scoped by
hand: 863 queue rows across five declared accounts, 509 of them dated in the
year being filed, 197 of those sitting against a period that already reconciled
clean, and 241 real card items on an account that exists twice because the bank
was re-linked. The tool has to reach 312 rather than 863.

`test_a_disconnect_is_never_permitted_with_a_loaded_queue` is the one to treat
as a stop. Disconnecting a feed deletes every item in that account's Pending and
For Review tabs, and for an account whose activity was never booked there is no
second copy anywhere. If that test ever passes for the wrong reason, the tool
will help somebody destroy a year of evidence in one action.
"""

import csv
import datetime as dt
import json
import os
import pathlib
import shutil
import subprocess
import sys
import tempfile

REPO = pathlib.Path(__file__).resolve().parent.parent
BOOKS = REPO / "bin" / "books.py"
EXAMPLE = REPO / "examples" / "acme-robotics"
sys.path.insert(0, str(REPO / "lib"))

OK, FAILURE, REFUSAL = 0, 1, 2

# The counts the fixture exists to reproduce.
TOTAL = 863
IN_FILING_YEAR = 509
AGAINST_RECONCILED = 197
IN_SCOPE = 312
REAL_CARD_ITEMS = 241
FILING_YEAR = 2025

# Written as escapes so that this file, which forbids them, does not contain one.
EM_DASH = "\u2014"
EN_DASH = "\u2013"

_WORK = None


# ------------------------------------------------------------------ helpers

def run(*args, workdir=None, expect=None):
    cmd = [sys.executable, str(BOOKS)]
    if workdir is not None:
        cmd += ["--workdir", str(workdir)]
    cmd += [str(a) for a in args]
    p = subprocess.run(cmd, capture_output=True, text=True, cwd=str(REPO))
    if expect is not None and p.returncode != expect:
        raise AssertionError(
            f"`books.py {' '.join(str(a) for a in args)}` exited {p.returncode}, "
            f"expected {expect}\n--- stdout ---\n{p.stdout}\n--- stderr ---\n{p.stderr}"
        )
    return p.returncode, p.stdout, p.stderr


def ensure_example():
    exports = EXAMPLE / "exports"
    if exports.is_dir() and any(exports.glob("*.xlsx")):
        return
    p = subprocess.run([sys.executable, str(EXAMPLE / "build.py")],
                       capture_output=True, text=True, cwd=str(REPO))
    if p.returncode != 0:
        raise AssertionError(f"could not build the example company:\n{p.stderr}")


# The five declared accounts. The unnumbered card is first on purpose: it is the
# one holding the live feed, so a queue export carrying the mask belongs to it,
# which is exactly the trap a bank re-link sets.
ACCOUNTS = [
    {"book": "Vantage Credit Card 3391", "label": "Vantage Credit Card 3391",
     "mask": "3391", "kind": "card", "feed": "live", "feed_last": "2026-09",
     "parser": "generic_csv", "institution": "", "note": "created by a re-link"},
    {"book": "2100", "label": "2100 Vantage Card 3391", "mask": "3391",
     "kind": "card", "feed": "none", "feed_last": "", "parser": "generic_csv",
     "institution": "", "note": "the numbered original, holds the reconciled history"},
    {"book": "1010", "label": "1010 Northgate Checking 7742", "mask": "7742",
     "kind": "bank", "feed": "live", "feed_last": "2026-09",
     "parser": "generic_csv", "institution": "", "note": ""},
    {"book": "1020", "label": "1020 Northgate Savings 7809", "mask": "7809",
     "kind": "bank", "feed": "dead", "feed_last": "2025-04",
     "parser": "generic_csv", "institution": "", "note": "eight months missing"},
    {"book": "1200", "label": "1200 Paylane Clearing", "mask": "9931",
     "kind": "bank", "feed": "live", "feed_last": "2026-09",
     "parser": "generic_csv", "institution": "", "note": ""},
]

VENDORS = ["NORTHGATE ACH CREDIT", "CLOUDWORKS INVOICE", "OFFICE DEPOT 4412",
           "PAYLANE PAYOUT", "STATE FILING FEE", "FREIGHTLINE LTD",
           "MERIDIAN INSURANCE", "TRAVEL BOOKING 88", "PARTS SUPPLY CO",
           "COFFEE AND CO"]


def _rows(start, end, count, first_amount, total=None):
    """`count` rows spread across a date range, with plain descriptors."""
    span = max((end - start).days, 1)
    out = []
    for i in range(count):
        d = start + dt.timedelta(days=(i * span) // max(count - 1, 1))
        amount = first_amount if i == 0 else round(first_amount + (i % 7) * 3.11, 2)
        out.append([d.isoformat(), f"{VENDORS[i % len(VENDORS)]} {i + 1:04d}",
                    f"{amount:.2f}"])
    if total is not None and out:
        got = sum(float(r[2]) for r in out)
        out[-1][2] = f"{float(out[-1][2]) + (total - got):.2f}"
    return out


def write_queue(path, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["Date", "Description", "Amount"])
        w.writerows(rows)


def build_queue(folder):
    """863 rows: 509 in the filing year, 197 of those against a reconciled month."""
    y, nxt = FILING_YEAR, FILING_YEAR + 1
    card = (
        # 241 real card items, March to December of the year being filed. This
        # account holds the live feed and nothing else in the file holds them.
        _rows(dt.date(y, 3, 1), dt.date(y, 12, 28), REAL_CARD_ITEMS, -41.20)
        # 300 dated next year, which belong to next year's return.
        + _rows(dt.date(nxt, 1, 2), dt.date(nxt, 9, 1), 300, -63.75)
    )
    checking = (
        # 197 re-downloads sitting against a period reconciled through June.
        _rows(dt.date(y, 1, 2), dt.date(y, 6, 29), AGAINST_RECONCILED, 812.44)
        + _rows(dt.date(y, 7, 1), dt.date(y, 12, 20), 42, -1204.05)
        + _rows(dt.date(nxt, 1, 5), dt.date(nxt, 8, 30), 30, -318.90)
    )
    savings = (
        _rows(dt.date(y, 5, 4), dt.date(y, 11, 30), 9, -9000.00, total=-120000.00)
        + _rows(dt.date(nxt, 2, 3), dt.date(nxt, 7, 4), 12, 55.10)
    )
    clearing = (
        _rows(dt.date(y, 2, 6), dt.date(y, 12, 4), 20, 2240.19)
        + _rows(dt.date(nxt, 1, 9), dt.date(nxt, 6, 6), 12, 91.44)
    )
    write_queue(folder / "for-review-3391.csv", card)
    write_queue(folder / "for-review-7742.csv", checking)
    write_queue(folder / "for-review-7809.csv", savings)
    write_queue(folder / "for-review-9931.csv", clearing)
    return len(card) + len(checking) + len(savings) + len(clearing)


def build_workdir(prefix):
    ensure_example()
    wd = pathlib.Path(tempfile.mkdtemp(prefix=prefix))
    run("init", workdir=wd, expect=OK)
    for src in sorted((EXAMPLE / "exports").glob("*.xlsx")):
        shutil.copy2(src, wd / "exports" / src.name)
    run("learn", workdir=wd, expect=OK)
    prof_path = wd / "profiles" / "mine.local.json"
    prof = json.loads(prof_path.read_text(encoding="utf-8"))
    prof["accounts"] = [dict(a) for a in ACCOUNTS]
    prof_path.write_text(json.dumps(prof, indent=2) + "\n", encoding="utf-8")
    return wd


def work():
    """A working directory shaped like the file this was built for, built once."""
    global _WORK
    if _WORK is not None:
        return _WORK
    wd = build_workdir("close-the-books-deadline-")
    written = build_queue(wd / "for-review")
    assert written == TOTAL, f"the fixture wrote {written} rows, not {TOTAL}"
    for field, value in (("basis", "accrual"),
                         ("end_use", f"the {FILING_YEAR} federal corporate return"),
                         ("materiality", "500.00")):
        run("answer", f"entity.{field}", value, workdir=wd, expect=OK)
    _WORK = wd
    return wd


def set_filing_year(wd):
    run("filing-year", str(FILING_YEAR), "--due", f"{FILING_YEAR + 1}-10-15",
        workdir=wd, expect=OK)


def record_reconciled(wd):
    run("intake", "--account", "1010", "--reconciled-through",
        f"{FILING_YEAR}-06-30", workdir=wd, expect=OK)


def scope_json(wd):
    return json.loads((wd / "reports" / "scope.json").read_text(encoding="utf-8"))


def load_cli():
    import importlib.util
    spec = importlib.util.spec_from_file_location("books_cli", BOOKS)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


# ----------------------------------------------------------- the filing year

def test_scope_refuses_until_somebody_says_which_year_is_being_filed():
    wd = work()
    code, _, err = run("scope", workdir=wd, expect=REFUSAL)
    assert "no filing year has been set" in err
    assert "books.py filing-year" in err
    assert "Traceback" not in err


def test_the_filing_year_is_never_inferred_from_an_export():
    """`learn` fills entity.fiscal_year from the period the exports cover.

    That is a fact about the download. Reading it as the year being filed would
    answer the question without anyone being asked it.
    """
    wd = work()
    prof = json.loads((wd / "profiles" / "mine.local.json").read_text(encoding="utf-8"))
    assert prof["entity"]["fiscal_year"], "learn does fill fiscal_year"
    live_path = wd / "answers" / "live-screen.json"
    if live_path.exists():
        data = json.loads(live_path.read_text(encoding="utf-8"))
        assert not (data.get("filing") or {}).get("year"), "nobody has said yet"


def test_filing_year_records_the_year_and_the_due_date():
    wd = work()
    set_filing_year(wd)
    code, out, _ = run("filing-year", workdir=wd, expect=OK)
    assert f"Filing {FILING_YEAR}" in out
    assert f"{FILING_YEAR + 1}-10-15" in out
    data = json.loads((wd / "answers" / "live-screen.json").read_text(encoding="utf-8"))
    assert data["filing"]["year"] == str(FILING_YEAR)
    assert data["filing"]["due"] == f"{FILING_YEAR + 1}-10-15"
    assert data["filing"]["year_on"], "the date it was given is stored with it"


# ------------------------------------------------------------------- scope

def test_the_queue_is_863_and_the_deadline_needs_312():
    wd = work()
    set_filing_year(wd)
    record_reconciled(wd)
    code, out, _ = run("scope", workdir=wd, expect=OK)
    data = scope_json(wd)
    assert data["total"] == TOTAL, data
    assert data["in_scope"] == IN_SCOPE, data
    assert data["deferred"] == TOTAL - IN_SCOPE, data
    assert data["counts"]["next_year"] == TOTAL - IN_FILING_YEAR, data
    assert data["counts"]["reconciled_period"] == AGAINST_RECONCILED, data
    assert f"{IN_SCOPE} of {TOTAL}" in out, out


def test_the_card_items_the_only_record_holds_are_all_in_scope():
    wd = work()
    set_filing_year(wd)
    record_reconciled(wd)
    run("scope", workdir=wd, expect=OK)
    by = scope_json(wd)["by_account"]
    assert by["Vantage Credit Card 3391"]["in_scope"] == REAL_CARD_ITEMS, by


def test_a_re_download_against_a_reconciled_month_is_set_aside():
    wd = work()
    set_filing_year(wd)
    record_reconciled(wd)
    code, out, _ = run("scope", workdir=wd, expect=OK)
    assert "already reconciled" in out
    assert "counts the same transaction twice" in out


def test_every_row_set_aside_is_written_out_rather_than_hidden():
    wd = work()
    set_filing_year(wd)
    record_reconciled(wd)
    run("scope", workdir=wd, expect=OK)
    md = (wd / "reports" / "deferred.md").read_text(encoding="utf-8")
    rows = [l for l in md.splitlines() if l.startswith("| 20")]
    assert len(rows) == TOTAL - IN_SCOPE, f"{len(rows)} rows written out"
    assert "Set aside, not dropped" in md


def test_catchup_batches_only_the_rows_this_deadline_needs():
    wd = work()
    set_filing_year(wd)
    record_reconciled(wd)
    code, out, _ = run("catchup", "--batch-size", "150", workdir=wd, expect=OK)
    assert f"{IN_SCOPE} of {TOTAL}" in out, out
    assert "set aside for now" in out
    from openpyxl import load_workbook
    total = 0
    for path in sorted((wd / "review").glob("batch-*.xlsx")):
        wb = load_workbook(path)
        ws = wb["Review"]
        header = next(r for r in range(1, 30)
                      if str(ws.cell(row=r, column=1).value or "").strip() == "Date")
        total += sum(1 for r in range(header + 1, ws.max_row + 1)
                     if ws.cell(row=r, column=1).value is not None)
        wb.close()
    assert total == IN_SCOPE, f"{total} rows reached the workbooks, expected {IN_SCOPE}"


# ------------------------------------------------------------------ intake

def test_intake_names_the_screen_and_says_what_it_still_lacks():
    wd = work()
    code, out, _ = run("intake", workdir=wd, expect=OK)
    assert "Transactions, Bank transactions" in out
    assert "Audit log" in out
    assert "Still missing" in out
    assert "blocks every command that produces figures for the return" in out


def test_intake_keeps_what_it_was_already_given():
    wd = work()
    run("intake", "--account", "1020", "--bank-balance", "37.14",
        "--as-of", f"{FILING_YEAR}-12-31", "--feed", "stopped",
        "--feed-last", f"{FILING_YEAR}-04-30", workdir=wd, expect=OK)
    run("intake", "--account", "1020", "--for-review", "21", workdir=wd, expect=OK)
    data = json.loads((wd / "answers" / "live-screen.json").read_text(encoding="utf-8"))
    node = data["accounts"]["1020"]
    assert node["bank_balance"] == "37.14"
    assert node["feed_state"] == "stopped"
    assert node["for_review_count"] == 21
    assert node["bank_balance_on"], "the date it was given is stored with it"
    code, out, _ = run("intake", workdir=wd, expect=OK)
    still = out.split("Still missing")[1]
    assert "1020 Northgate Savings 7809: the bank's own balance" not in still


def test_intake_refuses_an_account_the_profile_does_not_declare():
    wd = work()
    code, _, err = run("intake", "--account", "999999", "--for-review", "1",
                       workdir=wd, expect=REFUSAL)
    assert "there is no account" in err
    assert "Traceback" not in err


# ------------------------------------ gate one: never disconnect a loaded queue

def test_a_disconnect_is_never_permitted_with_a_loaded_queue():
    """The stop. Disconnecting deletes the only record those rows have.

    Do not make this pass by loosening the gate. The whole point is that the
    action is unavailable while anything is in the queue, whichever way it is
    asked for.
    """
    wd = work()
    run("intake", "--account", "Vantage Credit Card 3391", "--for-review", "541",
        workdir=wd, expect=OK)
    for action in ("disconnect", "merge"):
        code, out, err = run("merge-plan", "--duplicate", "Vantage Credit Card 3391",
                             "--action", action, workdir=wd, expect=REFUSAL)
        assert "goes ahead" in err, err
        assert "DELETES" in err, err
        assert "Traceback" not in err
        assert "The order, and it is the whole point" not in out, \
            "a plan was printed alongside the refusal"


def test_the_refusal_says_the_queue_is_the_only_record_of_those_transactions():
    wd = work()
    run("intake", "--account", "Vantage Credit Card 3391", "--for-review", "541",
        workdir=wd, expect=OK)
    code, _, err = run("merge-plan", "--duplicate", "Vantage Credit Card 3391",
                       workdir=wd, expect=REFUSAL)
    assert "only" in err and "record of it anywhere in the file" in err
    assert "Book them, then dispose of what is left, then disconnect, then merge" in err
    assert "books.py catchup" in err


def test_a_merge_is_refused_while_nobody_has_read_the_count_off_the_screen():
    wd = work()
    code, _, err = run("merge-plan", "--duplicate", "2100", workdir=wd, expect=REFUSAL)
    assert "nobody has said what is sitting in" in err
    assert "books.py intake --account 2100 --for-review N" in err


def test_a_stale_export_alone_still_stops_the_merge():
    """The screen says zero, the folder still holds the rows. One is out of date."""
    wd = work()
    run("intake", "--account", "Vantage Credit Card 3391", "--for-review", "0",
        workdir=wd, expect=OK)
    code, _, err = run("merge-plan", "--duplicate", "Vantage Credit Card 3391",
                       workdir=wd, expect=REFUSAL)
    assert "still holds" in err
    assert "out of date" in err.replace("\n  ", " ")


def test_merge_plan_gives_the_order_once_both_sources_read_zero():
    wd = build_workdir("close-the-books-merge-")
    for key in ("Vantage Credit Card 3391", "2100"):
        run("intake", "--account", key, "--for-review", "0", workdir=wd, expect=OK)
    code, out, _ = run("merge-plan", "--duplicate", "Vantage Credit Card 3391",
                       workdir=wd, expect=OK)
    assert "The order, and it is the whole point" in out
    assert out.index("Book everything in the queue") < out.index("disconnect the feed")
    assert "Doing 3 before 1 discards the only record" in out
    shutil.rmtree(wd, ignore_errors=True)


def test_merge_plan_finds_the_account_that_exists_twice():
    wd = work()
    code, out, _ = run("merge-plan", workdir=wd, expect=OK)
    assert "Vantage Credit Card 3391" in out
    assert "2100" in out
    assert "both answer to 3391" in out


# ------------------------ gate two: the queue is not the same as the period

def test_completeness_measures_the_queue_against_what_the_bank_says():
    wd = work()
    set_filing_year(wd)
    run("intake", "--account", "1020", "--bank-balance", "37.14",
        "--as-of", f"{FILING_YEAR}-12-31", workdir=wd, expect=OK)
    code, out, _ = run("completeness", workdir=wd, expect=FAILURE)
    assert "the bank says" in out
    assert "books plus the queue" in out
    assert "further from the bank" in out, out
    data = json.loads((wd / "reports" / "completeness.json").read_text(encoding="utf-8"))
    hit = next(f for f in data["findings"] if f["account"] == "1020")
    assert hit["moves_away"] is True, hit
    assert hit["open"] is True, hit


def test_the_queue_cannot_be_called_finished_while_a_feed_hole_stands():
    wd = work()
    set_filing_year(wd)
    run("intake", "--account", "1020", "--bank-balance", "37.14",
        "--as-of", f"{FILING_YEAR}-12-31", workdir=wd, expect=OK)
    code, _, err = run("attest", "--unbooked", "0", "--note", "queue is empty",
                       workdir=wd, expect=REFUSAL)
    flat = " ".join(err.split())
    assert "the queue cannot be called finished" in flat, err
    assert ("Finishing the queue and the period being complete are different things"
            in flat)
    assert "books.py completeness" in err
    assert "Traceback" not in err
    assert not (wd / "answers" / "attestation.json").exists()


def test_settling_a_finding_needs_the_document_that_explains_it():
    wd = work()
    set_filing_year(wd)
    code, _, err = run("completeness", "--settle", "1020", workdir=wd, expect=REFUSAL)
    assert "--evidence" in err
    flat = " ".join(err.split())
    assert "a plug destroys the only signal there was" in flat, err


def test_a_settled_finding_carries_its_document_a_date_and_a_person():
    wd = build_workdir("close-the-books-settle-")
    run("filing-year", str(FILING_YEAR), "--due", f"{FILING_YEAR + 1}-10-15",
        workdir=wd, expect=OK)
    code, out, _ = run("completeness", "--settle", "1010",
                       "--evidence", "Northgate statement 2025-12, page 2, wire in",
                       "--by", "the owner", workdir=wd)
    assert "Northgate statement 2025-12" in out
    data = json.loads((wd / "answers" / "live-screen.json").read_text(encoding="utf-8"))
    node = data["completeness_settled"]["1010"]
    assert node["evidence"].startswith("Northgate statement")
    assert node["on"] and node["by"] == "the owner"
    shutil.rmtree(wd, ignore_errors=True)


# --------------------- gate three: nothing while automation is still posting

def test_handoff_refuses_while_rules_are_still_posting_into_the_year():
    wd = work()
    set_filing_year(wd)
    code, _, err = run("handoff", workdir=wd, expect=REFUSAL)
    flat = " ".join(err.split())
    assert f"produces figures for the {FILING_YEAR} return" in flat
    assert "post without being seen" in flat
    assert "books.py intake --rules" in err
    assert "Traceback" not in err


def test_entries_refuses_while_rules_are_still_posting_into_the_year():
    wd = work()
    set_filing_year(wd)
    code, _, err = run("entries", "--through", f"{FILING_YEAR}-12-31",
                       workdir=wd, expect=REFUSAL)
    flat = " ".join(err.split())
    assert "An entry worksheet produces figures" in flat
    assert "still changing under it" in flat


def test_a_rule_that_only_categorizes_is_not_the_same_as_one_that_posts():
    wd = work()
    set_filing_year(wd)
    run("intake", "--rules", "15", "--auto-add", "15", workdir=wd, expect=OK)
    code, _, err = run("entries", "--through", f"{FILING_YEAR}-12-31",
                       workdir=wd, expect=REFUSAL)
    assert "15 of 15 bank rules post without anyone seeing them" in " ".join(err.split())
    run("intake", "--rules", "15", "--auto-add", "0", workdir=wd, expect=OK)
    code, out, err = run("entries", "--through", f"{FILING_YEAR}-12-31", workdir=wd)
    assert "post without anyone seeing" not in out + err


def test_the_automation_report_names_what_posted_after_the_last_person():
    wd = work()
    set_filing_year(wd)
    run("intake", "--last-human", f"{FILING_YEAR}-09-01", workdir=wd, expect=OK)
    run("intake", "--rules", "15", "--auto-add", "15", workdir=wd, expect=OK)
    code, _, err = run("handoff", workdir=wd, expect=REFUSAL)
    flat = " ".join(err.split())
    assert f"{FILING_YEAR}-09-01" in flat
    assert "the last day a person worked in this file" in flat
    assert "An export cannot say who posted a line. The audit log can." in flat


# ------------------------------------------------ the analyses, on their own

class _StubLedger:
    """Enough Ledger for the completeness arithmetic, and nothing more."""

    accounts = {}

    def __init__(self, balances):
        self._balances = balances

    def balance_as_of(self, key, cutoff):
        return self._balances.get(key)

    def balance_of(self, key):
        return self._balances.get(key)

    def opening_of(self, key):
        return None


def test_the_completeness_arithmetic_on_the_numbers_it_was_built_for():
    """Books (41,876.20), bank 37.14, unbooked (52,309.55).

    The queue moves the account away from the bank rather than toward it, which
    is the signal that transactions exist in neither place.
    """
    books = load_cli()
    from decimal import Decimal
    from closethebooks.model import BankLine
    from closethebooks.profile import AccountSpec, Entity, Profile

    prof = Profile(entity=Entity(name="Fixture", materiality="500.00"))
    prof.accounts = [AccountSpec(book="1020", label="Savings", mask="7809",
                                 kind="bank", feed="dead")]
    ledger = _StubLedger({"1020": Decimal("-41876.20")})
    queue = [BankLine(date=dt.date(FILING_YEAR, 6, 1), descriptor="x",
                      amount=Decimal("-52309.55"), account_key="1020")]
    live = {"accounts": {"1020": {"bank_balance": "37.14",
                                  "bank_balance_as_of": f"{FILING_YEAR}-12-31"}},
            "rules": {}, "people": {}}

    findings = books.analyze_completeness(prof, ledger, queue, [], live,
                                          filing_year=FILING_YEAR,
                                          materiality=Decimal("500.00"))
    f = findings[0]
    assert f.after == Decimal("-94185.75"), f.after
    assert f.difference == Decimal("94222.89"), f.difference
    assert f.moves_away is True
    assert f.open is True
    assert "further from the bank" in f.why


def test_scope_puts_a_row_in_exactly_one_group():
    """The already-booked test runs first, and on purpose.

    It is the only group whose members do damage when they are worked, and that
    is true whether or not they also fall outside the filing year. So the row
    dated before the year being filed goes on an account with no reconciliation
    date, which is the only way it can land in its own group.
    """
    books = load_cli()
    from decimal import Decimal
    from closethebooks.model import BankLine

    lines = [
        BankLine(date=dt.date(2024, 5, 1), descriptor="older", amount=Decimal("1"),
                 account_key="1200"),
        BankLine(date=dt.date(FILING_YEAR, 3, 1), descriptor="reconciled",
                 amount=Decimal("1"), account_key="1010"),
        BankLine(date=dt.date(FILING_YEAR, 9, 1), descriptor="in scope",
                 amount=Decimal("1"), account_key="1010"),
        BankLine(date=dt.date(FILING_YEAR + 1, 2, 1), descriptor="next",
                 amount=Decimal("1"), account_key="1010"),
    ]
    report = books.analyze_scope(
        lines, filing_year=FILING_YEAR,
        reconciled_through={"1010": dt.date(FILING_YEAR, 6, 30)})
    assert len(report.in_scope) == 1
    assert report.counts == {"earlier_year": 1, "reconciled_period": 1,
                             "next_year": 1}, report.counts
    assert len(report.in_scope) + len(report.deferred) == len(lines)


# ------------------------------------------------------------------- shape

def test_every_new_refusal_exits_two_and_reads_as_a_sentence():
    wd = work()
    for args in (("merge-plan", "--duplicate", "2100"),
                 ("completeness", "--settle", "1020")):
        code, out, err = run(*args, workdir=wd)
        assert code == REFUSAL, f"{args} exited {code}"
        assert "Traceback" not in err
        assert err.strip(), f"{args} refused with nothing to read"


def test_no_new_command_prints_a_balance_as_a_negative_number():
    import re as _re
    wd = work()
    set_filing_year(wd)
    run("intake", "--account", "1020", "--bank-balance", "37.14",
        "--as-of", f"{FILING_YEAR}-12-31", workdir=wd)
    for command in (("scope",), ("completeness",), ("filing-year",), ("intake",),
                    ("merge-plan",)):
        _, out, _ = run(*command, workdir=wd)
        for line in out.splitlines():
            low = line.lower()
            if "difference" in low or "netting" in low or "out by" in low:
                continue
            for token in _re.findall(r"[-\u2212]\s?\d[\d,]*\.\d{2}", line):
                raise AssertionError(
                    f"`books.py {command[0]}` printed a balance as a negative "
                    f"number: {token!r} in {line!r}")


def test_no_long_dashes_in_anything_the_new_commands_print():
    wd = work()
    set_filing_year(wd)
    for command in (("scope",), ("intake",), ("filing-year",), ("merge-plan",),
                    ("completeness",)):
        _, out, err = run(*command, workdir=wd)
        for text, where in ((out, "stdout"), (err, "stderr")):
            assert EM_DASH not in text and EN_DASH not in text, \
                f"books.py {command[0]} printed a long dash on {where}"


# --------------------------------------------------------------------- main

def main():
    fns = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    order = [
        "test_scope_refuses_until_somebody_says_which_year_is_being_filed",
        "test_the_filing_year_is_never_inferred_from_an_export",
        "test_filing_year_records_the_year_and_the_due_date",
        "test_the_queue_is_863_and_the_deadline_needs_312",
        "test_the_card_items_the_only_record_holds_are_all_in_scope",
        "test_a_re_download_against_a_reconciled_month_is_set_aside",
        "test_every_row_set_aside_is_written_out_rather_than_hidden",
        "test_catchup_batches_only_the_rows_this_deadline_needs",
        "test_intake_names_the_screen_and_says_what_it_still_lacks",
        "test_intake_keeps_what_it_was_already_given",
        "test_intake_refuses_an_account_the_profile_does_not_declare",
        "test_a_merge_is_refused_while_nobody_has_read_the_count_off_the_screen",
        "test_a_disconnect_is_never_permitted_with_a_loaded_queue",
        "test_the_refusal_says_the_queue_is_the_only_record_of_those_transactions",
        "test_a_stale_export_alone_still_stops_the_merge",
        "test_merge_plan_gives_the_order_once_both_sources_read_zero",
        "test_merge_plan_finds_the_account_that_exists_twice",
        "test_completeness_measures_the_queue_against_what_the_bank_says",
        "test_the_queue_cannot_be_called_finished_while_a_feed_hole_stands",
        "test_settling_a_finding_needs_the_document_that_explains_it",
        "test_a_settled_finding_carries_its_document_a_date_and_a_person",
        "test_handoff_refuses_while_rules_are_still_posting_into_the_year",
        "test_entries_refuses_while_rules_are_still_posting_into_the_year",
        "test_a_rule_that_only_categorizes_is_not_the_same_as_one_that_posts",
        "test_the_automation_report_names_what_posted_after_the_last_person",
        "test_the_completeness_arithmetic_on_the_numbers_it_was_built_for",
        "test_scope_puts_a_row_in_exactly_one_group",
        "test_every_new_refusal_exits_two_and_reads_as_a_sentence",
        "test_no_new_command_prints_a_balance_as_a_negative_number",
        "test_no_long_dashes_in_anything_the_new_commands_print",
    ]
    by_name = {f.__name__: f for f in fns}
    missing = [n for n in by_name if n not in order]
    assert not missing, f"tests not in the run order: {missing}"

    failed = 0
    for name in order:
        try:
            by_name[name]()
            print(f"PASS {name}")
        except Exception as exc:  # noqa: BLE001
            failed += 1
            print(f"FAIL {name}: {type(exc).__name__}: {exc}")
    if _WORK is not None and not os.environ.get("BOOKS_KEEP_WORKDIR"):
        shutil.rmtree(_WORK, ignore_errors=True)
    elif _WORK is not None:
        print(f"\nworking directory kept at {_WORK}")
    print(f"\n{len(order) - failed}/{len(order)} passed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
