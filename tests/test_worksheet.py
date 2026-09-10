"""Tests for the two things that decide whether 40 hand-typed entries ever finish.

Every fixture here is invented. Acme Robotics Inc. is not a real company, no
number came from a real ledger, and no file here came from a real bank.

The measured claim these tests defend: about 2 hours 8 minutes to type 40
two-line entries into QuickBooks Online US by hand, and the honest expectation
that a third of them never get typed. Two things move that number. Materiality
makes the list shorter (`materiality.roll_up`). Ticks make it resumable
(`je_worksheet.read_ticks`). Both have to be safe: a shorter list that moved
income between years, or a resume that says "done" about something untyped, is
worse than the long list it replaced.

Run either way:
    pytest tests/test_worksheet.py
    python3 tests/test_worksheet.py
"""

from __future__ import annotations

import datetime as dt
import sys
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "lib"))

from closethebooks import je_worksheet, materiality  # noqa: E402
from closethebooks.materiality import MaterialityError, Percentage  # noqa: E402
from closethebooks.model import Account, JournalLine, Ledger, ProposedEntry  # noqa: E402
from closethebooks.util import money  # noqa: E402

D = dt.date

CHART = {
    "1010": "Current Assets:Northgate Checking 7742",
    "1350": "Current Assets:Prepaid Expenses",
    "6120": "Operating Expenses:Software Subscriptions",
    "6300": "Operating Expenses:Contract Labor",
    "6410": "Operating Expenses:Bank Charges",
    "2100": "Liabilities:Accrued Payroll",
}


def resolve(key):
    return CHART.get(str(key))


# --------------------------------------------------------------- fixtures

def entry(number, date, amount, debit="6410", credit="1010", kind="bank_fees",
          basis="invented test fixture"):
    return ProposedEntry(
        date=date, number=number,
        lines=[(debit, amount, "0.00", f"{number} memo"),
               (credit, "0.00", amount, f"{number} memo")],
        kind=kind, basis=basis, batch_tag="01",
    )


def amortization(month, amount="208.33", year=2025):
    """One month of a twelve month prepaid, the recurring-template case."""
    return ProposedEntry(
        date=D(year, month, 28), number=f"JE-2{month:03d}",
        lines=[("6120", amount, "0.00", f"Amortize prepaid, month {month}"),
               ("1350", "0.00", amount, f"Amortize prepaid, month {month}")],
        kind="prepaid", basis="12-month prepaid schedule, signed order form",
        batch_tag="01",
    )


def amortization_family():
    """2,500.00 over twelve months: 11 x 208.33 and a 208.37 stub in month 12.

    This is the shape that used to split into a family of eleven and a
    stranger, which lost the recurring template for the whole family over four
    cents.
    """
    return [amortization(m) for m in range(1, 12)] + [amortization(12, "208.37")]


def small_ledger(revenue="400000.00"):
    """Just enough ledger for a percentage threshold to resolve against."""
    accounts = {
        "4000": Account(name="Product Revenue", number="4000", type="Income", role="revenue"),
        "1010": Account(name="Northgate Checking 7742", number="1010", type="Bank", role="bank"),
    }
    lines = [
        JournalLine(date=D(2025, 6, 30), account="1010", debit=revenue, credit="0.00"),
        JournalLine(date=D(2025, 6, 30), account="4000", debit="0.00", credit=revenue),
    ]
    return Ledger(accounts=accounts, lines=lines, company="Acme Robotics Inc.")


class FakeProfile:
    """A profile is a dataclass with an `entity`; only these two fields matter."""

    class Entity:
        def __init__(self, materiality="", fiscal_year=""):
            self.materiality = materiality
            self.fiscal_year = fiscal_year

    def __init__(self, materiality="", fiscal_year="2025"):
        self.entity = self.Entity(materiality, fiscal_year)


def worksheet_rows(path):
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb[je_worksheet.SHEET]
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    wb.close()
    return rows


def tick(path, indexes):
    """Mark the given 1-based entries done, the way a person would."""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb[je_worksheet.SHEET]
    seen = 0
    for row in ws.iter_rows():
        text = str(row[1].value or "")
        if je_worksheet.ENTRY_HEAD_RE.match(text.lstrip("'")):
            seen += 1
            if seen in indexes:
                row[0].value = "[x]"
    wb.save(path)
    wb.close()


# --------------------------------------------------- 1. reading a threshold

def test_threshold_is_parsed_in_every_form_a_human_writes_it():
    assert materiality.parse_threshold("1000") == money("1000.00")
    assert materiality.parse_threshold("1,000.00") == money("1000.00")
    assert materiality.parse_threshold("$1,000") == money("1000.00")
    assert materiality.parse_threshold(Decimal("500")) == money("500.00")
    assert materiality.parse_threshold(500) == money("500.00")

    pct = materiality.parse_threshold("1%")
    assert isinstance(pct, Percentage) and pct.percent and pct == Decimal("1")
    assert pct.base == "", "a bare percentage names no base, and must not pretend to"
    assert materiality.parse_threshold("0.5% of revenue").base == "revenue"
    assert materiality.parse_threshold("2% of total assets").base == "total_assets"

    # Blank in every disguise means "not set", never zero. Zero would be a
    # threshold; not set is a missing engagement fact.
    for blank in (None, "", "   ", "none", "N/A", "TBD"):
        assert materiality.parse_threshold(blank) is None

    for bad in ("about five hundred", "-100", "12%%"):
        try:
            materiality.parse_threshold(bad)
        except MaterialityError:
            pass
        else:
            raise AssertionError(f"{bad!r} was accepted as a threshold")


def test_a_percentage_threshold_resolves_against_a_named_base_or_says_it_cannot():
    ledger = small_ledger(revenue="400000.00")

    got = materiality.threshold_for(FakeProfile("1% of revenue"), ledger)
    assert got == money("4000.00")

    # Named base, but no ledger to compute it from: absent, and said out loud.
    detail = materiality.threshold_detail(FakeProfile("1% of revenue"), None)
    assert detail.threshold is None and not detail.available
    assert "not available" in detail.note and "guessed" in detail.note

    # A percentage that names no base at all is not a threshold yet.
    detail = materiality.threshold_detail(FakeProfile("1%"), ledger)
    assert detail.threshold is None
    assert "names no base" in detail.note

    # A base this engine cannot compute is refused rather than approximated.
    detail = materiality.threshold_detail(FakeProfile("1% of gross margin"), ledger)
    assert detail.threshold is None and "gross margin" in detail.note

    # And a flat amount needs no ledger at all.
    assert materiality.threshold_for(FakeProfile("500.00")) == money("500.00")


# ------------------------------------------------------- 2. the roll-up

def test_roll_up_combines_the_small_and_leaves_the_material_alone():
    entries = [
        entry("JE-1", D(2025, 1, 12), "42.00"),
        entry("JE-2", D(2025, 2, 12), "38.50"),
        entry("JE-3", D(2025, 3, 12), "45.25"),
        entry("JE-4", D(2025, 2, 20), "7500.00"),               # material
        entry("JE-5", D(2025, 3, 3), "120.00", debit="6300"),   # different account
    ]
    result = materiality.roll_up(entries, money("500.00"))

    assert result.summary["before"] == 5
    assert result.summary["after"] == 3
    assert result.saved_count == 2, "three small entries became one"
    assert result.summary["entries_replaced"] == 3

    numbers = [e.number for e in result.entries]
    assert "JE-4" in numbers, "a material entry is never rolled up"
    assert "JE-5" in numbers, "a lone small entry on other accounts stays itself"
    assert not any(n.startswith("JE-1") or n == "JE-2" or n == "JE-3" for n in numbers)

    rolled = [e for e in result.entries if e.number.startswith("RU-")]
    assert len(rolled) == 1
    combined = rolled[0]
    combined.check()
    assert combined.total_debits() == money("125.75"), "42.00 + 38.50 + 45.25"
    assert combined.total_debits() == combined.total_credits()
    assert combined.date == D(2025, 3, 12), "dated at the last entry it replaces"
    assert combined.kind == "bank_fees"
    assert [l[0] for l in combined.lines] == ["6410", "1010"], "the accounts do not move"

    # The benefit is a number, so it can be argued with.
    assert "5 entries became 3" in result.note and "2 fewer to type" in result.note


def test_roll_up_never_crosses_a_fiscal_year_end():
    """Combining across a year end moves income between years. Never quietly."""
    entries = [
        entry("JE-DEC1", D(2025, 12, 20), "60.00"),
        entry("JE-DEC2", D(2025, 12, 28), "70.00"),
        entry("JE-JAN1", D(2026, 1, 4), "65.00"),
        entry("JE-JAN2", D(2026, 1, 9), "55.00"),
    ]
    result = materiality.roll_up(entries, money("500.00"), period="quarter")

    rolled = [e for e in result.entries if e.number.startswith("RU-")]
    assert len(rolled) == 2, "one summary for 2025, one for 2026, never one for both"
    for combined in rolled:
        years = {int(p.split("-")[0]) for p in combined.basis.split()
                 if p[:4].isdigit() and p.count("-") == 2}
        assert len(years) == 1, f"a summary spans two fiscal years: {combined.basis}"
    assert {e.date.year for e in rolled} == {2025, 2026}
    assert sum(e.total_debits() for e in rolled) == money("250.00"), "nothing was lost"

    # A non-calendar year end moves the boundary, and the roll-up moves with
    # it. These four sit inside one calendar quarter, so a December year end
    # states them once. A January 31 year end splits them down the middle,
    # because January is the last month of one year and February is the first
    # month of the next, and combining them would move income between years.
    straddle = [
        entry("JE-JAN1", D(2025, 1, 8), "40.00"),
        entry("JE-JAN2", D(2025, 1, 24), "44.00"),
        entry("JE-FEB1", D(2025, 2, 6), "41.00"),
        entry("JE-FEB2", D(2025, 2, 19), "45.00"),
    ]
    calendar_fy = materiality.roll_up(straddle, money("500.00"))
    assert len([e for e in calendar_fy.entries if e.number.startswith("RU-")]) == 1
    assert calendar_fy.saved_count == 3

    jan_fy = materiality.roll_up(straddle, money("500.00"), fiscal_year_end_month=1)
    combined = [e for e in jan_fy.entries if e.number.startswith("RU-")]
    assert len(combined) == 2, "nothing may be combined across a January 31 year end"
    assert jan_fy.saved_count == 2
    assert sum(e.total_debits() for e in combined) == money("170.00"), "nothing was lost"
    assert {e.date.month for e in combined} == {1, 2}


def test_the_rolled_basis_names_every_entry_it_replaced():
    entries = [
        entry("JE-A", D(2025, 4, 2), "31.00", basis="Northgate wire fee, April statement"),
        entry("JE-B", D(2025, 4, 19), "31.00", basis="Northgate wire fee, April statement"),
        entry("JE-C", D(2025, 5, 6), "18.75", basis="Northgate returned item fee, May statement"),
    ]
    result = materiality.roll_up(entries, money("500.00"))
    combined = [e for e in result.entries if e.number.startswith("RU-")][0]

    for number, date, amount in (("JE-A", "2025-04-02", "31.00"),
                                 ("JE-B", "2025-04-19", "31.00"),
                                 ("JE-C", "2025-05-06", "18.75")):
        assert number in combined.basis, f"{number} is not recoverable from the summary"
        assert date in combined.basis, f"{number}'s date is missing"
        assert f"{number} {date} {amount}" in combined.basis, f"{number}'s amount is missing"

    assert "500.00" in combined.basis, "the threshold that justified it is stated"
    assert "Northgate wire fee" in combined.basis, "the original bases are carried forward"
    assert "Northgate returned item fee" in combined.basis

    # And the machine-readable side says the same thing.
    replaced = result.rolled[0]["replaced"]
    assert [r["number"] for r in replaced] == ["JE-A", "JE-B", "JE-C"]
    assert [r["amount"] for r in replaced] == ["31.00", "31.00", "18.75"]


def test_no_threshold_means_no_roll_up_and_a_sentence_saying_why():
    entries = [entry("JE-1", D(2025, 1, 12), "1.00"),
               entry("JE-2", D(2025, 1, 13), "2.00"),
               entry("JE-3", D(2025, 1, 14), "3.00")]

    result = materiality.roll_up(entries, None)
    assert result.saved_count == 0
    assert [e.number for e in result.entries] == ["JE-1", "JE-2", "JE-3"]
    assert result.rolled == []
    assert "No materiality threshold" in result.note
    assert "not permission" in result.note, "absence of a fact is not a licence"

    # The profile route says the same thing, and says what to do about it.
    detail = materiality.threshold_detail(FakeProfile(""))
    assert detail.threshold is None and not bool(detail)
    assert "missing engagement fact" in detail.note and "entity.materiality" in detail.note

    # An unresolved percentage is refused rather than treated as an amount:
    # comparing "1" percent against a 1.00 entry would roll up the wrong things.
    try:
        materiality.roll_up(entries, materiality.parse_threshold("1%"))
    except MaterialityError as exc:
        assert "has not been resolved" in str(exc)
    else:
        raise AssertionError("a percentage was compared to an entry as if it were an amount")


def test_roll_up_never_combines_across_different_accounts():
    """Where an amount lands is the whole content of the entry."""
    entries = [
        entry("JE-1", D(2025, 1, 5), "20.00", debit="6410"),
        entry("JE-2", D(2025, 1, 6), "20.00", debit="6410"),
        entry("JE-3", D(2025, 1, 7), "20.00", debit="6300"),
        entry("JE-4", D(2025, 1, 8), "20.00", debit="6300"),
    ]
    result = materiality.roll_up(entries, money("500.00"))
    rolled = [e for e in result.entries if e.number.startswith("RU-")]
    assert len(rolled) == 2, "same amounts, same period, different accounts, two summaries"
    assert {l[0] for e in rolled for l in e.lines} == {"6410", "6300", "1010"}
    for combined in rolled:
        accounts = {l[0] for l in combined.lines}
        assert accounts in ({"6410", "1010"}, {"6300", "1010"})

    # Different kinds are different facts too, even on the same accounts.
    mixed = [entry("JE-5", D(2025, 1, 5), "20.00", kind="bank_fees"),
             entry("JE-6", D(2025, 1, 6), "20.00", kind="wind_down")]
    assert materiality.roll_up(mixed, money("500.00")).saved_count == 0

    # And by="account_class" holds the class dimension as well, so a summary
    # never merges two classes into one untagged entry.
    classed = [
        ProposedEntry(date=D(2025, 1, 5), number="JE-C1", kind="bank_fees",
                      basis="invented test fixture", batch_tag="01",
                      lines=[("6410", "20.00", "0.00", "fee", "", "Hardware"),
                             ("1010", "0.00", "20.00", "fee", "", "Hardware")]),
        ProposedEntry(date=D(2025, 1, 6), number="JE-C2", kind="bank_fees",
                      basis="invented test fixture", batch_tag="01",
                      lines=[("6410", "24.00", "0.00", "fee", "", "Software"),
                             ("1010", "0.00", "24.00", "fee", "", "Software")]),
    ]
    assert materiality.roll_up(classed, money("500.00"), by="account_class").saved_count == 0
    loose = materiality.roll_up(classed, money("500.00"), by="account")
    assert loose.saved_count == 1, "on accounts alone the two classes do combine"
    assert [l[5] for l in loose.entries[0].lines] == ["", ""], (
        "a combined entry does not claim a class it cannot evidence")
    assert "more than one class" in loose.entries[0].basis, (
        "a dropped class is stated, never lost quietly")
    assert "Hardware, Software" in loose.entries[0].basis

    # Where every member agrees on the class, the class is carried through
    # rather than thrown away for being inconvenient.
    for e in classed:
        e.lines = [(a, d, c, m, n, "Hardware") for a, d, c, m, n, _ in e.lines]
    agreed = materiality.roll_up(classed, money("500.00"), by="account")
    assert [l[5] for l in agreed.entries[0].lines] == ["Hardware", "Hardware"]

    try:
        materiality.roll_up(classed, money("500.00"), by="whatever")
    except MaterialityError as exc:
        assert "must be one of" in str(exc)
    else:
        raise AssertionError("an unknown grouping was accepted")


# -------------------------------------- 3. the recurring family and the stub

def test_a_twelve_month_family_with_a_stub_is_one_family_with_the_stub_marked(tmp_path):
    path = tmp_path / "worksheet-batch-01.xlsx"
    wf = je_worksheet.write_worksheet(path, amortization_family(), resolve=resolve,
                                      batch_tag="01", meta={"company": "Acme Robotics Inc."})

    assert wf.entries == 12
    assert wf.repeating_groups == 1, "the rounding stub does not get its own family"
    assert wf.repeating_entries == 12
    assert wf.repeating_exact == 11
    assert wf.repeating_needing_edit == 1

    rows = worksheet_rows(path)
    text = [str(r[1]) for r in rows if r[1]]
    notes = [t for t in text if t.startswith(je_worksheet.RECURRING_NOTE_PREFIX)]
    assert len(notes) == 1, "one template covers all twelve"
    note = notes[0]
    assert "Make recurring" in note
    assert "right for 11 of the 12" in note, "how many the template covers exactly"
    assert "The other 1 needs the amount changed" in note, "how many need an edit"
    # The per-month amounts are in the block, so nothing has to be looked up.
    assert "01/28/2025 208.33" in note and "12/28/2025 208.37 (differs)" in note
    assert note.count("(differs)") == 1

    heads = [t for t in text if t.startswith("Entry ")]
    assert len(heads) == 12
    marked = [h for h in heads if je_worksheet.DIFFERS_MARK in h]
    assert len(marked) == 1 and marked[0].startswith("Entry 12 of 12")
    assert je_worksheet.DIFFERS_MARK in str(rows[[r[1] for r in rows].index(marked[0])][4])

    # The whole point: the family still gets the template, so the batch is
    # cheaper than typing twelve entries from scratch.
    scratch = 12 * (je_worksheet.SECONDS_PER_ENTRY + 2 * je_worksheet.SECONDS_PER_LINE)
    assert wf.seconds < scratch

    # Markdown says the same thing, because the terminal reader gets no less.
    md = je_worksheet.render_markdown(amortization_family(), resolve=resolve, batch_tag="01")
    assert md.count("- [ ] **Entry") == 12
    assert md.count(je_worksheet.DIFFERS_MARK) >= 1
    assert "208.37" in md


def test_a_family_stays_one_family_only_when_the_accounts_match():
    """Looser on amounts, not looser on accounts."""
    entries = amortization_family()
    entries.append(ProposedEntry(
        date=D(2025, 6, 28), number="JE-9001",
        lines=[("6120", "208.33", "0.00", "same amount, wrong credit account"),
               ("2100", "0.00", "208.33", "same amount, wrong credit account")],
        kind="prepaid", basis="invented test fixture", batch_tag="01"))
    blocks = je_worksheet.plan(entries, resolve=resolve, batch_tag="01")
    families = [f for b in blocks for f in b["families"]]
    assert len([f for f in families if f["repeating"]]) == 1
    assert sorted(len(f["views"]) for f in families) == [1, 12]


# ------------------------------------------------------- 4. resumability

def test_read_ticks_on_a_part_typed_worksheet_reports_the_right_position(tmp_path):
    path = tmp_path / "worksheet-batch-02.xlsx"
    entries = amortization_family() + [
        entry("JE-8001", D(2025, 3, 31), "1200.00", debit="6300", kind="payroll"),
        entry("JE-8002", D(2025, 6, 30), "900.00", debit="6300", kind="payroll"),
    ]
    wf = je_worksheet.write_worksheet(path, entries, resolve=resolve, batch_tag="02")
    assert wf.entries == 14

    tick(path, set(range(1, 10)))          # nine saved, then the phone rang
    progress = je_worksheet.read_ticks(path)

    assert progress.total == 14
    assert progress.done == 9
    assert progress.remaining == 5
    assert progress.next_index == 10
    assert progress.next_number == "JE-2010"
    assert "prepaid" in progress.next_label and "10/28/2025" in progress.next_label
    assert CHART["6120"] in progress.next_label
    assert not progress.finished

    # The estimate is of what is LEFT, and it is smaller than the whole job.
    assert 0 < progress.seconds_remaining < wf.seconds
    assert progress.estimate.startswith("about")
    assert "9 of 14" in progress.line() and "5 left" in progress.line()

    # The stub is still flagged on the way past, so a resumed session does not
    # copy the template's amount into the month that needs a different one.
    stub = [e for e in progress.entries if e["number"] == "JE-2012"][0]
    assert stub["differs"] is True and stub["done"] is False

    # Finishing the rest reads as finished, with nothing left to estimate.
    tick(path, set(range(1, 15)))
    done = je_worksheet.read_ticks(path)
    assert done.done == 14 and done.remaining == 0 and done.finished
    assert done.next_index == 0 and done.next_label == ""
    assert done.seconds_remaining == 0
    assert "not proof the entries are in QuickBooks" in done.line()


def test_read_ticks_on_an_untouched_worksheet_reports_zero_done(tmp_path):
    path = tmp_path / "worksheet-batch-03.xlsx"
    entries = [entry("JE-7001", D(2025, 1, 31), "310.00"),
               entry("JE-7002", D(2025, 2, 28), "410.00"),
               entry("JE-7003", D(2025, 3, 31), "510.00")]
    wf = je_worksheet.write_worksheet(path, entries, resolve=resolve, batch_tag="03")

    progress = je_worksheet.read_ticks(path)
    assert progress.done == 0
    assert progress.total == 3 and progress.remaining == 3
    assert progress.next_index == 1 and progress.next_number == "JE-7001"
    assert progress.unmarked == 0, "an untouched worksheet is full of [ ], not blanks"
    assert progress.seconds_remaining == wf.seconds, "nothing done, nothing off the estimate"
    assert "0 of 3" in progress.line()

    # A blank tick box counts as NOT done. Both errors cost something; typing an
    # entry twice shows up in the register, never typing it does not.
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb[je_worksheet.SHEET]
    for row in ws.iter_rows():
        if str(row[1].value or "").startswith("Entry 1 of"):
            row[0].value = None
    wb.save(path)
    wb.close()

    blanked = je_worksheet.read_ticks(path)
    assert blanked.done == 0 and blanked.next_index == 1
    assert blanked.unmarked == 1
    assert "counted as not done" in blanked.line()


# ------------------------------------------------------ 5. the header block

def test_the_header_says_the_count_the_estimate_and_the_importer_alternative(tmp_path):
    path = tmp_path / "worksheet-batch-04.xlsx"
    entries = [entry(f"JE-6{i:03d}", D(2025, 1, 1) + dt.timedelta(days=i * 7), "310.00")
               for i in range(40)]
    wf = je_worksheet.write_worksheet(path, entries, resolve=resolve, batch_tag="04",
                                      meta={"company": "Acme Robotics Inc."})
    assert wf.entries == 40

    rows = worksheet_rows(path)
    header = " ".join(str(r[1]) for r in rows[:14] if r[1])

    cost = [str(r[1]) for r in rows if r[1] and str(r[1]).startswith(je_worksheet.COST_PREFIX)]
    assert len(cost) == 1, "one line that prices the job, at the top, before any entry"
    assert "40 entries" in cost[0]
    assert "80 lines" in cost[0]
    assert wf.estimate in cost[0], "the same estimate the caller was handed"
    assert wf.estimate.startswith("about"), "an estimate is said as an estimate"

    assert je_worksheet.IMPORTER_NOTE in header
    assert "QuickBooks App Store" in header
    assert "priced per month" in header and "a single month is usually enough" in header
    # Their decision, not our recommendation, and no claim we cannot stand behind.
    assert "pick whichever" in je_worksheet.IMPORTER_NOTE
    assert "$" not in je_worksheet.IMPORTER_NOTE, "no price we have not verified"
    for named_app in ("SaasAnt", "Transaction Pro", "Business Importer", "Zed Axis"):
        assert named_app.lower() not in header.lower(), "no app is named"

    # Markdown carries the same two facts.
    md = je_worksheet.render_markdown(entries[:3], resolve=resolve, batch_tag="04")
    assert je_worksheet.COST_PREFIX in md and je_worksheet.IMPORTER_NOTE in md


# ------------------------------------ 6. the two halves working together

def test_a_forty_entry_batch_gets_shorter_and_stays_recoverable(tmp_path):
    """The whole point, end to end: shorter list, and every detail still there."""
    entries = amortization_family()                                   # 12 repeats
    entries += [entry(f"JE-F{i:02d}", D(2025, 1, 3) + dt.timedelta(days=i * 11), "38.50")
                for i in range(20)]                                   # 20 small bank fees
    entries += [entry(f"JE-M{i:02d}", D(2025, 2, 14) + dt.timedelta(days=i * 40),
                      "4200.00", debit="6300", kind="payroll") for i in range(8)]
    assert len(entries) == 40

    threshold = materiality.threshold_for(FakeProfile("500.00"))
    result = materiality.roll_up(entries, threshold, period="quarter")
    assert result.summary["before"] == 40
    assert result.summary["after"] < 40 and result.saved_count > 0

    # Nothing material was touched, and every replaced entry is still named.
    survivors = {e.number for e in result.entries}
    assert all(f"JE-M{i:02d}" in survivors for i in range(8))
    named = " ".join(e.basis for e in result.entries if e.number.startswith("RU-"))
    for i in range(20):
        number = f"JE-F{i:02d}"
        assert number in survivors or number in named, f"{number} vanished"

    before = je_worksheet.write_worksheet(tmp_path / "before-batch-01.xlsx", entries,
                                         resolve=resolve, batch_tag="01")
    after = je_worksheet.write_worksheet(tmp_path / "after-batch-01.xlsx", result.entries,
                                        resolve=resolve, batch_tag="01")
    assert after.entries == result.summary["after"]
    assert after.seconds < before.seconds, "a shorter list is a shorter afternoon"

    # And the shorter one is still resumable.
    tick(tmp_path / "after-batch-01.xlsx", {1, 2, 3})
    progress = je_worksheet.read_ticks(tmp_path / "after-batch-01.xlsx")
    assert progress.done == 3
    assert progress.remaining == after.entries - 3
    assert progress.next_index == 4


# --------------------------------------------------------------- runner

def _main():
    import inspect
    import tempfile
    import traceback

    tests = [(n, f) for n, f in sorted(globals().items())
             if n.startswith("test_") and callable(f)]
    failed = []
    for name, fn in tests:
        root = Path(tempfile.mkdtemp(prefix="ctb-worksheet-"))
        try:
            if "tmp_path" in inspect.signature(fn).parameters:
                fn(root)
            else:
                fn()
            print(f"  ok   {name}")
        except Exception:
            failed.append(name)
            print(f"  FAIL {name}")
            traceback.print_exc()
    print(f"\n{len(tests) - len(failed)} passed, {len(failed)} failed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(_main())
