"""Tests for the four QuickBooks import writers.

Every fixture here is invented. The company is Acme Robotics Inc., the vendors
are made up, and no file in this test suite came from a real bank or a real
client.

Run either way:
    pytest tests/test_writers.py
    python3 tests/test_writers.py
"""

from __future__ import annotations

import csv
import datetime as dt
import re
import os
import sys
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "lib"))

from closethebooks import bank_csv, je_csv, je_worksheet, review_workbook, rules_xlsx  # noqa: E402
from closethebooks.model import BankLine, ProposedEntry, Proposal, Rule  # noqa: E402
from closethebooks.util import money, plain, write_csv  # noqa: E402

D = dt.date

# The two payloads. The first is the classic DDE command execution string, the
# second is the one people forget: a bare arithmetic expression is a formula in
# Excel too, and it starts with a character no filter thinks about.
INJECT_CMD = "=cmd|' /C calc'!A0"
INJECT_MINUS = "-1+1"
FORMULA_LEADS = ("=", "+", "-", "@", "\t", "\r")

CHART = {
    "4015": "Current Assets:Mercury Checking (4015)",
    "6120": "Operating Expenses:Software Subscriptions",
    "6300": "Operating Expenses:Contract Labor",
    "2100": "Liabilities:Accrued Payroll",
}


def resolve(key):
    return CHART.get(str(key))


# --------------------------------------------------------------- fixtures

def acme_lines():
    return [
        BankLine(date=D(2026, 1, 5), descriptor="STRIPE TRANSFER ACME ROBOTICS", amount="4820.15"),
        BankLine(date=D(2026, 1, 7), descriptor="BOLTWORKS SUPPLY CO INVOICE 5512", amount="-1290.00"),
        BankLine(date=D(2026, 1, 9), descriptor="GEARHOUSE CLOUD MONTHLY", amount="-249.00"),
        BankLine(date=D(2026, 1, 22), descriptor="REFUND GEARHOUSE CLOUD", amount="249.00"),
    ]


def acme_entry(number="JE-1001", basis="Gearhouse Cloud invoice 2026-01, filed in source-documents"):
    return ProposedEntry(
        date=D(2026, 1, 31),
        number=number,
        lines=[
            ("6120", "249.00", "0.00", "Gearhouse Cloud, January"),
            ("4015", "0.00", "249.00", "Gearhouse Cloud, January"),
        ],
        memo="January software",
        kind="prepaid",
        basis=basis,
        batch_tag="01",
    )


def acme_amortization(month):
    """The same entry in a different month: the Make recurring case."""
    return ProposedEntry(
        date=D(2026, month, 28),
        number=f"JE-2{month:03d}",
        lines=[
            ("6120", "100.00", "0.00", f"Amortize prepaid, month {month}"),
            ("2100", "0.00", "100.00", f"Amortize prepaid, month {month}"),
        ],
        kind="prepaid",
        basis="12-month prepaid schedule, signed order form in source-documents",
        batch_tag="01",
    )


def acme_rules():
    return [
        Rule(id="gearhouse-cloud", account="6120", match_contains=("GEARHOUSE CLOUD",),
             direction="out", source="17 of 17 historical rows", confidence=1.0, support=17),
        Rule(id="boltworks", account="6300", match_contains=("BOLTWORKS SUPPLY",),
             direction="out", klass="Hardware", source="9 of 10 historical rows",
             confidence=0.9, support=9, conflicts=1),
    ]


def acme_proposals():
    lines = acme_lines()
    return [
        Proposal(line=lines[0], action="match", account="4015", account_full=CHART["4015"],
                 rule_id="stripe-payout", confidence=0.98, source="deposit ties to Stripe payout 2026-01-05"),
        Proposal(line=lines[1], action="add", account="6300", account_full=CHART["6300"],
                 klass="Hardware", rule_id="boltworks", confidence=0.9,
                 source="9 of 10 historical rows"),
        Proposal(line=lines[2], action="add", account="6120", account_full=CHART["6120"],
                 rule_id="gearhouse-cloud", confidence=1.0, source="17 of 17 historical rows"),
        Proposal(line=lines[3], action="question", question="Is this refund for the January charge?",
                 confidence=0.0, source="no rule matched", needs_human=True),
    ]


# ------------------------------------------------------------ shared checks

def read_csv(path):
    with open(path, newline="", encoding="utf-8") as fh:
        return list(csv.reader(fh))


# A plain negative number is not an injection: "-12.34" in an Amount column is
# the amount, and escaping it would break the import. The invariant is that
# every field which is NOT a bare number must be inert.
NUMERIC = re.compile(r"^-?\d+(\.\d+)?$")


def assert_no_live_formulas_csv(path):
    for row in read_csv(path):
        for field in row:
            if not field or NUMERIC.match(field):
                continue
            assert field[0] not in FORMULA_LEADS, f"live formula in CSV: {field!r}"


def assert_no_live_formulas_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=False)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v:
                    assert v[0] not in FORMULA_LEADS, f"live formula in {ws.title}!{cell.coordinate}: {v!r}"
    wb.close()


# ------------------------------------------------------------ bank_csv

def test_three_column_round_trip(tmp_path):
    path = tmp_path / "acme-jan-batch-01.csv"
    n = bank_csv.write_three_column(path, acme_lines())
    assert n == 4
    rows = read_csv(path)
    assert rows[0] == list(bank_csv.THREE_COLUMN_HEADER)
    assert rows[1] == ["01/05/2026", "STRIPE TRANSFER ACME ROBOTICS", "4820.15"]
    assert rows[2][2] == "-1290.00"          # money out stays negative
    assert rows[4][2] == "249.00"            # the refund is money in
    assert sum(money(r[2]) for r in rows[1:]) == money("3530.15")


def test_four_column_round_trip(tmp_path):
    path = tmp_path / "acme-jan-4col-batch-01.csv"
    n = bank_csv.write_four_column(path, acme_lines())
    assert n == 4
    rows = read_csv(path)
    assert rows[0] == list(bank_csv.FOUR_COLUMN_HEADER)
    # Credit is money in, Debit is money out, and never both on one row.
    assert rows[1] == ["01/05/2026", "STRIPE TRANSFER ACME ROBOTICS", "4820.15", ""]
    assert rows[2] == ["01/07/2026", "BOLTWORKS SUPPLY CO INVOICE 5512", "", "1290.00"]
    for r in rows[1:]:
        assert (r[2] == "") != (r[3] == ""), r
    credits = sum(money(r[2]) for r in rows[1:])
    debits = sum(money(r[3]) for r in rows[1:])
    assert credits - debits == money("3530.15")


def test_write_dispatches_on_layout(tmp_path):
    three, four = tmp_path / "a.csv", tmp_path / "b.csv"
    bank_csv.write(three, acme_lines(), layout="three")
    bank_csv.write(four, acme_lines(), layout="four")
    assert read_csv(three)[0] == list(bank_csv.THREE_COLUMN_HEADER)
    assert read_csv(four)[0] == list(bank_csv.FOUR_COLUMN_HEADER)
    try:
        bank_csv.write(tmp_path / "c.csv", acme_lines(), layout="five")
    except ValueError as exc:
        assert "five" in str(exc)
    else:
        raise AssertionError("an unknown layout should raise")


def test_render_matches_util_write_csv(tmp_path):
    """The size guard measures `render`, so it has to be byte identical."""
    lines = acme_lines()
    rows = bank_csv.three_column_rows(lines)
    a, b = tmp_path / "render.csv", tmp_path / "util.csv"
    a.write_bytes(bank_csv.render(bank_csv.THREE_COLUMN_HEADER, rows).encode("utf-8"))
    write_csv(b, list(bank_csv.THREE_COLUMN_HEADER), rows, safe_columns=bank_csv.SAFE_COLUMNS)
    assert a.read_bytes() == b.read_bytes()


def test_row_limit_refused_and_split(tmp_path):
    lines = []
    for i in range(bank_csv.MAX_ROWS + 1):
        lines.append(BankLine(date=D(2026, 1, 1) + dt.timedelta(days=i % 28),
                              descriptor=f"BOLTWORKS SUPPLY CO INVOICE {i}", amount="-10.00"))
    path = tmp_path / "too-many-batch-01.csv"
    try:
        bank_csv.write_three_column(path, lines)
    except bank_csv.UploadTooLarge as exc:
        assert "split_for_upload" in str(exc)
    else:
        raise AssertionError("over the row limit should refuse")
    assert not path.exists(), "a refused write must not leave a file behind"

    chunks = bank_csv.split_for_upload(lines)
    assert [len(c) for c in chunks] == [bank_csv.MAX_ROWS, 1]
    assert sum(len(c) for c in chunks) == len(lines)
    assert [l.descriptor for c in chunks for l in c] == [l.descriptor for l in lines]
    for i, chunk in enumerate(chunks):
        bank_csv.write_three_column(tmp_path / f"part-{i}-batch-01.csv", chunk)


def test_byte_limit_refused_and_split(tmp_path):
    """1,000 rows can still be too big: the size limit is the binding one."""
    fat = "GEARHOUSE CLOUD MONTHLY " + ("X" * 400)
    lines = [BankLine(date=D(2026, 2, 1), descriptor=f"{fat} {i}", amount="-1.00")
             for i in range(900)]
    try:
        bank_csv.write_three_column(tmp_path / "fat-batch-01.csv", lines)
    except bank_csv.UploadTooLarge as exc:
        assert "byte" in str(exc)
    else:
        raise AssertionError("over the size limit should refuse")

    chunks = bank_csv.split_for_upload(lines)
    assert len(chunks) > 1
    assert sum(len(c) for c in chunks) == len(lines)
    for i, chunk in enumerate(chunks):
        p = tmp_path / f"fat-{i}-batch-01.csv"
        bank_csv.write_three_column(p, chunk)
        assert os.path.getsize(p) <= bank_csv.MAX_BYTES


def test_bank_refuses_blank_description_and_zero(tmp_path):
    for bad in (BankLine(date=D(2026, 1, 5), descriptor="   ", amount="10.00"),
                BankLine(date=D(2026, 1, 5), descriptor="GEARHOUSE CLOUD", amount="0")):
        try:
            bank_csv.write_three_column(tmp_path / "bad-batch-01.csv", [bad])
        except bank_csv.BankLineError:
            pass
        else:
            raise AssertionError(f"should refuse {bad.descriptor!r} / {bad.amount}")


# ------------------------------------------------------------- je_csv

def test_journal_entry_balances(tmp_path):
    path = tmp_path / "je-batch-01.csv"
    n = je_csv.write_entries(path, [acme_entry()], resolve=resolve)
    assert n == 2
    rows = read_csv(path)
    assert rows[0] == list(je_csv.HEADER)
    body = rows[1:]
    assert {r[0] for r in body} == {"JE-1001"}, "one journal number across the lines"
    assert {r[1] for r in body} == {"01/31/2026"}
    assert [r[2] for r in body] == [CHART["6120"], CHART["4015"]]
    debits = sum(money(r[3]) for r in body)
    credits = sum(money(r[4]) for r in body)
    assert debits == credits == money("249.00")
    for r in body:
        assert (r[3] == "") != (r[4] == ""), "exactly one side per row"
        assert "[batch-01]" in r[5], "every description carries its batch tag"


def test_journal_entry_refuses_unbalanced(tmp_path):
    bad = acme_entry()
    bad.lines[1] = ("4015", "0.00", "248.00", "off by a dollar")
    try:
        je_csv.write_entries(tmp_path / "je-batch-01.csv", [bad], resolve=resolve)
    except ValueError as exc:
        assert "249.00" in str(exc) and "248.00" in str(exc)
    else:
        raise AssertionError("an unbalanced entry must be refused")
    assert not (tmp_path / "je-batch-01.csv").exists()


def test_journal_entry_refuses_no_basis(tmp_path):
    bad = acme_entry(basis="")
    try:
        je_csv.write_entries(tmp_path / "je-batch-01.csv", [bad], resolve=resolve)
    except ValueError as exc:
        assert "basis" in str(exc)
    else:
        raise AssertionError("an entry with no basis must be refused")


def test_journal_entry_refuses_no_lines(tmp_path):
    bad = acme_entry()
    bad.lines = []
    try:
        je_csv.write_entries(tmp_path / "je-batch-01.csv", [bad], resolve=resolve)
    except ValueError as exc:
        assert "no lines" in str(exc)
    else:
        raise AssertionError("an entry with no lines must be refused")


def test_journal_entry_refuses_unresolved_account(tmp_path):
    bad = acme_entry()
    bad.lines[0] = ("9999", "249.00", "0.00", "an account that is not in the chart")
    try:
        je_csv.write_entries(tmp_path / "je-batch-01.csv", [bad], resolve=resolve)
    except je_csv.UnresolvedAccount as exc:
        assert "9999" in str(exc)
    else:
        raise AssertionError("an unresolved account key must be refused")
    assert not (tmp_path / "je-batch-01.csv").exists()


def test_journal_entry_refuses_duplicate_number(tmp_path):
    try:
        je_csv.write_entries(tmp_path / "je-batch-01.csv",
                             [acme_entry(), acme_entry()], resolve=resolve)
    except je_csv.EntryError as exc:
        assert "JE-1001" in str(exc)
    else:
        raise AssertionError("two entries sharing a journal number must be refused")


def test_journal_entry_refuses_missing_batch_tag(tmp_path):
    e = acme_entry()
    e.batch_tag = ""
    try:
        je_csv.write_entries(tmp_path / "je-batch-01.csv", [e], resolve=resolve)
    except je_csv.EntryError as exc:
        assert "batch" in str(exc)
    else:
        raise AssertionError("no batch tag must be refused")
    # ... unless the caller supplies one for the whole file.
    assert je_csv.write_entries(tmp_path / "je2-batch-07.csv", [e],
                                resolve=resolve, batch_tag="07") == 2
    assert "[batch-07]" in read_csv(tmp_path / "je2-batch-07.csv")[1][5]


def test_split_sure_and_review():
    sure_entry = acme_entry(number="JE-1001")
    guess = acme_entry(number="JE-1002", basis="assumed contract labor from the vendor name")
    marked = acme_entry(number="JE-1003", basis="? waiting on the founder to confirm the split")
    sure, review = je_csv.split_sure_and_review([sure_entry, guess, marked])
    assert [e.number for e in sure] == ["JE-1001"]
    assert [e.number for e in review] == ["JE-1002", "JE-1003"]

    # An explicit confidence, if a future entry carries one, decides instead.
    scored = acme_entry(number="JE-1004")
    scored.confidence = 0.4
    assert je_csv.split_sure_and_review([scored]) == ([], [scored])

    # And a broken entry stops the run rather than being sorted into a pile.
    broken = acme_entry(number="JE-1005", basis="")
    try:
        je_csv.split_sure_and_review([broken])
    except ValueError:
        pass
    else:
        raise AssertionError("splitting must not launder an unpostable entry")


# ----------------------------------------------------------- rules_xlsx

def test_rules_without_a_template_refuses_an_import_file(tmp_path):
    """The corrected design. A rules import file cannot be synthesized.

    The real template is three columns, Rule Name / Rule Condition / Rule
    Outputs, and the last two hold QuickBooks' own generated encoding. So with
    no template to learn from, the writer produces a plain-language list a
    person types in, and says so.
    """
    from openpyxl import load_workbook
    path = tmp_path / "rules-batch-01.xlsx"
    rf = rules_xlsx.write_rules(path, acme_rules(), resolve=resolve)
    assert rf == 2, "still an int of rules written"
    assert rf.kind == "hand-entry"
    assert rf.importable is False
    assert rf.confidence == 0.0
    assert any("no template" in n for n in rf.notes)

    wb = load_workbook(path)
    ws = wb[rules_xlsx.SHEET]
    assert [c.value for c in ws[1]] == list(rules_xlsx.HAND_HEADER)

    # Ordered by how many transactions each rule would catch, so the person
    # typing them gets the valuable ones done before they lose patience.
    assert [ws.cell(row=r, column=8).value for r in (2, 3)] == [17, 9]
    first = [c.value for c in ws[2]]
    assert first[0] == "CTB gearhouse-cloud"
    assert first[1] == "Money out"
    assert first[2] == 'Description contains "GEARHOUSE CLOUD"'
    assert first[3] == CHART["6120"]
    assert first[6] == "No", "auto-add is off, always"
    assert first[8] == rules_xlsx.TICK, "a tick column for the person typing"
    assert [c.value for c in ws[3]][5] == "Hardware"

    # Nothing anywhere in the file pretends to be the import template.
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            assert list(row[:3]) != list(rules_xlsx.HEADER), "no fake import sheet"

    intro = " ".join(
        str(c.value or "") for row in wb[rules_xlsx.INTRO_SHEET].iter_rows() for c in row)
    assert "retroactively" in intro, "rules never touch the existing backlog"
    assert "Batch actions" in intro and "Modify selected" in intro
    assert "Export rules" in intro
    wb.close()


def test_rules_with_a_template_match_its_shape(tmp_path):
    """With the user's own export in hand, copy its columns and its encoding."""
    from openpyxl import Workbook, load_workbook

    template = tmp_path / "exported-rules.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Bank Rules"
    ws.append(list(rules_xlsx.HEADER))
    ws.append(["Coffee", "0~0~STARBUCKS", "1~0~Meals and Entertainment"])
    ws.append(["Cloud", "0~0~AWS", "1~0~Computer and Internet"])
    wb.save(template)
    wb.close()

    shape = rules_xlsx.read_template(template)
    assert shape.columns == rules_xlsx.HEADER
    assert shape.rows_sampled == 2
    assert shape.separators[1] == "~" and shape.slots[1] == 2
    assert 0 < shape.confidence <= rules_xlsx.MAX_CONFIDENCE

    path = tmp_path / "rules-import-batch-01.xlsx"
    rf = rules_xlsx.write_rules(path, acme_rules(), template=template, resolve=resolve)
    assert rf == 2
    assert rf.kind == "import" and rf.importable is True
    assert 0 < rf.confidence <= rules_xlsx.MAX_CONFIDENCE, "confidence is reported, never 1.0"

    wb = load_workbook(path)
    assert wb.sheetnames == ["Bank Rules"], "an import file carries nothing but its table"
    out = wb["Bank Rules"]
    assert [c.value for c in out[1]] == list(rules_xlsx.HEADER), "the template's columns"
    body = [[c.value for c in out[r]] for r in (2, 3)]
    assert body[0][0] == "CTB gearhouse-cloud"
    # The encoding is copied, with our text in the slot the template showed us.
    assert body[0][1] == "0~0~GEARHOUSE CLOUD"
    assert body[0][2] == "1~0~" + CHART["6120"]
    assert body[1][1] == "0~0~BOLTWORKS SUPPLY"
    for row in body:
        assert len(row) == len(rules_xlsx.HEADER)
    wb.close()

    # And the hand-entry list is still written beside it, in case it is rejected.
    assert Path(rf.hand_path).exists()
    wb = load_workbook(rf.hand_path)
    assert [c.value for c in wb[rules_xlsx.SHEET][1]] == list(rules_xlsx.HAND_HEADER)
    wb.close()


def test_rules_template_with_no_rules_in_it_is_refused(tmp_path):
    """A header teaches column names and nothing about the encoding."""
    from openpyxl import Workbook
    template = tmp_path / "empty-export.xlsx"
    wb = Workbook()
    wb.active.append(list(rules_xlsx.HEADER))
    wb.save(template)
    wb.close()
    try:
        rules_xlsx.write_rules(tmp_path / "r-batch-01.xlsx", acme_rules(),
                               template=template, resolve=resolve)
    except rules_xlsx.TemplateUnusable as exc:
        assert "at least two rules" in str(exc)
    else:
        raise AssertionError("an empty template cannot teach an encoding")


def test_rules_markdown_list(tmp_path):
    path = tmp_path / "rules-batch-01.md"
    rf = rules_xlsx.write_rules(path, acme_rules(), resolve=resolve)
    assert rf == 2 and rf.importable is False
    text = path.read_text(encoding="utf-8")
    assert "retroactively" in text
    assert "CTB gearhouse-cloud" in text
    assert text.index("CTB gearhouse-cloud") < text.index("CTB boltworks"), "17 catches first"


def test_rule_names_are_deterministic_and_unique(tmp_path):
    from openpyxl import load_workbook
    twins = [
        Rule(id="gearhouse-cloud", account="6120", match_contains=("GEARHOUSE CLOUD",)),
        Rule(id="gearhouse-cloud", account="6120", match_contains=("GEARHOUSE CLOUD MONTHLY",)),
    ]
    a, b = tmp_path / "r1-batch-01.xlsx", tmp_path / "r2-batch-01.xlsx"
    rules_xlsx.write_rules(a, twins, resolve=resolve)
    rules_xlsx.write_rules(b, twins, resolve=resolve)
    names = []
    for p in (a, b):
        wb = load_workbook(p)
        names.append([wb[rules_xlsx.SHEET].cell(row=r, column=1).value for r in (2, 3)])
        wb.close()
    assert names[0] == names[1], "the same rules must produce the same names"
    assert names[0] == ["CTB gearhouse-cloud", "CTB gearhouse-cloud 2"]
    assert all(len(x) <= rules_xlsx.MAX_RULE_NAME for x in names[0])


def test_rules_refusals(tmp_path):
    regex_only = Rule(id="regex-only", account="6120", match_regex=r"gearhouse.*cloud")
    try:
        rules_xlsx.write_rules(tmp_path / "r-batch-01.xlsx", [regex_only], resolve=resolve)
    except rules_xlsx.RuleNotExportable as exc:
        assert "regular expression" in str(exc)
    else:
        raise AssertionError("a regex-only rule cannot be exported")

    too_many = Rule(id="wide", account="6120",
                    match_contains=tuple(f"VENDOR {i}" for i in range(6)))
    try:
        rules_xlsx.write_rules(tmp_path / "r2-batch-01.xlsx", [too_many], resolve=resolve)
    except rules_xlsx.RuleNotExportable as exc:
        assert str(rules_xlsx.MAX_CONDITIONS) in str(exc)
    else:
        raise AssertionError("more conditions than QuickBooks accepts must be refused")

    unknown = Rule(id="nowhere", account="9999", match_contains=("GEARHOUSE",))
    try:
        rules_xlsx.write_rules(tmp_path / "r3-batch-01.xlsx", [unknown], resolve=resolve)
    except je_csv.UnresolvedAccount as exc:
        assert "9999" in str(exc)
    else:
        raise AssertionError("an unresolved category must be refused")


# --------------------------------------------------------- je_worksheet

def worksheet_rows(path):
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb[je_worksheet.SHEET]
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    wb.close()
    return rows


def test_worksheet_one_block_per_entry_with_running_total_and_tick(tmp_path):
    path = tmp_path / "worksheet-batch-01.xlsx"
    entries = [acme_entry(), acme_amortization(1)]
    wf = je_worksheet.write_worksheet(path, entries, resolve=resolve, batch_tag="01",
                                      meta={"company": "Acme Robotics Inc."})
    assert wf.entries == 2 and wf.lines == 4
    assert wf.seconds > 0 and wf.estimate.startswith("about")

    rows = worksheet_rows(path)
    ticks = [r for r in rows if r[0] == je_worksheet.TICK]
    assert len(ticks) == 2, "one tick per entry, so a break cannot lose your place"
    heads = [str(r[1]) for r in rows if r[0] == je_worksheet.TICK]
    assert heads == ["Entry 1 of 2", "Entry 2 of 2"]

    # The block asks for the fields in the order the QuickBooks screen does.
    labels = [str(r[1]) for r in rows if r[1]]
    first = labels.index("Entry 1 of 2")
    assert labels[first + 1] == "Journal date"
    assert labels[first + 2] == "Journal no."
    assert labels[first + 3] == "Account"

    # A running total on every line, and a bold must-equal line under the block.
    must = [r for r in rows if r[1] == je_worksheet.MUST_EQUAL]
    assert len(must) == 2, "one MUST EQUAL line per entry"
    for r in must:
        assert money(r[2]) == money(r[3]), "debits equal credits on the check line"
    assert money(must[0][2]) == money("249.00")
    assert money(must[1][2]) == money("100.00")

    lines = [r for r in rows if r[1] in (CHART["6120"], CHART["4015"], CHART["2100"])]
    assert len(lines) == 4
    running = [(money(r[7]), money(r[8])) for r in lines]
    assert running[0] == (money("249.00"), money("0.00"))
    assert running[1] == (money("249.00"), money("249.00")), "the running total closes"
    assert running[3] == (money("100.00"), money("100.00"))

    # Every description still carries its batch tag, and the basis is shown.
    assert all("[batch-01]" in str(r[4]) for r in lines)
    body = " ".join(str(v) for r in rows for v in r if v)
    assert "12-month prepaid schedule" in body
    assert je_worksheet.CLICK_PATH in body
    assert "journal entry import is currently not an option" in body


def test_worksheet_groups_repeats_and_says_to_make_them_recurring(tmp_path):
    path = tmp_path / "worksheet-batch-02.xlsx"
    entries = [acme_amortization(m) for m in range(1, 13)]
    entries.insert(6, acme_entry())          # a stranger in the middle of the run
    wf = je_worksheet.write_worksheet(path, entries, resolve=resolve, batch_tag="02")
    assert wf.entries == 13
    assert wf.repeating_groups == 1
    assert wf.repeating_entries == 12

    rows = worksheet_rows(path)
    text = [str(r[1]) for r in rows if r[1]]
    notes = [t for t in text if t.startswith(je_worksheet.RECURRING_NOTE_PREFIX)]
    assert len(notes) == 1, "one note for the one family of identical entries"
    assert "Make recurring" in notes[0]
    assert "11" in notes[0], "the other eleven come off the template"

    # The twelve identical entries sit together, so the template gets reused
    # once rather than being rebuilt after every interruption.
    order = [t for t in text if t.startswith("Entry ")]
    numbers = [str(r[2]) for r in rows if str(r[1]) == "Journal no."]
    amort = [i for i, n in enumerate(numbers) if n.startswith("JE-2")]
    assert amort == list(range(min(amort), min(amort) + 12)), "the repeats are contiguous"
    assert len(order) == 13

    # And the estimate is smaller than typing all thirteen from scratch.
    from_scratch = 13 * (je_worksheet.SECONDS_PER_ENTRY + 2 * je_worksheet.SECONDS_PER_LINE)
    assert wf.seconds < from_scratch


def test_worksheet_markdown_mirrors_the_workbook():
    entries = [acme_amortization(m) for m in range(1, 4)] + [acme_entry()]
    text = je_worksheet.render_markdown(entries, resolve=resolve, batch_tag="01")
    assert text.count("- [ ] **Entry") == 4, "a checkbox per entry"
    assert text.count("**" + je_worksheet.MUST_EQUAL) == 4
    assert je_worksheet.RECURRING_NOTE_PREFIX in text
    assert "Make recurring" in text
    assert CHART["6120"] in text and "[batch-01]" in text
    assert "100.00 debits = 100.00 credits" in text


def test_worksheet_refuses_what_the_import_refuses(tmp_path):
    bad = acme_entry()
    bad.lines[1] = ("4015", "0.00", "248.00", "off by a dollar")
    try:
        je_worksheet.write_worksheet(tmp_path / "w-batch-01.xlsx", [bad], resolve=resolve)
    except ValueError as exc:
        assert "249.00" in str(exc) and "248.00" in str(exc)
    else:
        raise AssertionError("a worksheet is not a lower bar than an import file")
    assert not (tmp_path / "w-batch-01.xlsx").exists()

    unknown = acme_entry()
    unknown.lines[0] = ("9999", "249.00", "0.00", "an account that is not in the chart")
    try:
        je_worksheet.write_worksheet(tmp_path / "w2-batch-01.xlsx", [unknown], resolve=resolve)
    except je_csv.UnresolvedAccount as exc:
        assert "9999" in str(exc)
    else:
        raise AssertionError("an unresolved account would be typed in by hand and created")


def test_formula_injection_worksheet(tmp_path):
    e = ProposedEntry(
        date=D(2026, 3, 31), number="JE-3001",
        lines=[("6120", "12.34", "0.00", INJECT_CMD, INJECT_MINUS, INJECT_CMD),
               ("4015", "0.00", "12.34", INJECT_MINUS)],
        kind="prepaid", basis="invented test fixture", batch_tag="01",
    )
    path = tmp_path / "inject-worksheet-batch-01.xlsx"
    je_worksheet.write_worksheet(path, [e], resolve=resolve, batch_tag="01")
    assert_no_live_formulas_xlsx(path)
    rows = worksheet_rows(path)
    lines = [r for r in rows if r[1] in (CHART["6120"], CHART["4015"])]
    assert lines[0][4].startswith("'" + INJECT_CMD)
    assert lines[0][5] == "'" + INJECT_MINUS
    assert money(lines[0][2]) == money("12.34"), "amounts stay numbers"


# ------------------------------------------------------ je_csv by region

def test_us_csv_also_writes_the_worksheet_and_says_so(tmp_path):
    """QuickBooks Online US has no journal entry import, so never ship only a CSV."""
    path = tmp_path / "je-batch-01.csv"
    ef = je_csv.write_entries(path, [acme_entry()], resolve=resolve)
    assert ef == 2, "still an int of rows written"
    assert ef.region == "US"
    assert ef.importable is False
    assert ef.worksheet and Path(ef.worksheet).exists()
    assert "worksheet" in ef.message and "cannot import" in ef.message
    assert Path(ef.note_path).exists()
    note = Path(ef.note_path).read_text(encoding="utf-8")
    assert "journal entry import is currently not an option" in note
    assert je_csv.INTUIT_JE_URL in note
    assert Path(ef.worksheet).name in note

    rows = worksheet_rows(ef.worksheet)
    assert [r for r in rows if r[0] == je_worksheet.TICK], "the worksheet is a real one"


def test_canada_and_uk_get_an_importable_csv(tmp_path):
    for region in ("CA", "UK", "canada"):
        path = tmp_path / f"je-{region}-batch-01.csv"
        ef = je_csv.write_entries(path, [acme_entry()], resolve=resolve, region=region)
        assert ef == 2
        assert ef.importable is True
        assert ef.worksheet == "", "no worksheet where the import actually works"
        assert "Import data" in ef.message
        assert "NOT work in the US version" in Path(ef.note_path).read_text(encoding="utf-8")

    try:
        je_csv.write_entries(tmp_path / "je-x-batch-01.csv", [acme_entry()],
                             resolve=resolve, region="Narnia")
    except je_csv.UnknownRegion as exc:
        assert "Narnia" in str(exc)
    else:
        raise AssertionError("a region we have no facts about must not be guessed")


# ------------------------------------------------------- review_workbook

def test_review_workbook_round_trip(tmp_path):
    path = tmp_path / "batch-01.xlsx"
    meta = {"company": "Acme Robotics Inc.", "period": "January 2026",
            "prepared_by": "close-the-books"}
    rf = review_workbook.write_review(path, acme_proposals(), "01", meta)
    assert rf.rows == 4
    assert rf.needs_human == 1
    assert rf.batch_tag == "01"
    assert len(rf.row_hash) == 64

    decisions = review_workbook.read_decisions(path)
    assert len(decisions) == 4
    assert decisions[0]["date"] == D(2026, 1, 5)
    assert decisions[0]["amount"] == money("4820.15")
    assert decisions[1]["amount"] == money("-1290.00")
    assert decisions[1]["account"] == CHART["6300"]
    assert decisions[1]["klass"] == "Hardware"
    assert [d["action"] for d in decisions] == ["match", "add", "add", "question"]
    assert all(d["founder_decision"] == "" for d in decisions)
    assert decisions[3]["needs_human"] is True
    assert decisions[3]["question"].startswith(review_workbook.NEEDS_HUMAN_MARKER)

    # The founder fills the column in, in the spreadsheet.
    fill_decisions(path, {0: "approve", 1: "approve", 2: "skip", 3: "answer"})
    back = review_workbook.read_decisions(path)
    assert [d["decision"] for d in back] == ["approve", "approve", "skip", "answer"]
    assert all(d["valid"] for d in back)
    assert review_workbook.unrecognized(back) == []

    fill_decisions(path, {0: "aprove"})
    typo = review_workbook.read_decisions(path)
    assert typo[0]["valid"] is False
    assert [d["row"] for d in review_workbook.unrecognized(typo)] == [typo[0]["row"]]


def fill_decisions(path, by_index):
    """Type into the founder_decision column the way a spreadsheet would."""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb[review_workbook.SHEET]
    header = review_workbook._find_header(ws)
    col = len(review_workbook.COLUMNS)
    for i, value in by_index.items():
        ws.cell(row=header + 1 + i, column=col, value=value)
    wb.save(path)
    wb.close()


def resave(path):
    """Open and save with no edits, the way Excel does when a founder peeks."""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    wb.save(path)
    wb.close()


def test_row_hash_stable_and_sensitive(tmp_path):
    path = tmp_path / "batch-02.xlsx"
    rf = review_workbook.write_review(path, acme_proposals(), "02", {"company": "Acme Robotics Inc."})
    before = rf.row_hash
    assert review_workbook.row_hash(path) == before

    resave(path)
    assert review_workbook.row_hash(path) == before, "an open-and-save must not break an approval"

    # A cell that carries no decision does not invalidate one.
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb[review_workbook.SHEET]
    header = review_workbook._find_header(ws)
    ws.cell(row=header + 1, column=review_workbook.COLUMNS.index("Why") + 1, value="reworded")
    ws.column_dimensions["B"].width = 99
    wb.save(path)
    wb.close()
    assert review_workbook.row_hash(path) == before, "only decision cells are hashed"

    # A decision does.
    fill_decisions(path, {0: "approve"})
    after_decision = review_workbook.row_hash(path)
    assert after_decision != before

    # So does changing what the row would post.
    wb = load_workbook(path)
    ws = wb[review_workbook.SHEET]
    ws.cell(row=header + 2, column=review_workbook.COLUMNS.index("Account") + 1,
            value=CHART["6120"])
    wb.save(path)
    wb.close()
    assert review_workbook.row_hash(path) not in (before, after_decision)


def test_review_workbook_shape(tmp_path):
    from openpyxl import load_workbook
    path = tmp_path / "batch-03.xlsx"
    rf = review_workbook.write_review(path, acme_proposals(), "03",
                                      {"company": "Acme Robotics Inc.", "period": "January 2026"})
    wb = load_workbook(path)
    ws = wb[review_workbook.SHEET]
    assert [ws.cell(row=rf.header_row, column=c).value for c in range(1, 11)] == list(review_workbook.COLUMNS)
    assert ws.freeze_panes == f"A{rf.header_row + 1}"
    # The instruction block sits above the table and says the one thing that matters.
    intro = " ".join(str(ws.cell(row=r, column=1).value or "") for r in range(1, rf.header_row))
    assert "Acme Robotics Inc." in intro
    assert "founder_decision" in intro
    assert "terminal" in intro and "QuickBooks" in intro
    assert "NEEDS HUMAN" in intro
    # A dropdown on the decision column, and the Question column wraps.
    dvs = ws.data_validations.dataValidation
    assert len(dvs) == 1
    assert "approve" in dvs[0].formula1 and "skip" in dvs[0].formula1
    assert str(dvs[0].sqref).startswith("J")
    q = ws.cell(row=rf.header_row + 1, column=review_workbook.COLUMNS.index("Question") + 1)
    assert q.alignment.wrap_text is True
    assert ws.column_dimensions["B"].width and ws.column_dimensions["B"].width > 20
    # The flagged row is visibly filled.
    flagged = ws.cell(row=rf.header_row + 4, column=1)
    assert flagged.fill.fgColor.rgb not in (None, "00000000"), "needs_human rows are shaded"
    assert review_workbook.META_SHEET in wb.sheetnames
    wb.close()


# ------------------------------------------------------- formula injection

def test_formula_injection_bank_csv(tmp_path):
    lines = [
        BankLine(date=D(2026, 3, 1), descriptor=INJECT_CMD, amount="-12.34"),
        BankLine(date=D(2026, 3, 2), descriptor=INJECT_MINUS, amount="56.78"),
    ]
    for layout in ("three", "four"):
        path = tmp_path / f"inject-{layout}-batch-01.csv"
        bank_csv.write(path, lines, layout=layout)
        raw = path.read_text(encoding="utf-8")
        assert "'" + INJECT_CMD in raw
        assert "'" + INJECT_MINUS in raw
        assert_no_live_formulas_csv(path)
        # The escape must not have leaked into the amount columns.
        rows = read_csv(path)
        assert rows[1][2:] in (["-12.34"], ["", "12.34"])
        assert rows[2][2:] in (["56.78"], ["56.78", ""])


def test_formula_injection_je_csv(tmp_path):
    e = ProposedEntry(
        date=D(2026, 3, 31), number="JE-2001",
        lines=[("6120", "12.34", "0.00", INJECT_CMD, INJECT_MINUS, INJECT_CMD),
               ("4015", "0.00", "12.34", INJECT_MINUS)],
        basis="invented test fixture", batch_tag="01",
    )
    path = tmp_path / "inject-je-batch-01.csv"
    je_csv.write_entries(path, [e], resolve=resolve)
    raw = path.read_text(encoding="utf-8")
    assert "'" + INJECT_CMD in raw
    assert "'" + INJECT_MINUS in raw
    assert_no_live_formulas_csv(path)
    rows = read_csv(path)
    assert rows[1][3] == "12.34" and rows[2][4] == "12.34", "amounts stay unescaped"


def test_formula_injection_rules_xlsx(tmp_path):
    from openpyxl import load_workbook
    r = Rule(id=INJECT_MINUS, account="6120", match_contains=(INJECT_CMD,), direction="out")
    r.payee = INJECT_CMD                     # payees come from statement text too
    path = tmp_path / "inject-rules-batch-01.xlsx"
    rules_xlsx.write_rules(path, [r], resolve=resolve)
    assert_no_live_formulas_xlsx(path)
    wb = load_workbook(path)
    ws = wb[rules_xlsx.SHEET]
    assert ws.cell(row=2, column=5).value == "'" + INJECT_CMD
    assert INJECT_CMD in ws.cell(row=2, column=3).value      # the payload survives, inert
    wb.close()


def test_formula_injection_review_workbook(tmp_path):
    from openpyxl import load_workbook
    lines = [
        BankLine(date=D(2026, 3, 1), descriptor=INJECT_CMD, amount="-12.34"),
        BankLine(date=D(2026, 3, 2), descriptor=INJECT_MINUS, amount="56.78"),
    ]
    proposals = [
        Proposal(line=lines[0], action="question", question=INJECT_CMD, needs_human=True,
                 source=INJECT_MINUS),
        Proposal(line=lines[1], action="add", account=INJECT_MINUS, account_full=INJECT_MINUS,
                 klass=INJECT_CMD, confidence=0.5, source="invented test fixture"),
    ]
    path = tmp_path / "inject-batch-01.xlsx"
    review_workbook.write_review(path, proposals, "01",
                                 {"company": INJECT_CMD, "note": INJECT_MINUS})
    assert_no_live_formulas_xlsx(path)
    wb = load_workbook(path)
    ws = wb[review_workbook.SHEET]
    header = review_workbook._find_header(ws)
    assert ws.cell(row=header + 1, column=2).value == "'" + INJECT_CMD
    assert ws.cell(row=header + 2, column=2).value == "'" + INJECT_MINUS
    # Amounts are numbers, not text, so nothing was mangled making them safe.
    assert money(ws.cell(row=header + 1, column=3).value) == money("-12.34")
    wb.close()
    decisions = review_workbook.read_decisions(path)
    assert decisions[0]["description"] == "'" + INJECT_CMD
    assert review_workbook.row_hash(path) == review_workbook.row_hash(path)


# --------------------------------------------------------------- runner

def _main():
    import inspect
    import tempfile
    import traceback

    tests = [(n, f) for n, f in sorted(globals().items())
             if n.startswith("test_") and callable(f)]
    failed = []
    for name, fn in tests:
        root = Path(tempfile.mkdtemp(prefix="ctb-"))
        try:
            if "tmp_path" in inspect.signature(fn).parameters:
                fn(root)
            else:
                fn()
            print(f"  ok   {name}")
        except Exception:
            failed.append(name)
            print(f"  FAIL {name}")
            traceback.print_exc()
    print(f"\n{len(tests) - len(failed)} passed, {len(failed)} failed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(_main())
