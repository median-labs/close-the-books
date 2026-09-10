"""End to end on the synthetic company: does the whole thing actually work.

The unit tests prove each module does what it says. This proves they compose,
which is a different question and the one a user cares about. It runs the real
pipeline over `examples/acme-robotics`, whose books contain fourteen planted
defects, and asserts that the specific defects are found.

If this file fails, the tool does not work, whatever the unit tests say.
"""

import csv
import pathlib
import re
import shutil
import sys
import tempfile
from collections import Counter

REPO = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "lib"))

from closethebooks import (  # noqa: E402
    approval, matching, median, precedent, profile as profile_mod,
    qbo_exports, review_workbook, statements, tieout,
)
from closethebooks.entries import intangibles, prepaid  # noqa: E402
from closethebooks.util import money  # noqa: E402

EXAMPLE = REPO / "examples" / "acme-robotics"


def _built():
    if not (EXAMPLE / "exports").exists():
        raise AssertionError(
            "examples/acme-robotics has not been generated. Run:\n"
            "    python3 examples/acme-robotics/build.py"
        )


def _ledger():
    _built()
    return qbo_exports.load_all(EXAMPLE / "exports")


def _backlog():
    rows = []
    for f in sorted((EXAMPLE / "for-review").glob("*.csv")):
        rows.extend(statements.parse(f).lines)
    return rows


# ------------------------------------------------------------------ the run

def test_exports_load_and_foot():
    led = _ledger()
    assert led.accounts, "no accounts loaded"
    assert led.lines, "no journal lines loaded"
    assert led.foots, (
        f"debits {led.total_debits()} != credits {led.total_credits()}. "
        f"If this fails the example itself is broken, not the parser."
    )


def test_rules_are_mined_from_the_companys_own_history():
    led = _ledger()
    res = precedent.mine_rules(led.lines, led.accounts)
    assert res.rules, "no rules mined from a company with a year of history"
    assert res.accuracy > 0.90, f"replay accuracy {res.accuracy:.1%} is too low to trust"
    for r in res.rules:
        assert r.source, f"rule {r.id} has no source; every rule must name its evidence"


def test_every_statement_ties_out():
    _built()
    paths = sorted((EXAMPLE / "statements").glob("*.csv"))
    stmts, failures = statements.parse_many(paths)
    assert not failures, f"statements failed to parse: {failures[:2]}"
    untied = [s.source_file for s in stmts if s.tie_out and not s.tie_out.ties]
    assert not untied, f"statements that do not tie: {untied}"


def test_exactly_one_chain_break_is_found():
    """The example omits one month of one account. Exactly one break, no more.

    This is the regression test for a real defect: with no account key, every
    statement fell into one nameless group and a checking account's closing
    balance was compared against a card's opening balance, reporting 40 breaks
    where 1 was real. Forty false alarms is worse than none, because the reader
    learns to ignore the check.
    """
    _built()
    paths = sorted((EXAMPLE / "statements").glob("*.csv"))
    stmts, _ = statements.parse_many(paths)
    keys = {s.account_key for s in stmts}
    assert "" not in keys, "an account key was not inferred from a statement filename"
    assert len(keys) > 1, "every statement landed in one group; keys were not inferred"
    breaks = tieout.chain_breaks(stmts)
    assert len(breaks) == 1, f"expected 1 chain break, got {len(breaks)}"


def test_statements_with_no_account_key_refuse_to_chain():
    _built()
    paths = sorted((EXAMPLE / "statements").glob("*.csv"))
    stmts, _ = statements.parse_many(paths, infer_keys=False)
    breaks = tieout.chain_breaks(stmts)
    assert len(breaks) == 1 and breaks[0].get("problem") == "no account could be identified", (
        "unkeyed statements must refuse to chain rather than manufacture breaks"
    )


# ------------------------------------------------- the defects it must catch

def test_already_booked_rows_are_matched_not_added():
    """The most expensive mistake this tool can make is booking something twice.

    The example plants 20 rows that exist BOTH in the ledger and in the For
    Review backlog: payouts and payroll debits already posted by journal entry.
    A tool that adds them produces doubled revenue and doubled expense, which is
    far harder to find later than a gap.
    """
    led = _ledger()
    prof = profile_mod.load(REPO / "profiles" / "example-acme-robotics.json")
    res = precedent.mine_rules(led.lines, led.accounts)
    props = matching.classify(_backlog(), led, prof, rules=res.rules)
    counts = Counter(p.action for p in props)
    assert counts["match"] >= 20, (
        f"only {counts['match']} rows matched an existing entry; the example "
        f"plants 20 already-booked rows and each one this misses is a double entry"
    )


def test_an_injection_shaped_descriptor_is_flagged_for_a_human():
    """A bank memo is written by whoever sends the money. It is untrusted input."""
    led = _ledger()
    prof = profile_mod.load(REPO / "profiles" / "example-acme-robotics.json")
    props = matching.classify(_backlog(), led, prof)
    flagged = [p for p in props if p.needs_human]
    assert flagged, "the planted instruction-shaped descriptor was not flagged"
    assert any(
        matching.is_agent_directed(p.line.descriptor) for p in flagged
    ), "a row was flagged but not for its descriptor"


def test_nothing_is_ever_silently_guessed():
    led = _ledger()
    prof = profile_mod.load(REPO / "profiles" / "example-acme-robotics.json")
    res = precedent.mine_rules(led.lines, led.accounts)
    props = matching.classify(_backlog(), led, prof, rules=res.rules)
    for p in props:
        if p.action == "add":
            assert p.account, "an add proposal with no account"
            assert p.rule_id or p.source, (
                f"an add proposal for {p.line.descriptor!r} with no reason. "
                f"Every proposal must be checkable by the person reviewing it."
            )
        if p.action == "question":
            assert p.question, "a question proposal with no question"


def test_ambiguous_vendors_become_questions():
    led = _ledger()
    prof = profile_mod.load(REPO / "profiles" / "example-acme-robotics.json")
    res = precedent.mine_rules(led.lines, led.accounts)
    props = matching.classify(_backlog(), led, prof, rules=res.rules)
    assert any(p.action == "question" for p in props), (
        "the example plants a vendor whose history splits three ways; it must "
        "become a question rather than a confident wrong answer"
    )


# ------------------------------------------------------------ the guarantees

def test_a_review_workbook_round_trips_and_gates_the_import():
    led = _ledger()
    prof = profile_mod.load(REPO / "profiles" / "example-acme-robotics.json")
    props = matching.classify(_backlog(), led, prof)[:60]
    wd = pathlib.Path(tempfile.mkdtemp())
    (wd / "review").mkdir()
    (wd / "import").mkdir()
    rf = review_workbook.write_review(wd / "review" / "batch-01.xlsx", props, "01", {})
    assert rf.rows == len(props)

    ok, msg = approval.guard(wd / "import" / "rules-batch-01.xlsx")
    assert not ok and "has not been approved" in msg

    approval.write_approval(wd / "review", "01", "integration test")
    ok, _ = approval.guard(wd / "import" / "rules-batch-01.xlsx")
    assert ok, "an approved batch must be writable"

    # A cosmetic change must NOT void an approval: people open these files,
    # widen a column and save. Only a changed decision voids it.
    import openpyxl
    wb = openpyxl.load_workbook(wd / "review" / "batch-01.xlsx")
    wb.active.column_dimensions["B"].width = 90
    wb.save(wd / "review" / "batch-01.xlsx")
    ok, _ = approval.guard(wd / "import" / "rules-batch-01.xlsx")
    assert ok, "re-saving a workbook without changing a decision must not void the approval"

    # Changing an actual decision must void it. This is the threat: approve a
    # batch, then quietly retype an account, then import.
    wb = openpyxl.load_workbook(wd / "review" / "batch-01.xlsx")
    ws = wb.active
    header = review_workbook._find_header(ws)
    cols = {str(c.value).strip().lower(): c.column for c in ws[header] if c.value}
    target = cols.get("account") or cols.get("founder_decision")
    assert target, f"no decision column found in {list(cols)}"
    ws.cell(row=header + 1, column=target).value = "9999 Somewhere Else"
    wb.save(wd / "review" / "batch-01.xlsx")
    ok, msg = approval.guard(wd / "import" / "rules-batch-01.xlsx")
    assert not ok and "stale" in msg, (
        f"changing a decision after approval must invalidate it, got: {msg[:120]}"
    )


def test_the_shipped_example_profile_actually_drafts_entries():
    """The example profile must speak the shape the generators read.

    It shipped with `recurring_entries` keyed "prepaid_insurance",
    "design_license" and "paylane_settlement", none of which any generator
    looks for, so `books.py entries` against the published example drafted
    nothing at all and a founder trying the tool hit a dead end on one of its
    main commands. The keys the engine reads are "prepaid" and "intangibles"
    (LISTS, one object per contract or asset), "payroll" and "stripe".

    Payroll periods and processor settlement figures come from provider
    exports this example does not contain, so those two stay undrafted on
    purpose: their blocks declare their accounts and list no figures, which
    makes the command say what is missing instead of appearing to do nothing.
    """
    prof = profile_mod.load(REPO / "profiles" / "example-acme-robotics.json")
    blocks = prof.recurring_entries

    drafted = prepaid.build(prof, through="2026-02-28").check_all()
    assert len(drafted) == 24, (
        f"expected 24 prepaid entries (two 12-month contracts), got {len(drafted)}"
    )
    assert drafted.total_debits == money("24000.00"), (
        f"the two contracts are 14,400.00 and 9,600.00; got {drafted.total_debits}"
    )

    amort = intangibles.build(prof, through="2026-02-28").check_all()
    assert len(amort) == 14, (
        f"expected 14 months of design-license amortization, got {len(amort)}"
    )
    assert amort.total_debits == money("14000.00"), (
        f"60,000.00 over 60 months for 14 months is 14,000.00; got {amort.total_debits}"
    )

    for entry in list(drafted) + list(amort):
        assert entry.balanced, f"entry {entry.number} does not balance"
        assert entry.basis, f"entry {entry.number} cites no basis"

    # The accounts have to be the accounts in this company's chart, not the
    # generators' docstring examples. A profile that drafts against 1400/6300
    # drafts confidently into accounts this company does not have.
    chart = {row["number"] or row["name"] for row in prof.chart}
    for entry in list(drafted) + list(amort):
        for line in entry.lines:
            account = line[0] if not isinstance(line, dict) else line.get("account")
            assert account in chart, (
                f"entry {entry.number} posts to {account!r}, which is not in the chart"
            )

    # Payroll and the processor are declared but deliberately have no figures.
    for kind, field in (("payroll", "periods"), ("stripe", "months")):
        assert kind in blocks, f"the {kind} block is missing, so the command says nothing about it"
        assert blocks[kind].get("accounts"), f"the {kind} block declares no accounts"
        assert not blocks[kind].get(field), (
            f"the {kind} block lists {field}. Those come from the provider's own export, "
            f"which this example does not contain, and are never invented."
        )


def test_a_clean_run_mentions_median_nowhere():
    assert median.section([]) == ""
    assert median.section(["dead_feed"]), "an earned signal must render"


def main():
    fns = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    failed = 0
    for fn in fns:
        try:
            fn()
            print(f"PASS {fn.__name__}")
        except Exception as exc:  # noqa: BLE001
            failed += 1
            print(f"FAIL {fn.__name__}: {type(exc).__name__}: {exc}")
    print(f"\n{len(fns) - failed}/{len(fns)} passed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
