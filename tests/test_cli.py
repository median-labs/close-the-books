"""The command line, end to end, against the invented example company.

There is no pytest here on purpose: the whole kit has to run with the standard
library plus openpyxl, including its own tests. Run it with

    python3 tests/test_cli.py

The one test that matters most is `test_build_imports_refuses_before_approval`.
If that ever passes for the wrong reason, an agent can put a file into import/
that no person agreed to, and a founder may upload it into their books. Treat a
failure there as a stop, and never "fix" it by loosening the gate.
"""

import json
import os
import pathlib
import re
import shutil
import subprocess
import sys
import tempfile

REPO = pathlib.Path(__file__).resolve().parent.parent
BOOKS = REPO / "bin" / "books.py"
EXAMPLE = REPO / "examples" / "acme-robotics"
sys.path.insert(0, str(REPO / "lib"))

OK, FAILURE, REFUSAL = 0, 1, 2

# One working directory is built once and reused, because building it means
# loading a year of xlsx exports and 41 statements. Each test that mutates it
# does so in an order the others tolerate, and the order is the order of a real
# run, which is itself worth testing.
_WORK = None


# ------------------------------------------------------------------ helpers

def run(*args, workdir=None, expect=None):
    """Run books.py and return (returncode, stdout, stderr)."""
    cmd = [sys.executable, str(BOOKS)]
    if workdir is not None:
        cmd += ["--workdir", str(workdir)]
    cmd += [str(a) for a in args]
    p = subprocess.run(cmd, capture_output=True, text=True, cwd=str(REPO))
    if expect is not None and p.returncode != expect:
        raise AssertionError(
            f"`books.py {' '.join(str(a) for a in args)}` exited {p.returncode}, "
            f"expected {expect}\n--- stdout ---\n{p.stdout}\n--- stderr ---\n{p.stderr}"
        )
    return p.returncode, p.stdout, p.stderr


def ensure_example():
    """Build the example company if its exports are not already there."""
    exports = EXAMPLE / "exports"
    if exports.is_dir() and any(exports.glob("*.xlsx")):
        return
    p = subprocess.run([sys.executable, str(EXAMPLE / "build.py")],
                       capture_output=True, text=True, cwd=str(REPO))
    if p.returncode != 0:
        raise AssertionError(f"could not build the example company:\n{p.stderr}")


def work():
    """A working directory holding the example company's files, built once."""
    global _WORK
    if _WORK is not None:
        return _WORK
    ensure_example()
    wd = pathlib.Path(tempfile.mkdtemp(prefix="close-the-books-cli-"))
    run("init", workdir=wd, expect=OK)
    for folder in ("exports", "statements", "for-review"):
        for src in sorted((EXAMPLE / folder).glob("*")):
            if src.is_file():
                shutil.copy2(src, wd / folder / src.name)
    _WORK = wd
    return wd


def profile_path(wd):
    return wd / "profiles" / "mine.local.json"


def mark_every_row(path, decision="approve"):
    """Fill the founder_decision column, the way a founder would in Excel."""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb["Review"]
    header = None
    for r in range(1, 30):
        if str(ws.cell(row=r, column=1).value or "").strip() == "Date":
            header = r
            break
    assert header is not None, f"{path} has no header row"
    n = 0
    for r in range(header + 1, ws.max_row + 1):
        if ws.cell(row=r, column=1).value is None:
            continue
        ws.cell(row=r, column=10).value = decision
        n += 1
    wb.save(path)
    return n


def answer_engagement(wd):
    for field, value in (("basis", "accrual"),
                         ("end_use", "the 2025 federal corporate return"),
                         ("deadline", "2026-04-15"),
                         ("materiality", "500.00")):
        run("answer", f"entity.{field}", value, workdir=wd, expect=OK)
    say_which_year(wd)


def say_which_year(wd):
    """Which year is being filed. Nothing counts the queue until this is said."""
    run("filing-year", "2025", "--due", "2026-04-15", workdir=wd, expect=OK)


def stop_the_automation(wd):
    """The owner turned the auto-adding rules off and said so."""
    run("intake", "--rules", "0", "--auto-add", "0", workdir=wd, expect=OK)


def close_completeness(wd):
    """Name a document against every open completeness finding, as an owner would."""
    run("completeness", workdir=wd)
    path = wd / "reports" / "completeness.json"
    if not path.exists():
        return
    data = json.loads(path.read_text(encoding="utf-8"))
    for finding in data["findings"]:
        if not finding["open"]:
            continue
        run("completeness", "--settle", finding["account"],
            "--evidence", f"statement for {finding['account']}, "
                          f"the wire in on 2025-11-04", workdir=wd)


def ready_to_file(wd):
    say_which_year(wd)
    stop_the_automation(wd)
    close_completeness(wd)


# --------------------------------------------------------------------- init

def test_init_creates_the_layout():
    wd = pathlib.Path(tempfile.mkdtemp(prefix="close-the-books-init-"))
    code, out, _ = run("init", workdir=wd, expect=OK)
    for folder in ("exports", "statements", "for-review", "review", "import",
                   "answers", "questions", "reports", "handoff", "profiles"):
        assert (wd / folder).is_dir(), f"init did not create {folder}/"
        assert folder in out, f"init did not say what goes in {folder}/"
    # It must be safe to run twice: a founder who reruns it should not lose work.
    (wd / "exports" / "keep.txt").write_text("keep me")
    run("init", workdir=wd, expect=OK)
    assert (wd / "exports" / "keep.txt").exists(), "init destroyed an existing file"
    shutil.rmtree(wd, ignore_errors=True)


def test_init_says_nothing_posts_without_the_owner():
    wd = pathlib.Path(tempfile.mkdtemp(prefix="close-the-books-init2-"))
    _, out, _ = run("init", workdir=wd, expect=OK)
    assert "approve" in out.lower()
    assert "quickbooks" in out.lower()
    shutil.rmtree(wd, ignore_errors=True)


# -------------------------------------------------------------------- learn

def test_learn_writes_a_loadable_profile():
    wd = work()
    code, out, err = run("learn", workdir=wd, expect=OK)
    path = profile_path(wd)
    assert path.exists(), f"learn wrote no profile at {path}"

    from closethebooks import profile as profile_mod
    prof = profile_mod.load(path)          # raises if it is not usable
    assert prof.entity.name == "Acme Robotics Inc."
    assert len(prof.chart) > 20, f"only {len(prof.chart)} accounts in the chart"
    assert prof.what_goes_where, "no rules were mined from the history"
    for rule in prof.what_goes_where:
        assert rule.source, f"rule {rule.id} names no evidence"
    assert prof.accounts, "no bank or card accounts were declared"


def test_learn_prints_a_replay_accuracy():
    wd = work()
    _, out, _ = run("learn", workdir=wd, expect=OK)
    assert "REPLAY ACCURACY" in out, "learn did not report the replay accuracy"
    assert "%" in out
    # The number is the point, so it has to be a real measurement with a
    # denominator, not the word "accurate".
    assert " of " in out


def test_learn_reports_where_activity_stops_and_which_feeds_are_dead():
    wd = work()
    _, out, _ = run("learn", workdir=wd, expect=OK)
    lowered = out.lower()
    assert "activity stops" in lowered
    assert "feed" in lowered
    assert "2026-02" in out, "the last month with activity is not reported"


def test_learn_asks_for_the_engagement_facts():
    wd = work()
    _, out, _ = run("learn", workdir=wd, expect=OK)
    for field in ("end_use", "deadline", "materiality"):
        assert f"entity.{field}" in out, f"learn did not ask for {field}"


# ------------------------------------------------------------------ catchup

def test_catchup_refuses_until_the_engagement_is_answered():
    wd = work()
    run("learn", workdir=wd, expect=OK)
    code, out, err = run("catchup", "--batch-size", "150", workdir=wd, expect=REFUSAL)
    combined = out + err
    assert "end_use" in combined
    assert "Traceback" not in combined


def test_catchup_writes_a_review_workbook():
    wd = work()
    answer_engagement(wd)
    code, out, err = run("catchup", "--batch-size", "150", workdir=wd, expect=OK)
    wb = wd / "review" / "batch-01.xlsx"
    assert wb.exists(), "catchup wrote no batch-01.xlsx"

    from closethebooks import review_workbook
    rows = review_workbook.read_decisions(wb)
    assert rows, "the review workbook has no rows"
    assert len(rows) <= 150, f"batch-01 holds {len(rows)} rows, over the batch size"
    for row in rows:
        assert row["action"] in ("add", "match", "transfer", "question"), row["action"]
        # No account and no class is ever guessed: a row proposing to add
        # something must carry the rule that says where it goes.
        if row["action"] == "add":
            assert row["why"], f"row {row['row']} proposes an account with no reason"

    # Progress is stated with its denominator, never as a bare count.
    assert " of " in out
    assert "254" in out, "the queue count is not reported"


def test_catchup_flags_a_descriptor_that_is_aimed_at_an_automated_system():
    wd = work()
    answer_engagement(wd)
    _, out, _ = run("catchup", "--batch-size", "150", workdir=wd, expect=OK)
    assert "needs a human" in out.lower(), (
        "the example plants an instruction-shaped bank descriptor and catchup did "
        "not flag it")


def test_catchup_says_what_it_set_aside_and_why_each_group_waits():
    """The queue is not the job. What this deadline needs is the job."""
    wd = work()
    answer_engagement(wd)
    code, out, _ = run("catchup", "--batch-size", "150", workdir=wd, expect=OK)
    assert "this deadline needs" in out
    assert "set aside for now" in out
    assert "belongs to the 2026 return" in out
    deferred = wd / "reports" / "deferred.md"
    assert deferred.exists(), "the rows it set aside were not written out"
    text = deferred.read_text(encoding="utf-8")
    assert "Set aside, not dropped" in text
    scope = json.loads((wd / "reports" / "scope.json").read_text(encoding="utf-8"))
    assert scope["in_scope"] + scope["deferred"] == scope["total"]
    assert scope["deferred"] > 0, "the example carries rows dated after 2025"


# ------------------------------------------------- the approval gate

def test_build_imports_refuses_before_approval():
    """THE test. Nothing reaches import/ that a person has not approved."""
    wd = work()
    before = sorted(p.name for p in (wd / "import").glob("*"))
    code, out, err = run("build-imports", "--batch", "batch-01",
                         workdir=wd, expect=REFUSAL)

    combined = out + err
    assert "not been approved" in combined, (
        f"the refusal does not say why:\n{combined}")
    assert "approve batch-01" in combined, (
        "the refusal does not say what would change the answer")
    assert "Traceback" not in combined, "a refusal must never be a traceback"

    after = sorted(p.name for p in (wd / "import").glob("*"))
    assert after == before, f"files reached import/ without an approval: {after}"


def test_approve_refuses_a_workbook_with_no_decisions():
    wd = work()
    code, out, err = run("approve", "batch-01", workdir=wd, expect=REFUSAL)
    assert "no decisions" in (out + err).lower()


def test_approve_then_build_imports_succeeds():
    wd = work()
    marked = mark_every_row(wd / "review" / "batch-01.xlsx")
    assert marked > 0

    code, out, _ = run("approve", "batch-01", workdir=wd, expect=OK)
    assert (wd / "review" / "batch-01.APPROVED").exists(), "no approval was recorded"
    assert "approved" in out.lower()
    assert f"{marked:,}" in out or str(marked) in out, "the row count is not stated"
    assert "build-imports" in out, "approve does not say what happens next"

    code, out, _ = run("build-imports", "--batch", "batch-01", workdir=wd, expect=OK)
    written = sorted(p.name for p in (wd / "import").glob("*"))
    assert written, "build-imports wrote nothing after a live approval"
    assert any("batch-01" in name for name in written), (
        f"no file names its batch: {written}")
    assert " of " in out, "build-imports states a count with no denominator"


def test_an_edited_workbook_voids_its_own_approval():
    """An approval is of specific decisions, not of a filename."""
    wd = work()
    assert (wd / "review" / "batch-01.APPROVED").exists(), "run the approve test first"
    mark_every_row(wd / "review" / "batch-01.xlsx", decision="skip")
    code, out, err = run("build-imports", "--batch", "batch-01",
                         workdir=wd, expect=REFUSAL)
    combined = out + err
    assert "changed since" in combined or "stale" in combined.lower(), combined
    # Put it back, so the later tests see the state the run actually reached.
    mark_every_row(wd / "review" / "batch-01.xlsx", decision="approve")
    run("approve", "batch-01", workdir=wd, expect=OK)


def test_every_file_in_import_names_a_batch():
    wd = work()
    for path in (wd / "import").glob("*"):
        from closethebooks.approval import batch_tag
        assert batch_tag(path) is not None, (
            f"{path.name} is in import/ and names no batch, so nobody can tell what "
            f"they agreed to")


# ------------------------------------------------------ friendly failures

def test_a_command_with_no_profile_gives_a_friendly_error():
    wd = pathlib.Path(tempfile.mkdtemp(prefix="close-the-books-bare-"))
    run("init", workdir=wd, expect=OK)
    code, out, err = run("tieout", workdir=wd, expect=REFUSAL)
    combined = out + err
    assert "Traceback" not in combined, f"a missing profile produced a traceback:\n{combined}"
    assert "no profile" in combined.lower()
    assert "learn" in combined, "the error does not say how to make one"
    shutil.rmtree(wd, ignore_errors=True)


def test_a_missing_folder_gives_a_friendly_error():
    wd = pathlib.Path(tempfile.mkdtemp(prefix="close-the-books-empty-"))
    code, out, err = run("learn", workdir=wd, expect=REFUSAL)
    combined = out + err
    assert "Traceback" not in combined
    assert "exports" in combined
    assert "init" in combined
    shutil.rmtree(wd, ignore_errors=True)


def test_an_unreadable_date_is_a_refusal_and_not_a_traceback():
    wd = work()
    code, out, err = run("entries", "--through", "the-fifth-of-never",
                         workdir=wd, expect=REFUSAL)
    combined = out + err
    assert "Traceback" not in combined
    assert "date" in combined.lower()


def test_debug_adds_a_traceback_without_changing_the_exit_code():
    """--debug is accepted before or after the command, and is purely additive."""
    plain = subprocess.run(
        [sys.executable, str(BOOKS), "verify", "/nonexistent/evidence.json"],
        capture_output=True, text=True, cwd=str(REPO))
    assert plain.returncode == REFUSAL
    assert "Traceback" not in plain.stderr, "an expected condition showed a traceback"

    for args in (("--debug", "verify", "/nonexistent/evidence.json"),
                 ("verify", "/nonexistent/evidence.json", "--debug")):
        p = subprocess.run([sys.executable, str(BOOKS), *args],
                           capture_output=True, text=True, cwd=str(REPO))
        assert p.returncode == plain.returncode, (
            f"--debug changed the exit code from {plain.returncode} to "
            f"{p.returncode}; a caller switching on it would get a different answer")
        assert "Traceback" in p.stderr, f"--debug showed no traceback: {args}"


# ------------------------------------------------------------------- help

def test_help_works_for_every_subcommand():
    code, out, _ = run("--help", expect=OK)
    commands = [
        "init", "learn", "tieout", "coverage", "catchup", "approve",
        "build-imports", "fill-gaps", "entries", "reclass", "wind-down",
        "reconcile", "check", "attest", "questions", "answer", "handoff", "verify",
    ]
    for name in commands:
        assert name in out, f"{name} is missing from the top-level help"
        code, text, err = run(name, "--help", expect=OK)
        assert text.strip(), f"{name} --help printed nothing"
        assert "usage:" in text, f"{name} --help has no usage line"


def test_help_reads_like_english():
    """A non-programmer has to be able to read this."""
    _, out, _ = run("--help", expect=OK)
    for phrase in ("QuickBooks", "exit codes", "approve"):
        assert phrase in out, f"the help never mentions {phrase!r}"
    # No bare identifiers standing in for a description.
    assert "cmd_" not in out and "argparse" not in out


def test_no_arguments_prints_the_help_rather_than_an_error():
    p = subprocess.run([sys.executable, str(BOOKS)], capture_output=True,
                       text=True, cwd=str(REPO))
    assert p.returncode == OK
    assert "usage:" in p.stdout


# ------------------------------------------------------------- exit codes

def test_exit_codes_are_as_specified():
    wd = work()

    # 0: it ran and what it checked passed.
    run("learn", workdir=wd, expect=OK)

    # 2: it refused, because something it needs is missing.
    run("approve", "batch-does-not-exist", workdir=wd, expect=REFUSAL)
    run("build-imports", "--batch", "batch-02", workdir=wd, expect=REFUSAL)
    run("verify", "/nonexistent/evidence.json", workdir=wd, expect=REFUSAL)

    # 1: it ran and something is wrong. The example company's books contain
    # planted defects, so the exit tests must fail on them.
    code, out, err = run("check", workdir=wd)
    assert code == FAILURE, (
        f"check exited {code} on books with planted defects\n{out}\n{err}")
    assert "FAIL" in out


def test_a_refusal_is_distinguishable_from_a_break():
    """A caller must be able to tell 'it said no' from 'it broke'."""
    wd = work()
    refused, _, _ = run("build-imports", "--batch", "batch-02", workdir=wd)
    failed, _, _ = run("check", workdir=wd)
    assert refused == REFUSAL
    assert failed == FAILURE
    assert refused != failed


# ------------------------------------------------- the rest of the sequence

def test_tieout_reports_a_chain_break_the_example_plants():
    wd = work()
    code, out, err = run("tieout", workdir=wd)
    assert code == FAILURE, "the example is missing a statement and tieout passed"
    assert "chain break" in out.lower()
    assert " of " in out, "tieout states a count with no denominator"
    assert "2025-06" in out or "2025-05" in out


def test_coverage_names_the_month_with_no_statement():
    wd = work()
    code, out, _ = run("coverage", workdir=wd)
    assert code == FAILURE, "a month has no statement and coverage passed"
    assert "2025-06" in out
    assert "cannot check" in out.lower()


def test_reconcile_states_a_difference_for_every_cell_including_zeros():
    wd = work()
    code, out, _ = run("reconcile", "--from", "2025-01-01", "--to", "2025-12-31",
                       workdir=wd)
    assert code in (OK, FAILURE)
    assert "zeros included" in out
    assert "cannot check" in out.lower()
    assert (wd / "reports" / "reconciliation.md").exists()


def test_check_prints_a_measured_number_for_every_test():
    wd = work()
    code, out, _ = run("check", workdir=wd)
    assert code == FAILURE
    for number in range(1, 11):
        assert f"\n  {number} " in out or f"\n  {number}  " in out, (
            f"exit test {number} is missing from the table")
    assert "measured" in out
    assert (wd / "reports" / "exit-tests.md").exists()


# ------------------------------------------------- statements: asking for them

def _no_statements_workdir():
    """A working directory with exports and a profile, and an EMPTY statements/.

    Which is the normal state on the first day of a catch-up, and the state four
    of the main commands used to refuse outright in.
    """
    ensure_example()
    wd = pathlib.Path(tempfile.mkdtemp(prefix="close-the-books-nostmt-"))
    run("init", workdir=wd, expect=OK)
    for src in sorted((EXAMPLE / "exports").glob("*")):
        if src.is_file():
            shutil.copy2(src, wd / "exports" / src.name)
    run("learn", workdir=wd, expect=OK)
    say_which_year(wd)
    stop_the_automation(wd)
    assert not any((wd / "statements").iterdir()), "statements/ must start empty"
    return wd


def test_statements_names_the_accounts_the_months_and_the_reason():
    wd = work()
    code, out, _ = run("statements", workdir=wd, expect=OK)
    assert "account(s) need statements" in out
    assert "download" in out
    # A reason per account, and the reasons differ.
    assert "feed" in out.lower()
    assert "wrong side" in out.lower()
    path = wd / "reports" / "statement-request.md"
    assert path.exists(), "the request is written, not only printed"
    written = path.read_text(encoding="utf-8")
    assert "# Statements needed" in written
    assert "Months to download" in written
    assert "## Account by account" in written
    # An account that needs nothing says so, because a list that asks for
    # everything is a list nobody works. The example holds every month for two
    # of its accounts.
    assert "need nothing" in written


def test_the_four_commands_are_reachable_with_no_statements_at_all():
    wd = _no_statements_workdir()
    try:
        # A tie-out compares the books against a statement, so it still refuses.
        # What it refuses WITH is the difference: the accounts and months to get.
        code, out, err = run("tieout", workdir=wd)
        assert code == REFUSAL
        assert "account(s) need statements" in out
        assert "download" in out
        assert (wd / "reports" / "statement-request.md").exists()
        assert "books.py check" in err, "it points at what does run without one"

        # These three do real work with no statement and say what is missing.
        for command in ("reconcile", "check"):
            code, out, _ = run(command, workdir=wd)
            assert code == FAILURE, f"{command} should report, not refuse"
            assert "No statements in" in out, command
            assert "books.py statements" in out, command
        code, out, _ = run("handoff", workdir=wd)
        assert code in (OK, FAILURE), "handoff should run, not refuse"
        assert (wd / "handoff" / "handoff.md").exists()
        assert (wd / "handoff" / "statement-request.md").exists()
        doc = (wd / "handoff" / "handoff.md").read_text(encoding="utf-8")
        assert "Statements this still needs" in doc
    finally:
        shutil.rmtree(wd, ignore_errors=True)


def test_the_trial_balance_tie_out_runs_with_no_statements():
    """Exit test 1 needs no statement, and used to be unreachable without one."""
    wd = _no_statements_workdir()
    try:
        code, out, _ = run("check", workdir=wd)
        assert code == FAILURE
        assert "trial balance" in out.lower()
        assert "\n  1 " in out or "\n  1  " in out, "exit test 1 did not run"
        assert "trial balance total" in out
    finally:
        shutil.rmtree(wd, ignore_errors=True)


# ------------------------------------------- balances read on their own side

def test_no_command_ever_prints_a_balance_as_a_negative_number():
    """A liability is 3,118,447.25 Cr, never (3,118,447.25) and never negative.

    The Difference column is exempt and is checked for separately: a difference
    is not a balance and its sign says which way two figures disagree.
    """
    wd = work()
    for command, args in (("reconcile", ("--from", "2025-01-01", "--to", "2025-12-31")),
                          ("check", ()),
                          ("statements", ()),
                          ("scope", ()),
                          ("completeness", ()),
                          ("filing-year", ()),
                          ("intake", ()),
                          ("merge-plan", ())):
        _, out, _ = run(command, *args, workdir=wd)
        for line in out.splitlines():
            if "difference" in line.lower() or "out by" in line.lower():
                continue
            for token in re.findall(r"[-\u2212]\s?\d[\d,]*\.\d{2}", line):
                raise AssertionError(
                    f"`books.py {command}` printed a balance as a negative "
                    f"number: {token!r} in {line!r}")
    _, out, _ = run("check", workdir=wd)
    assert " Cr" in out and " Dr" in out, "balances carry their side"


def test_a_balance_on_the_wrong_side_is_a_named_finding_in_the_terminal():
    wd = work()
    _, out, _ = run("check", workdir=wd)
    assert "Balances on the wrong side" in out
    # The example plants a card that goes to a debit balance. It is named, its
    # normal side is stated, and the finding says what settles it.
    assert "credit-normal" in out
    assert "card statement" in out.lower() or "bank statement" in out.lower()
    assert "settles it" in out
    written = (wd / "reports" / "exit-tests.md").read_text(encoding="utf-8")
    assert "Balances on the wrong side" in written


def test_attest_records_a_person_and_a_date():
    wd = work()
    ready_to_file(wd)
    code, out, _ = run("attest", "--unbooked", "0",
                       "--note", "checked every account, For Review empty",
                       workdir=wd, expect=OK)
    path = wd / "answers" / "attestation.json"
    assert path.exists()
    data = json.loads(path.read_text())
    assert data["unbooked_count"] == 0
    assert data["by"], "nobody is named as saying it"
    assert data["on"], "the attestation carries no date"


def test_questions_asks_at_most_ten_and_states_the_denominator():
    wd = work()
    code, out, _ = run("questions", "--round", "1", workdir=wd, expect=OK)
    path = wd / "questions" / "round-1.md"
    assert path.exists()
    text = path.read_text()
    asked = text.count("\n## ")
    assert 0 < asked <= 10, f"round 1 asks {asked} questions"
    assert " of " in out, "the question count has no denominator"


def test_round_two_does_not_open_while_round_one_is_open():
    wd = work()
    code, out, err = run("questions", "--round", "2", workdir=wd, expect=REFUSAL)
    assert "round 1" in (out + err).lower()


def test_answer_records_the_date_and_who_said_it():
    wd = work()
    ids = json.loads((wd / "questions" / "round-1.json").read_text())["ids"]
    code, out, _ = run("answer", ids[0], "a customer invoice", "--on", "2026-03-01",
                       "--by", "the owner", workdir=wd, expect=OK)
    log = wd / "answers" / "answers.jsonl"
    assert log.exists()
    last = json.loads(log.read_text().strip().split("\n")[-1])
    assert last["id"] == ids[0]
    assert last["on"] == "2026-03-01"
    assert last["by"] == "the owner"


def test_wind_down_refuses_and_names_what_is_unanswered():
    wd = work()
    code, out, err = run("wind-down", workdir=wd, expect=REFUSAL)
    combined = out + err
    for key in ("target_date", "receivable_plan", "safe_terms", "cap_table",
                "resolution_date"):
        assert key in combined, f"wind-down did not name {key}"
    assert "Traceback" not in combined


def test_fill_gaps_refuses_to_write_imports_without_an_approval():
    wd = work()
    before = sorted(p.name for p in (wd / "import").glob("*"))
    code, out, err = run("fill-gaps", workdir=wd, expect=REFUSAL)
    combined = out + err
    assert "approve batch-gaps" in combined, combined
    after = sorted(p.name for p in (wd / "import").glob("*"))
    assert after == before, f"fill-gaps wrote into import/ unapproved: {after}"
    assert (wd / "review" / "batch-gaps.xlsx").exists(), (
        "fill-gaps refused without leaving anything to approve")


def test_fill_gaps_writes_upload_files_once_approved():
    wd = work()
    mark_every_row(wd / "review" / "batch-gaps.xlsx")
    run("approve", "batch-gaps", workdir=wd, expect=OK)
    code, out, _ = run("fill-gaps", workdir=wd, expect=OK)
    csvs = sorted(p.name for p in (wd / "import").glob("bank-*batch-gaps.csv"))
    assert csvs, "no upload files were written after the approval"
    # newline="" so the CRLF QuickBooks wants survives being read back.
    with open(wd / "import" / csvs[0], newline="", encoding="utf-8") as fh:
        first = fh.read()
    assert first.startswith("Date,Description,Amount\r\n"), (
        f"the upload file is not in the three-column layout QuickBooks accepts: "
        f"{first[:60]!r}")
    assert " of " in out


def test_handoff_writes_the_evidence_and_verify_reads_it():
    wd = work()
    ready_to_file(wd)
    code, out, err = run("handoff", "--out", str(wd / "handoff"), workdir=wd)
    assert code in (OK, FAILURE), (code, err)
    for name in ("evidence.json", "exit-tests.md", "change-log.md", "handoff.md"):
        assert (wd / "handoff" / name).exists(), f"handoff wrote no {name}"
    assert list((wd / "handoff").glob("*-handoff.zip")), "no archive was written"

    ledger = json.loads((wd / "handoff" / "evidence.json").read_text())
    assert ledger["sources"], "the evidence ledger declares no sources"
    assert ledger["figures"], "the evidence ledger declares no figures"
    for source in ledger["sources"]:
        completeness = source["completeness"]
        assert completeness["expected"] is not None
        assert completeness["returned"] is not None
    for figure in ledger["figures"]:
        assert figure["derivation"], f"{figure['label']} has no derivation"

    # verify reads it, and reports a failure as a failure. The example company's
    # clearing account is not zero, so this is expected to be a stop.
    code, out, _ = run("verify", str(wd / "handoff" / "evidence.json"))
    assert code in (OK, FAILURE)
    assert " of " in out, "verify states a count with no denominator"


def test_the_change_log_names_the_approval_each_decision_came_under():
    wd = work()
    text = (wd / "handoff" / "change-log.md").read_text()
    assert "batch-01" in text
    assert "Approved by" in text, "the change log does not name the approval"


# -------------------------------------------------------------------- main

def main():
    fns = [v for k, v in sorted(globals().items())
           if k.startswith("test_") and callable(v)]
    # The sequence tests build on each other, so they run in the order a real
    # catch-up runs rather than alphabetically.
    order = [
        "test_init_creates_the_layout",
        "test_init_says_nothing_posts_without_the_owner",
        "test_learn_writes_a_loadable_profile",
        "test_learn_prints_a_replay_accuracy",
        "test_learn_reports_where_activity_stops_and_which_feeds_are_dead",
        "test_learn_asks_for_the_engagement_facts",
        "test_catchup_refuses_until_the_engagement_is_answered",
        "test_catchup_writes_a_review_workbook",
        "test_catchup_flags_a_descriptor_that_is_aimed_at_an_automated_system",
        "test_catchup_says_what_it_set_aside_and_why_each_group_waits",
        "test_build_imports_refuses_before_approval",
        "test_approve_refuses_a_workbook_with_no_decisions",
        "test_approve_then_build_imports_succeeds",
        "test_an_edited_workbook_voids_its_own_approval",
        "test_every_file_in_import_names_a_batch",
        "test_a_command_with_no_profile_gives_a_friendly_error",
        "test_a_missing_folder_gives_a_friendly_error",
        "test_an_unreadable_date_is_a_refusal_and_not_a_traceback",
        "test_debug_adds_a_traceback_without_changing_the_exit_code",
        "test_help_works_for_every_subcommand",
        "test_help_reads_like_english",
        "test_no_arguments_prints_the_help_rather_than_an_error",
        "test_exit_codes_are_as_specified",
        "test_a_refusal_is_distinguishable_from_a_break",
        "test_tieout_reports_a_chain_break_the_example_plants",
        "test_coverage_names_the_month_with_no_statement",
        "test_reconcile_states_a_difference_for_every_cell_including_zeros",
        "test_check_prints_a_measured_number_for_every_test",
        "test_statements_names_the_accounts_the_months_and_the_reason",
        "test_the_four_commands_are_reachable_with_no_statements_at_all",
        "test_the_trial_balance_tie_out_runs_with_no_statements",
        "test_no_command_ever_prints_a_balance_as_a_negative_number",
        "test_a_balance_on_the_wrong_side_is_a_named_finding_in_the_terminal",
        "test_attest_records_a_person_and_a_date",
        "test_questions_asks_at_most_ten_and_states_the_denominator",
        "test_round_two_does_not_open_while_round_one_is_open",
        "test_answer_records_the_date_and_who_said_it",
        "test_wind_down_refuses_and_names_what_is_unanswered",
        "test_fill_gaps_refuses_to_write_imports_without_an_approval",
        "test_fill_gaps_writes_upload_files_once_approved",
        "test_handoff_writes_the_evidence_and_verify_reads_it",
        "test_the_change_log_names_the_approval_each_decision_came_under",
    ]
    by_name = {f.__name__: f for f in fns}
    missing = [n for n in by_name if n not in order]
    assert not missing, f"tests not in the run order: {missing}"

    failed = 0
    for name in order:
        fn = by_name[name]
        try:
            fn()
            print(f"PASS {name}")
        except Exception as exc:  # noqa: BLE001
            failed += 1
            print(f"FAIL {name}: {type(exc).__name__}: {exc}")
    if _WORK is not None and not os.environ.get("BOOKS_KEEP_WORKDIR"):
        shutil.rmtree(_WORK, ignore_errors=True)
    elif _WORK is not None:
        print(f"\nworking directory kept at {_WORK}")
    print(f"\n{len(order) - failed}/{len(order)} passed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
